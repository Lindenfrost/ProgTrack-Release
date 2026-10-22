"""Focused, disposable Core/dialog/Plots witnesses for #247, #265–#273 and #293–#296.

No live backend, launcher, or Heritage implementation is exercised.
"""

from __future__ import annotations

import copy
import importlib.util
import inspect
import json
import os
import pickle
import traceback
import tempfile
import unittest
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from time import perf_counter
from unittest.mock import patch

import numpy as np
from matplotlib.colors import to_rgba
from matplotlib.markers import MarkerStyle

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PyQt6.QtCore import Qt, QEvent
from PyQt6 import sip
from PyQt6.QtGui import QIcon, QPixmap, QValidator
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QDialog, QComboBox, QLineEdit, QTabWidget,
    QTabBar, QWidget, QPushButton, QListWidget, QListWidgetItem, QVBoxLayout,
    QGroupBox,
    QScrollArea, QCheckBox, QRadioButton,
    QAbstractItemView,
)
from matplotlib._pylab_helpers import Gcf
from Plugins.core.measurement_input import MeasurementValidator, parse_measurement, format_measurement
from Plugins.core.role_block_presets import RoleBlockPresetRegistry, normalize_event_definition
from Plugins.core.backend import ProgTrackBackend
from tests.test_current_backend_runtime import runtime_paths

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('progtrack_phase023_repairs', ROOT / 'ProgTrack.v.0.2.3.py')
MAIN = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MAIN)


class Records:
    def __init__(self):
        self.data = {}

    def get(self, namespace, key, default=None):
        return copy.deepcopy(self.data.get((namespace, key), default))

    def put(self, namespace, key, value, **kwargs):
        self.data[namespace, key] = copy.deepcopy(value)
        return 1


class CoreRepairTests(unittest.TestCase):
    def test_2s_156_all_editor_locales_parent_marker_save_and_reload(self):
        """Real Core dialogs/SQLite; locale/role checks do not need human repetition."""
        app = self._307_sqlite_host(population=0)
        app.master_track.user_db.add_user('test-5', 'Disposable-2S-only!', role='master')
        app.master_track._current_username = 'test-5'
        app.master_track._current_role = 'master'
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app)
        app._resolve_animal_key = MAIN.ProgTrackApp._resolve_animal_key.__get__(app)
        app.has_heritage_plugin = False
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'parenting']
        identities = {name: MAIN.animal_identity_key(name, 'Mus musculus',
            '01.01.2020' if name == 'test-1' else '01.01.2010', 'Institute')
            for name in ('test-1', 'test-2', 'test-3', 'test-4')}
        cases = list(self.weight_editor_cases()) + [('_dlg_female_animal', MAIN.Role.AMME.value)]
        for language in ('en', 'de', 'it', 'ru'):
            app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
            for method, role in cases:
                with self.subTest(language=language, editor=method, role=role):
                    raw = app.backend.load_core_data()
                    child = self.record(role)
                    child.update(name='test-1', ipid='IPID-2S-1', id='ID-2S-1',
                                 species='Mus musculus', eizellspenderin=identities['test-2'],
                                 samenspender=identities['test-3'])
                    raw['animals'][identities['test-1']] = app._serialize_core_record(identities['test-1'], child)
                    for key, sex, archived, died in (
                        ('test-2', 'Female', True, False),
                        ('test-3', 'Male', False, False),
                        ('test-4', 'Female', True, True),
                    ):
                        parent = {'name': key, 'ipid': f'IPID-2S-{key}', 'id': f'ID-2S-{key}',
                                  'species': 'Mus musculus', 'sex': sex, 'rolle': 'breeding_animal',
                                  'birth_date': '01.01.2010', 'origin': 'Institute', 'archived': archived}
                        if died:
                            parent['death_date'] = '01.01.2022'
                        raw['archived_animals' if archived else 'animals'][identities[key]] = parent
                    app.backend.save_core_data(raw)  # Explicit disposable scenario, not UI acceptance.
                    MAIN.ProgTrackApp._load_persistence(app)
                    kwargs = {'role_value': role} if method in (
                        '_dlg_basic_animal_role', '_dlg_versuchstier') else {}
                    if method == '_dlg_female_animal':
                        kwargs['default_role'] = role
                    dialog = self.open_dialog(app, method, name=identities['test-1'], **kwargs)
                    combo = next(c for c in dialog.findChildren(MAIN.AnimalRelationshipCombo)
                                 if c.currentData() == identities['test-2'])
                    self.assertNotIn(MAIN.DECEASED_STATUS_SYMBOL, combo.currentText())
                    index = combo.findData(identities['test-4'])
                    self.assertGreaterEqual(index, 0)
                    self.assertIn(MAIN.DECEASED_STATUS_SYMBOL, combo.itemText(index))
                    combo.setCurrentIndex(index)
                    errors = []
                    with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                        next(b for b in dialog.findChildren(QPushButton)
                             if b.text() == app.messages['button.save']).click()
                    self.assertFalse(errors or app.errors, (errors, app.errors))
                    self.assertEqual(app.backend.load_core_data()['animals'][identities['test-1']]['eizellspenderin'], identities['test-4'])
                    MAIN.ProgTrackApp._load_persistence(app)
                    reopened = self.open_dialog(app, method, name=identities['test-1'], **kwargs)
                    self.assertIn(identities['test-4'], [c.currentData() for c in reopened.findChildren(MAIN.AnimalRelationshipCombo)])
                    reopened.reject()

    def test_2s_315_316_severity_set_clear_real_audit_and_export(self):
        """Replace R02c/d/e with real editor Save, actual log parsing and PDF content."""
        from Plugins.Master_Track.dialogs import AuditLogsDialog
        from pypdf import PdfReader
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app.has_medi_track_plugin = True  # Severity is shown only with Medi Track installed.
        key = app.backend.records.get('seed', 'shire-colony')['named_animals']['Frodo']
        fixture = app.backend.load_core_data()
        fixture['animals'][key]['in_experiment'] = False
        fixture['animals'][key]['severity'] = ''
        app.backend.save_core_data(fixture)  # Separate explicit scenario setup.
        MAIN.ProgTrackApp._load_persistence(app)
        errors = []
        for desired in ('SV1', ''):
            dialog = self.open_dialog(app, '_dlg_versuchstier', name=key)
            selector = next(combo for combo in dialog.findChildren(QComboBox) if combo.findData('SV1') >= 0)
            index = selector.findData(desired) if desired else 0
            self.assertGreaterEqual(index, 0)
            selector.setCurrentIndex(index)
            with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                next(b for b in dialog.findChildren(QPushButton)
                     if b.text() == app.messages['button.save']).click()
            self.assertEqual(errors, [])
            self.assertFalse(app.errors, app.errors)
            self.assertEqual(app.backend.load_core_data()['animals'][key].get('severity', ''), desired)
        audit = AuditLogsDialog(app, app.messages, str(ROOT / 'Plugins' / 'Master_Track'),
                               audit_dir=app.master_track.audit_logs_dir())
        self.hosts.append(audit)
        rows = audit._load_rows()
        severity = [row for row in rows if row.get('_raw_parameter') == 'severity']
        self.assertEqual(len(severity), 2)
        self.assertTrue(any('SV1' in row['new'] for row in severity))
        self.assertTrue(any('SV1' in row['old'] and 'SV1' not in row['new'] for row in severity))
        self.assertTrue(all(row['user'] == 'Admin' for row in severity))
        self.assertFalse(any(row.get('_raw_parameter') == 'events' for row in rows))
        app._save_persistence()
        self.assertEqual(audit._load_rows(), rows, 'unchanged Save must not create audit rows')
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'audit-2s.pdf'
            audit._export_rows_to_pdf(str(target), rows, datetime.now().date(), datetime.now().date())
            text = ' '.join(' '.join(page.extract_text() or '' for page in PdfReader(target).pages).split())
            self.assertIn('SV1', text)
            self.assertIn('Admin', text)

    def test_2s_272_306_real_credentials_direct_switch_and_idle_warning_logout(self):
        """2S: actual authentication and Qt timer delivery, not login-return stubs."""
        from Plugins.Master_Track.dialogs import LoginDialog
        from PyQt6.QtCore import QTimer
        from PyQt6.QtTest import QTest
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        self._306_sidebar_shell(app)
        app._apply_style_settings({})
        master = app.master_track
        password = 'Disposable-2S-only!'
        for username, role in (('test-1', 'master'), ('test-2', 'user')):
            master.user_db.add_user(username, password, role=role)
            master.user_db.set_password(username, password, must_change=False)
        witness = next(iter(app.animals))
        failures = []
        def login(username, *, invalid=False):
            outgoing = master.current_username
            app.selected_animals = [witness]
            app._plot_selection_order = [witness]
            app._edit_selection_order = [witness]
            def submit():
                try:
                    dialog = QApplication.activeModalWidget()
                    self.assertIsInstance(dialog, LoginDialog)
                    dialog.username_edit.setText(username)
                    dialog.pw_edit.setText('wrong' if invalid else password)
                    dialog.login_btn.click()
                    if invalid:
                        self.assertIsNone(dialog.logged_in_user)
                        self.assertTrue(dialog.error_label.text())
                        self.assertEqual(master.current_username, outgoing)
                        dialog.reject()
                except BaseException as exc:
                    failures.append(exc)
                    modal = QApplication.activeModalWidget()
                    if modal:
                        modal.reject()
            QTimer.singleShot(0, submit)
            app._do_master_login()
            self.assertEqual(failures, [])
            if invalid:
                self.assertEqual(master.current_username, outgoing)
                self.assertEqual(app.selected_animals, [witness])
            else:
                self.assertEqual(master.current_username, username)
                for field in ('selected_animals', '_plot_selection_order', '_edit_selection_order'):
                    self.assertEqual(getattr(app, field), [])
        login('test-1')
        login('test-2', invalid=True)
        login('test-2')
        # Assert configured minute-to-ms behavior before compressing only the
        # test timers. Production timeouts/security settings are not changed.
        master._effective_timeout_minutes = 1
        master._start_idle_timer()
        self.assertEqual(master._idle_timer.interval(), 60_000)
        self.assertEqual(master._warning_timer.interval(), 45_000)
        master._warning_timer.start(5)
        master._idle_timer.start(60_000)
        deadline = perf_counter() + 5
        while master._idle_warning_dialog is None and perf_counter() < deadline:
            QTest.qWait(10)
        self.assertIsNotNone(master._idle_warning_dialog)
        keep = next(b for b in master._idle_warning_dialog.findChildren(QPushButton)
                    if b.text() == app.messages['master_track.timeout.keep_logged_in'])
        keep.click()
        self.assertTrue(master.is_logged_in)
        self.assertIsNone(master._idle_warning_dialog)
        self.assertEqual(master._idle_timer.interval(), 60_000)
        app.selected_animals = [witness]
        master._warning_timer.start(5)
        master._idle_timer.start(60)
        deadline = perf_counter() + 5
        while master.is_logged_in and perf_counter() < deadline:
            QTest.qWait(10)
        self.assertFalse(master.is_logged_in)
        self.assertEqual(app.selected_animals, [])
        self.assertFalse(master._idle_timer.isActive())
        self.assertFalse(master._warning_timer.isActive())
        log = Path(master.audit_logs_dir()) / datetime.now().strftime('audit_%Y-%m.log')
        text = log.read_text(encoding='utf-8')
        self.assertIn('[test-2] [timeout]', text)
        self.assertIn('[test-2] [logout]', text)

    def test_2s_247_316_actual_project_transition_medi_summary_and_restart(self):
        """2S corrected Q10e: real shared usage, never fake Measurement history."""
        from Plugins.Medi_Track.medi_track_widget import MediStore, MediTrackWidget
        from Plugins.Animal_Reports.report_index import AnimalReportIndex
        from Plugins.core.project_periods import current_event_counts, event_counts_by_period
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app._schedule_post_animal_save_project_updates = (
            MAIN.ProgTrackApp._schedule_post_animal_save_project_updates.__get__(app))
        key = app.backend.records.get('seed', 'shire-colony')['named_animals']['Frodo']
        record = app.animals[key]
        self.assertEqual(record['project'], 'Ringbearer')
        MAIN.ensure_project_periods(record)
        old_period = record['project_period_id']
        previous = current_event_counts(record).get('measurement', 0)
        for ordinal in range(3):
            record['events'].append(dict(event_id=f'2s-period-{ordinal}', event_type='measurement',
                datum=datetime(2025, 1, 1 if ordinal < 2 else 2),
                recorded_role=record['rolle'], project_period_id=old_period))
        app._save_persistence()
        widget = MediTrackWidget(app, MediStore(app.backend), app.messages, 'en')
        self.hosts.append(widget)
        widget.show_animal(key)
        label = app.messages.get('event.measurement', 'Measurement')
        self.assertIn(f'{label} {previous + 3}', widget._lbl_project_usage.text())
        self.assertFalse(any(e.get('status_type') == 'measurement' and
                             str(e.get('source_event_id', '')).startswith('2s-period')
                             for e in widget.store.get_entries(key)))
        record = app.animals[key]
        record.update(in_experiment=False, project='Shire')
        # This is the validated reason/date normally returned by the exit
        # dialog; the transition/history command itself is real, not stubbed.
        app._pending_experiment_exit = {
            'date': datetime.now().strftime(MAIN.DATE_FORMAT), 'reason': '2S project reassignment'}
        app._schedule_post_animal_save_project_updates(key, 'Ringbearer', 'Shire',
            old_in_experiment=True, new_in_experiment=False)
        app._save_persistence()
        record = app.animals[key]
        new_period = record['project_period_id']
        self.assertNotEqual(new_period, old_period)
        self.assertEqual(current_event_counts(record), {})
        self.assertEqual(event_counts_by_period(record)[old_period]['measurement'], previous + 3)
        widget.show_animal(key)
        self.assertIn('Shire', widget._lbl_project.text())
        self.assertIn(f'{label} {previous + 3}', widget._lbl_project_usage.text())
        record['events'].append(dict(event_id='2s-new-period', event_type='measurement',
            datum=datetime.now(), recorded_role=record['rolle'], project_period_id=new_period))
        app._save_persistence()
        self.assertEqual(current_event_counts(app.animals[key]), {'measurement': 1})
        index = AnimalReportIndex(app.animals[key], lambda token: token, date_format=MAIN.DATE_FORMAT)
        self.assertEqual(index.event_counts_through(datetime.now().date())['measurement'], (1, 1))
        before = app.backend.load_core_data()
        app.backend.close()
        app.backend = ProgTrackBackend(app.backend.paths, acquire_process_lock=False)
        self.addCleanup(app.backend.close)
        MAIN.ProgTrackApp._load_persistence(app)
        app._save_persistence()  # No warm-up or historical rewrite after reopen.
        self.assertEqual(app.backend.load_core_data(), before)
        widget.store = MediStore(app.backend)
        widget.show_animal(key)
        self.assertIn(f'{label} 1', widget._lbl_project_usage.text())
        self.assertIn(f'{label} {previous + 3}', widget._lbl_project_usage.text())

    def test_321_closed_custom_history_callback_ignores_deleted_qt_controls(self):
        app = self.host()
        scroll, frame = QScrollArea(), QWidget()
        sip.delete(scroll)
        MAIN.ProgTrackApp._cap_history_viewport(app, scroll, frame, 2)
        scroll = QScrollArea()
        sip.delete(frame)
        MAIN.ProgTrackApp._cap_history_viewport(app, scroll, frame, 2)
        scroll.deleteLater()

    def test_321_real_custom_checkbox_catalogs_normalizers_and_settings_rebuild(self):
        from Plugins.core.role_block_presets import normalize_experimental_block
        app = self._306_conventions()
        app._merge_user_messages = lambda *_a: None
        original = app.backend.records.get('configuration', 'role-block-presets')
        blocks = []
        for language in ('en', 'de', 'it', 'ru'):
            app.lang = language
            app._load_messages(language)
            for kind in ('counting', 'limiting'):
                for render in ('line', 'symbol'):
                    with self.subTest(language=language, kind=kind, render=render):
                        initial = dict(name='Test Clinical', event_type='test_clinical', kind=kind,
                                       render_mode=render, default_maximum=10)
                        def edit(dialog):
                            control = dialog.findChild(QCheckBox, 'customExperimentalBlockMirrorMedical')
                            self.assertIsNotNone(control)
                            self.assertFalse(control.isChecked())
                            self.assertEqual(control.text(), app.messages[
                                'settings.role_setup.custom_blocks.mirror_medical_history'])
                            control.setChecked(True)
                            return QDialog.DialogCode.Accepted
                        with patch.object(QDialog, 'exec', edit):
                            result = MAIN.StyleSettingsDialog._custom_experimental_block_editor(
                                SimpleNamespace(messages=app.messages), app, initial)
                        self.assertTrue(result['mirror_in_medical_history'])
                        self.assertEqual((result['kind'], result['render_mode']), (kind, render))
                        self.assertNotIn('mirror_in_medical_history', initial)
                        self.assertTrue(normalize_experimental_block(result)['mirror_in_medical_history'])
                        self.assertTrue(normalize_event_definition(result)['mirror_in_medical_history'])
                        if language == 'en':
                            result['event_type'] = f'test_{kind}_{render}'
                            result['name'] = result['event_type']
                            result.pop('id', None)
                            result['mirror_in_medical_history'] = render == 'symbol'
                            blocks.append(normalize_experimental_block(result))
        with self.assertRaisesRegex(ValueError, 'boolean'):
            normalize_event_definition(dict(name='Invalid', mirror_in_medical_history='false'))
        prior = next(block for block in blocks if block['mirror_in_medical_history'])
        def disable(dialog):
            dialog.findChild(QCheckBox, 'customExperimentalBlockMirrorMedical').setChecked(False)
            return QDialog.DialogCode.Accepted
        with patch.object(QDialog, 'exec', disable):
            changed = MAIN.StyleSettingsDialog._custom_experimental_block_editor(
                SimpleNamespace(messages=app.messages), app, prior)
        self.assertEqual(changed['id'], prior['id'])
        self.assertEqual(changed['stable_id'], prior['stable_id'])
        self.assertEqual(changed['event_type'], prior['event_type'])
        self.assertEqual(changed['revision'], prior['revision'] + 1)
        app.lang = 'en'
        app._load_messages('en')
        dialog = MAIN.StyleSettingsDialog(app, app.messages)
        self.hosts.append(dialog)
        dialog._custom_experimental_blocks.extend(blocks)
        dialog.reject()
        self.assertEqual(app.backend.records.get('configuration', 'role-block-presets'), original)
        dialog = MAIN.StyleSettingsDialog(app, app.messages)
        self.hosts.append(dialog)
        dialog._custom_experimental_blocks.extend(blocks)
        dialog.accept()
        app._apply_accepted_conventions(dialog)
        self.assertFalse(app.errors, app.errors)
        registry = RoleBlockPresetRegistry(app.backend)
        definitions = {item['event_type']: item for item in registry.event_definitions()}
        for block in blocks:
            self.assertEqual(definitions[block['event_type']]['mirror_in_medical_history'],
                             block['render_mode'] == 'symbol')
        reloaded = MAIN.StyleSettingsDialog(app, app.messages)
        self.hosts.append(reloaded)
        self.assertEqual(reloaded._prepared_preset_settings(), registry.prepare_payload(
            registry.presets(), experimental_blocks=registry.experimental_blocks(),
            event_definitions=registry.event_definitions()))
        before = app.backend.records.get('configuration', 'role-block-presets')
        with patch.object(app, '_can_configure_animal_roles', return_value=False), \
             patch.object(app.backend.records, 'put_many') as write:
            self.assertFalse(app._save_animal_role_definitions(
                app.animal_role_registry.roles(), preset_payload=reloaded._prepared_preset_settings()))
            write.assert_not_called()
        self.assertEqual(app.backend.records.get('configuration', 'role-block-presets'), before)

    def test_321_real_clinical_save_adapters_and_retained_annotations(self):
        questions = patch.object(MAIN.QMessageBox, 'question', return_value=MAIN.QMessageBox.StandardButton.Yes)
        questions.start()
        self.addCleanup(questions.stop)
        from Plugins.Medi_Track.medi_track_widget import MediStore
        from Plugins.core.backend.clinical_mirroring import NAMESPACE
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        # Use the production shared configuration, not a fabricated destination row.
        definition = normalize_event_definition(dict(name='Test Clinical', event_type='test_clinical',
            kind='counting', mirror_in_medical_history=True))
        app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [definition]})
        cases = [('_dlg_offspring', MAIN.Role.OFFSPRING.value),
                 ('_dlg_zuchttier', MAIN.Role.ZUCHTTIER.value),
                 ('_dlg_versuchstier', MAIN.Role.EXPERIMENTAL.value),
                 ('_dlg_versuchstier', MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING),
                 ('_dlg_female_animal', MAIN.Role.SPENDER.value),
                 ('_dlg_basic_animal_role', 'custom_role')]
        for method, role in cases:
            with self.subTest(adapter=method, role=role):
                if role == 'custom_role':
                    key = app.backend.records.get('seed', 'shire-colony')['named_animals']['Frodo']
                    raw = app.backend.load_core_data()
                    raw['animals'][key]['rolle'] = role
                    app.backend.save_core_data(raw)
                    MAIN.ProgTrackApp._load_persistence(app)
                    app._active_custom_event_definitions_for_role = lambda *_a, **_k: [definition]
                else:
                    key = next(k for k, r in app.animals.items() if app._animal_role_value(r) == role)
                before = copy.deepcopy(app.backend.load_core_data()['animals'][key])
                kwargs = {'role_value': role} if method in ('_dlg_versuchstier', '_dlg_basic_animal_role') else {}
                dialog = self.open_dialog(app, method, name=key, **kwargs)
                label_key = {'_dlg_offspring': 'dialog.offspring.button.new_event',
                    '_dlg_zuchttier': 'dialog.zuchttier.button.new_event',
                    '_dlg_versuchstier': 'dialog.versuchstier.button.new_event',
                    '_dlg_basic_animal_role': 'dialog.basic_role.button.new_event'}.get(method)
                label = app.messages[label_key] if label_key else app.messages['button.new_item'].format(
                    title=app.messages.get('dialog.female_animal.tab.events', 'Events'))
                next(b for b in dialog.findChildren(QPushButton) if b.text() == label).click()
                date, combo, _delete = self.event_editor_rows_247(dialog)['']
                token = 'test_clinical' if role == 'custom_role' else 'surgery'
                index = combo.findData(token)
                if index < 0 and role != 'custom_role':
                    token = next(combo.itemData(i) for i in range(combo.count())
                                 if combo.itemData(i) in ('pregnancy', 'birth', 'embryo_transfer'))
                    index = combo.findData(token)
                self.assertGreaterEqual(index, 0, [combo.itemData(i) for i in range(combo.count())])
                combo.setCurrentIndex(index)
                date.setText('01.10.2026')
                errors = []
                with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                    next(b for b in dialog.findChildren(QPushButton) if b.text() == app.messages['button.save']).click()
                self.assertFalse(errors or app.errors, (errors, app.errors))
                entries = MediStore(app.backend).get_entries(key)
                clinical = [entry for entry in entries if entry.get('source_namespace') == 'core.events']
                self.assertEqual(len(clinical), 1)
                self.assertEqual(clinical[0]['status_type'], token)
                self.assertEqual(clinical[0]['actor_login'], 'Admin')
                original = copy.deepcopy(clinical[0])
                new = next(e for e in app.animals[key]['events']
                           if e['event_id'] not in {r['event_id'] for r in before['events']})
                # Real editor correction then deletion must preserve canonical metadata.
                for delete in (False, True):
                    dialog = self.open_dialog(app, method, name=key, **kwargs)
                    date, _combo, remove = self.event_editor_rows_247(dialog)[new['event_id']]
                    if delete:
                        with patch.object(app, '_show_message_raw', return_value=MAIN.QMessageBox.StandardButton.Yes):
                            remove.click()
                    else:
                        date.setText('02.10.2026')
                    with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                        next(b for b in dialog.findChildren(QPushButton) if b.text() == app.messages['button.save']).click()
                    self.assertFalse(errors or app.errors, (errors, app.errors))
                ledger = app.backend.records.get(NAMESPACE, key)
                self.assertEqual(ledger['entries'][0], original)
                self.assertEqual([e['entry_type'] for e in ledger['entries']],
                                 ['core_event', 'core_correction', 'core_deletion'])

        # Original 2S S04f: use the actual Edit Role command to restore the
        # built-in role after the custom source was corrected/deleted.
        before_role = copy.deepcopy(app.backend.load_core_data()['animals'][key])
        before_ledger = copy.deepcopy(app.backend.records.get(NAMESPACE, key))
        from Plugins.core.animal_roles import AnimalRoleRegistry
        app.animal_role_registry = AnimalRoleRegistry(app.backend)
        app.category_tab = QTabBar(app)
        for label in range(8):
            app.category_tab.addTab(str(label))
        def restore_role(dialog):
            combo = dialog.findChild(QComboBox)
            index = combo.findData(MAIN.Role.EXPERIMENTAL.value)
            self.assertGreaterEqual(index, 0)
            combo.setCurrentIndex(index)
            return QDialog.DialogCode.Accepted
        with patch.object(QDialog, 'exec', restore_role):
            app._dlg_change_sort(key)
        self.assertFalse(app.errors, app.errors)
        restored = app.backend.load_core_data()['animals'][key]
        self.assertEqual(restored['rolle'], MAIN.Role.EXPERIMENTAL.value)
        self.assertEqual(restored['events'], before_role['events'])
        self.assertEqual(restored['gewicht'], before_role['gewicht'])
        self.assertEqual(app.backend.records.get(NAMESPACE, key), before_ledger)

    def test_321_clinical_table_pdf_xlsx(self):
        from Plugins.Medi_Track.medi_track_widget import MediStore, MediTrackWidget
        from tests.test_current_backend_runtime import CurrentBackendRuntimeTests
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from PyQt6.QtGui import QFontDatabase
        from matplotlib import get_data_path
        # Windows' offscreen platform does not enumerate system fonts: without
        # an explicit application font Qt prints tofu outlines, not PDF text.
        font_path = Path(os.environ.get('WINDIR', '')) / 'Fonts' / 'arial.ttf'
        if not font_path.is_file():
            font_path = Path(get_data_path()) / 'fonts' / 'ttf' / 'DejaVuSans.ttf'
        font_id = QFontDatabase.addApplicationFont(str(font_path))
        self.assertGreaterEqual(font_id, 0)
        self.addCleanup(QFontDatabase.removeApplicationFont, font_id)
        old_font = self.qt.font()
        self.qt.setFont(MAIN.QFont('Arial', 10))
        self.addCleanup(self.qt.setFont, old_font)
        app = self._307_sqlite_host(population=2)
        definition = normalize_event_definition(dict(name='Test Clinical', event_type='test_clinical',
            kind='counting', mirror_in_medical_history=True))
        app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [definition]})
        actor = CurrentBackendRuntimeTests._321_actor()
        for stage in ('original', 'correction', 'deletion'):
            snapshot = app.backend.load_core_data()
            key = next(iter(snapshot['animals']))
            events = snapshot['animals'][key]['events']
            if stage == 'original':
                events.append(dict(event_id='export321', event_type='test_clinical', datum='2026-01-01'))
            elif stage == 'correction':
                next(e for e in events if e['event_id'] == 'export321')['datum'] = '2026-01-02'
            else:
                events[:] = [e for e in events if e['event_id'] != 'export321']
            app.backend.save_core_data(snapshot, clinical_context=actor)
        app.animals = app.backend.load_core_data()['animals']
        original = app.backend.records.get('medical.core-events', key)['entries'][0]
        # Actual table and real export readers, not expected destination injection.
        widget = MediTrackWidget(app, MediStore(app.backend), app.messages, 'en')
        self.hosts.append(widget)
        widget.show_animal(key)
        self.assertGreaterEqual(widget._table.rowCount(), 3)
        text = '\n'.join(widget._table.item(row, col).text()
                         for row in range(widget._table.rowCount()) for col in range(5))
        self.assertIn(app.messages['medi_track.core_correction'], text)
        self.assertIn(original['id'], text)
        with tempfile.TemporaryDirectory() as directory:
            pdf, xlsx = Path(directory) / 'clinical.pdf', Path(directory) / 'clinical.xlsx'
            widget._export_animal_to_pdf(key, str(pdf))
            widget._export_animal_to_xlsx(key, str(xlsx), include_documents=False)
            reader = PdfReader(pdf)
            pdf_text = '\n'.join(p.extract_text() for p in reader.pages)
            self.assertTrue(pdf_text.strip(), repr(reader.pages[0].get_contents().get_data()[:1800]))
            book = load_workbook(xlsx)
            sheet_text = '\n'.join(str(v) for row in book.active.values for v in row if v is not None)
            book.close()
            for result in (pdf_text, sheet_text):
                readable = ' '.join(result.split())
                self.assertIn('Test Clinical', readable)
                self.assertIn(app.messages['medi_track.core_correction'], readable)
                self.assertIn(app.messages['medi_track.core_deletion'], readable)
                self.assertIn(original['id'], ''.join(result.split()), 'PDF line wrapping must not lose source identity')
                for value in ('2026-01-01', '2026-01-02', 'export321'):
                    self.assertIn(value, readable)

    def test_319_catalog_coverage_production_loader_and_literal_overrides(self):
        captions = {'en': 'New event', 'de': 'Neues Ereignis',
                    'it': 'Nuovo evento', 'ru': 'Новое событие'}
        key = 'dialog.basic_role.button.new_event'
        app = self.host()
        for language, expected in captions.items():
            with self.subTest(language=language):
                def unique_keys(pairs):
                    counts = Counter(name for name, _value in pairs)
                    for changed in (key, 'table.header.event_type', 'op_planner.label.event_type',
                                    'settings.role_setup.custom_blocks.event_type',
                                    'surgery_planner.label.event_type'):
                        self.assertEqual(counts[changed], 1, (language, changed))
                    return dict(pairs)
                catalog = json.loads((ROOT / f'lang/messages_{language}.json').read_text(
                    encoding='utf-8'), object_pairs_hook=unique_keys)
                self.assertEqual(catalog[key], expected)
                app.backend.records.put('configuration', f'user-message-overrides:{language}',
                                        {'role.experimental': 'Test-Role 2Q'})
                stored = copy.deepcopy(app.backend.records.data)
                app.lang = language
                app._load_messages(language)
                self.assertEqual(app.messages[key], expected)
                self.assertEqual(app.messages['role.experimental'], 'Test-Role 2Q')
                self.assertEqual(app.backend.records.data, stored)
                if language == 'de':
                    for label_key in ('table.header.event_type', 'op_planner.label.event_type',
                                      'settings.role_setup.custom_blocks.event_type'):
                        self.assertEqual(app.messages[label_key], 'Ereignisart')
                    self.assertEqual(app.messages['surgery_planner.label.event_type'], 'Ereignisart:')

    def test_319_custom_editors_rebuild_localized_controls_without_mutating_history(self):
        app = self.host()
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom_audit_event']
        definition = normalize_event_definition({
            'event_type': 'custom_audit_event', 'block_id': 'custom_audit_event',
            'name': 'Test-Event 2Q', 'kind': 'counting', 'active': True})
        app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [definition]})
        animal = self.record('custom_role')
        animal['events'] = [{'event_id': 'stable-319', 'event_type': 'custom_audit_event',
                             'datum': datetime(2026, 1, 1), 'recorded_role': 'custom_role'}]
        app.animals['Subject'] = animal
        # A normal initialized app already has its identity convention catalog.
        app._load_identity_conventions()
        before = copy.deepcopy((app.animals, app.archived, app.backend.records.data))
        for language in ('en', 'de', 'it', 'ru', 'en'):
            with self.subTest(language=language):
                app.lang = language
                app._load_messages(language)
                dlg = self.open_dialog(app, '_dlg_basic_animal_role')
                dlg.show()
                tabs = dlg.findChild(QTabWidget)
                event_tab = next(tabs.widget(index) for index in range(tabs.count())
                                 if tabs.tabText(index) == app.messages['dialog.animal.tab.events'])
                tabs.setCurrentWidget(event_tab)
                self.qt.processEvents()
                add = next(button for button in event_tab.findChildren(QPushButton)
                           if button.text() == app.messages['dialog.basic_role.button.new_event'])
                self.assertTrue(add.isVisible() and add.isEnabled())
                self.assertIn(app.messages['table.header.event_type'],
                              [label.text() for label in event_tab.findChildren(MAIN.QLabel)])
                rows = self.event_editor_rows_247(dlg)
                original = rows['stable-319'][1]
                self.assertEqual(original.currentData(), 'custom_audit_event')
                self.assertEqual(original.currentText(), 'Test-Event 2Q')
                self.assertFalse(original.isEnabled())
                add.click()
                self.qt.processEvents()
                new_combo = next(combo for combo in event_tab.findChildren(QComboBox)
                                 if combo is not original)
                self.assertEqual(new_combo.currentText(), 'Test-Event 2Q')
                self.assertEqual(new_combo.currentData(), 'custom_audit_event')
                dlg.reject()
                self.assertEqual((app.animals, app.archived, app.backend.records.data), before)
                self.assertEqual(app._test_audit_entries, [])

    def test_319_builtin_event_buttons_and_headers_remain_localized(self):
        for language in ('en', 'de', 'it', 'ru'):
            for method, role, key in (
                ('_dlg_offspring', 'offspring', 'dialog.offspring.button.new_event'),
                ('_dlg_zuchttier', 'breeding_animal', 'dialog.zuchttier.button.new_event'),
                ('_dlg_versuchstier', 'experimental', 'dialog.versuchstier.button.new_event')):
                with self.subTest(language=language, editor=method):
                    app = self.host()
                    app.lang = language
                    app._load_messages(language)
                    app.animals['Subject'] = self.record(role)
                    before = copy.deepcopy(app.animals)
                    dlg = self.open_dialog(app, method)
                    self.assertIn(app.messages[key], [b.text() for b in dlg.findChildren(QPushButton)])
                    self.assertIn(app.messages['table.header.event_type'],
                                  [label.text().removeprefix('<b>').removesuffix('</b>')
                                   for label in dlg.findChildren(MAIN.QLabel)])
                    dlg.reject()
                    self.assertEqual(app.animals, before)

    def test_319_custom_block_editor_localizes_caption_not_canonical_choices(self):
        app = self.host()
        initial = {'name': 'Test-Event 2Q', 'kind': 'counting', 'event_type': 'custom.audit_event'}
        original = copy.deepcopy(initial)
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                app.lang = language
                app._load_messages(language)
                captured = []
                with patch.object(QDialog, 'exec', lambda dlg: captured.append(dlg) or 0):
                    result = MAIN.StyleSettingsDialog._custom_experimental_block_editor(
                        SimpleNamespace(messages=app.messages), app, initial)
                self.assertIsNone(result)
                self.assertEqual(initial, original)
                dlg = captured[0]
                self.hosts.append(dlg)
                combo = dlg.findChild(QComboBox, 'customExperimentalBlockEventTypeCombo')
                self.assertEqual(dlg.layout().labelForField(combo).text(),
                                 app.messages['settings.role_setup.custom_blocks.event_type'])
                self.assertEqual([combo.itemData(index) for index in range(combo.count())],
                                 ['counting', 'limiting'])
                self.assertEqual(combo.currentData(), 'counting')
                self.assertEqual(dlg.findChild(QLineEdit, 'customExperimentalBlockName').text(),
                                 'Test-Event 2Q')

    def _318_admin_host(self):
        from Plugins.Cage__Track.cage_track_widget import CageTrackPlugin
        previous_font = self.qt.font()
        self.qt.setFont(MAIN.QFont('Arial', 10))
        self.addCleanup(self.qt.setFont, previous_font)
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app.has_cage_track_plugin = True
        app.cage_track_plugin = CageTrackPlugin(app)
        app._is_projects_track_active = lambda: True
        for method in ('_validate_identity_species_for_save', '_validate_existing_identity_for_save',
                       '_name_species_conflict', '_confirm_species_change_once',
                       '_capture_death_lifecycle'):
            setattr(app, method, getattr(MAIN.ProgTrackApp, method).__get__(app))
        return app

    def _318_open(self, app, role, name=None, read_only=False):
        captured = []
        with patch.object(app, '_install_new_animal_id_preview', wraps=app._install_new_animal_id_preview) as preview, \
             patch.object(QDialog, 'exec', lambda dlg: captured.append(dlg) or 0):
            app._dialog_for_role_value(role)(name, read_only=read_only)
        self.assertEqual(len(captured), 1)
        dialog = captured[0]
        self.hosts.append(dialog)
        controls = dict(preview.call_args.kwargs, id_widget=preview.call_args.args[0])
        dialog.show()
        for _ in range(3):
            self.qt.processEvents()
        return dialog, controls

    def _318_fill_new(self, app, dialog, controls, name):
        from PyQt6.QtTest import QTest
        for key, text in (('name_widget', name), ('birth_widget', '01.01.2020')):
            widget = controls[key]
            QTest.mouseClick(widget, Qt.MouseButton.LeftButton)
            QTest.keyClicks(widget, text)
            self.assertEqual(widget.text(), text)
        species = controls['species_widget']
        index = species.findData('Callitrix jacchus')
        self.assertGreaterEqual(index, 0)
        species.setCurrentIndex(index)
        origin = controls['origin_widget']
        origin.setCurrentIndex(origin.findData('DPZ'))
        dates = [field for field in dialog.findChildren(QLineEdit)
                 if field.placeholderText() == app.messages['form.placeholder.date_short']]
        death = next(field for field in dates if field is not controls['birth_widget'])
        QTest.mouseClick(death, Qt.MouseButton.LeftButton)
        QTest.keyClicks(death, '02.01.2020')
        self.assertEqual(death.text(), '02.01.2020')
        death.selectAll()
        QTest.keyClick(death, Qt.Key.Key_Backspace)
        self.assertEqual(death.text(), '')
        return next(button for button in dialog.findChildren(QPushButton)
                    if button.text() == app.messages['button.save'])

    def test_318_moving_middle_form_rows_preserves_occupied_rows_and_layout_items(self):
        app = self.host()
        body = QWidget(app)
        form = MAIN.QFormLayout(body)
        before = QLineEdit('before')
        moved = QWidget()
        anchor = QLineEdit('anchor')
        after = QLineEdit('after')
        form.addRow('Before', before)
        form.addRow(moved)
        form.addRow('Anchor', anchor)
        form.addRow('After', after)
        app._move_form_row_after(form, moved, anchor)
        self.assertEqual(form.rowCount(), 4)
        self.assertEqual(form.getWidgetPosition(moved), (2, MAIN.QFormLayout.ItemRole.SpanningRole))
        self.assertEqual(form.getWidgetPosition(after)[0], 3)
        app._move_form_row_after(form, after, before)
        self.assertEqual(form.rowCount(), 4)
        self.assertEqual(form.getWidgetPosition(after)[0], 1)
        self.assertEqual(form.labelForField(after).text(), 'After')
        dates = MAIN.QHBoxLayout()
        dates.addWidget(QLineEdit('date one'))
        dates.addWidget(QLineEdit('date two'))
        form.addRow('Dates', dates)
        label = form.labelForField(dates)
        app._move_form_row_after(form, label, before)
        self.assertEqual(form.rowCount(), 5)
        self.assertIs(form.itemAt(1, MAIN.QFormLayout.ItemRole.FieldRole).layout(), dates)
        app._move_form_row_after(form, moved, moved)
        self.assertEqual(form.rowCount(), 5)

    def test_318_seeded_admin_new_dates_have_real_input_hit_ownership(self):
        from Plugins.core.animal_dialog_sections import AnimalDialogSection
        app = self._318_admin_host()
        for role in (MAIN.Role.EXPERIMENTAL.value, MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING):
            with self.subTest(role=role):
                dialog = self.open_dialog(app, '_dlg_versuchstier', name=None, role_value=role)
                dialog.show()
                for _ in range(3):
                    self.qt.processEvents()
                fields = [field for field in dialog.findChildren(QLineEdit)
                          if field.placeholderText() == app.messages['form.placeholder.date_short']]
                self.assertEqual(len(fields), 2)
                birth, death = fields
                address = next(section for section in dialog.findChildren(AnimalDialogSection)
                               if section.property('section_id') == 'address')
                states = []
                for field in fields:
                    point = field.mapToGlobal(field.rect().center())
                    hit = self.qt.widgetAt(point)
                    states.append({'enabled': field.isEnabled(), 'read_only': field.isReadOnly(),
                                   'hit': type(hit).__name__, 'same': hit is field,
                                   'rect': [point.x(), point.y(), field.width(), field.height()]})
                print('ISSUE_318_INPUT ' + json.dumps({'role': role, 'fields': states,
                                                     'identity_permission': app._master_can('core.edit_animal_identity')}))
                self.assertTrue(all(state['enabled'] and not state['read_only'] and state['same'] for state in states))
                self.assertFalse(birth.text() or death.text())
                self.assertTrue(address.isVisible())
                dialog.close()

    def test_318_real_seeded_dispatch_typed_creation_save_reopen_and_immutable_identity(self):
        app = self._318_admin_host()
        for role, name in ((MAIN.Role.EXPERIMENTAL.value, 'test-2'),
                           (MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING, 'test-4')):
            with self.subTest(role=role):
                self.assertFalse(any(record.get('name') == name for record in app.animals.values()))
                dialog, controls = self._318_open(app, role)
                save = self._318_fill_new(app, dialog, controls, name)
                self.assertTrue(save.isEnabled())
                save.click()
                self.qt.processEvents()
                self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted, app.errors)
                key = next(key for key, record in app.animals.items() if record.get('name') == name)
                expected = copy.deepcopy(app.animals[key])
                self.assertEqual((expected['rolle'], expected['birth_date'], expected.get('death_date', '')),
                                 (role, '01.01.2020', ''))
                self.assertTrue(expected['id'])
                self.assertEqual(app.backend.load_core_data()['animals'][key]['birth_date'], '01.01.2020')
                reopened, fields = self._318_open(app, role, name=key)
                self.assertEqual(fields['birth_widget'].text(), '01.01.2020')
                self.assertFalse(fields['name_widget'].isEnabled() and not fields['name_widget'].isReadOnly())
                fields['birth_widget'].setText('02.01.2020')
                next(button for button in reopened.findChildren(QPushButton)
                     if button.text() == app.messages['button.save']).click()
                self.assertEqual(app.animals[key], expected)
                self.assertNotEqual(reopened.result(), QDialog.DialogCode.Accepted)
                reopened.close()

    def test_318_new_invalid_missing_birth_cancel_denied_create_and_read_only_stay_safe(self):
        app = self._318_admin_host()
        before = copy.deepcopy(app.animals)
        for role in (MAIN.Role.EXPERIMENTAL.value, MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING):
            with self.subTest(role=role):
                dialog, controls = self._318_open(app, role)
                save = self._318_fill_new(app, dialog, controls, 'test-cancel')
                for value in ('', 'not-a-date', '31.02.2020'):
                    controls['birth_widget'].setText(value)
                    save.click()
                    self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)
                    self.assertEqual(app.animals, before)
                dialog.close()
        app.master_track._current_username = None
        app.master_track._current_role = 'guest'
        with patch.object(QDialog, 'exec', side_effect=AssertionError('denied create opened editor')):
            app._dlg_new_animal()
        self.assertEqual(app.animals, before)
        self._316_admin(app)
        target = next(key for key, record in app.animals.items()
                      if app._animal_role_value(record) == MAIN.Role.EXPERIMENTAL.value)
        dialog, controls = self._318_open(app, MAIN.Role.EXPERIMENTAL.value, target, read_only=True)
        self.assertFalse(controls['birth_widget'].isEnabled())
        save = next(button for button in dialog.findChildren(QPushButton) if button.text() == app.messages['button.save'])
        self.assertFalse(save.isEnabled())
        dialog.close()
        self.assertEqual(app.animals, before)

    def test_318_address_status_geometry_repeated_toggles_resize_and_no_row_drift(self):
        from PyQt6.QtTest import QTest
        from Plugins.core.animal_dialog_sections import AnimalDialogSection
        app = self._318_admin_host()
        for role in (MAIN.Role.EXPERIMENTAL.value, MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING):
            with self.subTest(role=role):
                dialog, controls = self._318_open(app, role)
                address = next(section for section in dialog.findChildren(AnimalDialogSection)
                               if section.property('section_id') == 'address')
                toggle = address.findChild(QPushButton, 'animalDialogSectionToggle')
                scroll = dialog.findChild(QScrollArea, 'animalDialogFormScroll')
                form = scroll.widget().layout()
                rows = form.rowCount()
                states = dialog.findChildren(QCheckBox)
                baseline = None
                for width in (1000, 620, 1000):
                    dialog.resize(width, dialog.height())
                    for expanded in (True, False, True, False):
                        scroll.ensureWidgetVisible(toggle)
                        self.qt.processEvents()
                        self.assertIs(self.qt.widgetAt(toggle.mapToGlobal(toggle.rect().center())), toggle)
                        if address.is_expanded() != expanded:
                            QTest.mouseClick(toggle, Qt.MouseButton.LeftButton)
                        self.assertEqual(address.is_expanded(), expanded)
                        for _ in range(3):
                            self.qt.processEvents()
                        self.assertEqual(form.rowCount(), rows)
                        self.assertEqual(form.getWidgetPosition(address)[1], MAIN.QFormLayout.ItemRole.SpanningRole)
                        rect = MAIN.QtCore.QRect(address.mapToGlobal(address.rect().topLeft()), address.size())
                        for status in states:
                            if status.isVisible():
                                status_rect = MAIN.QtCore.QRect(status.mapToGlobal(status.rect().topLeft()), status.size())
                                self.assertFalse(rect.intersects(status_rect), status.text())
                        point = controls['birth_widget'].mapToGlobal(controls['birth_widget'].rect().center())
                        self.assertIs(self.qt.widgetAt(point), controls['birth_widget'])
                        if width == 1000 and not expanded:
                            current = (address.y(), address.height(), scroll.widget().sizeHint().height())
                            if baseline is None:
                                baseline = current
                            self.assertEqual(current, baseline)
                output = os.environ.get('PROGTRACK_318_EVIDENCE_DIR')
                if output:
                    Path(output).mkdir(parents=True, exist_ok=True)
                    dialog.grab().save(str(Path(output) / f'{role}-collapsed.png'))
                    address.set_expanded(True)
                    for _ in range(3):
                        self.qt.processEvents()
                    dialog.grab().save(str(Path(output) / f'{role}-expanded.png'))
                dialog.close()

    def test_318_real_address_save_guard_uses_host_permission_and_fails_closed(self):
        app = self._318_admin_host()
        plugin = app.cage_track_plugin
        occupants = plugin.store.load_data()['occupants']
        key = next(key for key, occupant in occupants.items()
                   if key in app.animals and occupant.get('cage_id') not in (None, '', 'cage_unassigned'))
        address = plugin.get_current_address(key)
        self.assertTrue(address['cage_id'])
        with patch.object(plugin.store, 'set_address_from_dialog', wraps=plugin.store.set_address_from_dialog) as save:
            plugin.save_address_from_dialog(key, {})
            self.assertEqual(save.call_count, 1)
        self.assertFalse(any(plugin.get_current_address(key).values()))
        app.master_track._current_username = None
        app.master_track._current_role = 'guest'
        before_errors = len(app.errors)
        with patch.object(plugin.store, 'set_address_from_dialog', side_effect=AssertionError('denied address write')):
            plugin.save_address_from_dialog(key, {})
            self.assertEqual(len(app.errors), before_errors, 'An unchanged address must not demand write permission')
            plugin.save_address_from_dialog(key, address)
        self.assertIn(('denied',), app.errors)
        plugin.app = SimpleNamespace(_show_permission_denied=lambda: app.errors.append(('missing-policy',)))
        with patch.object(plugin.store, 'set_address_from_dialog', side_effect=AssertionError('missing policy wrote address')):
            plugin.save_address_from_dialog(key, address)
        self.assertIn(('missing-policy',), app.errors)

    def _316_sql_state(self, app):
        with app.backend.adapter.transaction() as connection:
            return {table: [tuple(row) for row in connection.execute(query).fetchall()]
                    for table, query in {
                        'animals': 'SELECT ipid,record_json,revision,updated_at FROM animals ORDER BY ipid',
                        'measurements': 'SELECT measurement_id,animal_ipid,kind,measured_at,value_json,sample_id,source_id,created_at FROM measurements ORDER BY measurement_id',
                        'events': 'SELECT event_id,animal_ipid,payload_json,created_at FROM animal_events ORDER BY event_id',
                        'settings': "SELECT value_json,revision,updated_at FROM installation WHERE key='core.settings'",
                    }.items()}

    def _316_admin(self, app):
        app.master_track._current_username = 'Admin'
        app.master_track._current_role = app.master_track.user_db.get_user('Admin')['role']
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app)
        app._resolve_animal_key = MAIN.ProgTrackApp._resolve_animal_key.__get__(app)

    def _315_add_ui_events(self, app, method, key, *, role=None, count=3):
        kwargs = {'role_value': role} if role and method in (
            '_dlg_versuchstier', '_dlg_basic_animal_role') else {}
        dialog = self.open_dialog(app, method, name=key, **kwargs)
        label_key = {
            '_dlg_offspring': 'dialog.offspring.button.new_event',
            '_dlg_zuchttier': 'dialog.zuchttier.button.new_event',
            '_dlg_versuchstier': 'dialog.versuchstier.button.new_event',
            '_dlg_basic_animal_role': 'dialog.basic_role.button.new_event',
        }.get(method)
        label = (app.messages[label_key] if label_key else app.messages['button.new_item'].format(
            title=app.messages.get('dialog.female_animal.tab.events', 'Events')))
        add = next(b for b in dialog.findChildren(QPushButton) if b.text() == label)
        for index in range(count):
            add.click()
            self.qt.processEvents()
            date, combo, _delete = self.event_editor_rows_247(dialog)['']
            self.assertFalse(combo._progtrack_source_event, 'new rows must not be pre-stamped')
            token = 'custom.audit_315' if method == '_dlg_basic_animal_role' else 'measurement'
            choice = combo.findData(token)
            if choice < 0:
                choice = next(i for i in range(combo.count()) if combo.itemData(i))
            combo.setCurrentIndex(choice)
            date.setText('01.10.2026' if index < 2 else '02.10.2026')
        errors = []
        with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
            next(b for b in dialog.findChildren(QPushButton)
                 if b.text() == app.messages['button.save']).click()
        self.assertFalse(errors or app.errors, (errors, app.errors))

    def test_315_real_new_event_save_all_five_adapters_and_experimental_offspring(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        cases = [('_dlg_offspring', MAIN.Role.OFFSPRING.value),
                 ('_dlg_zuchttier', MAIN.Role.ZUCHTTIER.value),
                 ('_dlg_versuchstier', MAIN.Role.EXPERIMENTAL.value),
                 ('_dlg_versuchstier', MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING),
                 ('_dlg_female_animal', MAIN.Role.SPENDER.value),
                 ('_dlg_basic_animal_role', 'custom_role')]
        for method, role in cases:
            with self.subTest(adapter=method, role=role):
                if role == 'custom_role':
                    key = app.backend.records.get('seed', 'shire-colony', {})['named_animals']['Frodo']
                    raw = app.backend.load_core_data()
                    raw['animals'][key]['rolle'] = role
                    app.backend.save_core_data(raw)  # Explicit configured role fixture.
                    MAIN.ProgTrackApp._load_persistence(app)
                    app._active_custom_event_definitions_for_role = lambda *_a, **_k: [
                        {'event_type': 'custom.audit_315', 'name': 'Test-Event', 'kind': 'counting'}]
                else:
                    key = next(k for k, r in app.animals.items() if app._animal_role_value(r) == role)
                old_ids = {e['event_id'] for e in app.animals[key]['events']}
                receipts, calls = [], []
                save = app.backend.save_core_data
                def capture(*args, **kwargs):
                    receipt = save(*args, **kwargs)
                    receipts.append(receipt)
                    return receipt
                app._master_audit = lambda *args: calls.append(args)
                with patch.object(app.backend, 'save_core_data', side_effect=capture):
                    self._315_add_ui_events(app, method, key, role=role)
                self.assertEqual(len(receipts), 1)
                record = app.animals[key]
                added = [e for e in record['events'] if e['event_id'] not in old_ids]
                self.assertEqual(len(added), 3)
                self.assertEqual(len({e['event_id'] for e in added}), 3)
                self.assertTrue(all(e['recorded_role'] == role for e in added))
                self.assertEqual(added[0]['datum'], added[1]['datum'])
                encoded = app._serialize_core_record(key, record)
                stored = app.backend.load_core_data()['animals'][key]
                self.assertEqual(encoded, stored)
                self.assertEqual(receipts[0].acknowledged_records['animals', key], stored)
                self.assertEqual(receipts[0].audit_after['animals'][key], stored)
                self.assertEqual(sum('parameter=events;' in a[2] for a in calls), 1)
                with app.backend.adapter.transaction() as connection:
                    payloads = [json.loads(row[0]) for row in connection.execute(
                        'SELECT payload_json FROM animal_events WHERE animal_ipid=?', (key,))]
                self.assertEqual({e['event_id'] for e in payloads}, {e['event_id'] for e in stored['events']})
                self.assertTrue(all(e['recorded_role'] == role for e in payloads if e['event_id'] not in old_ids))

    def test_315_exact_seed_new_events_unrelated_weight_noop_and_restart(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        names = app.backend.records.get('seed', 'shire-colony', {})['named_animals']
        frodo, adamanta = names['Frodo'], names['Adamanta']
        self._315_add_ui_events(app, '_dlg_versuchstier', frodo)
        baseline, rows = app.backend.load_core_data(), self._316_sql_state(app)
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        record, collection = app.animals[frodo], app.animals[frodo]['events']
        row = collection[0]
        # Reproduce the original stale-candidate omission despite repaired UI producers.
        row.pop('recorded_role')
        weight = app.animals[adamanta]['gewicht'][0]
        weight['wert'] += .2
        app._save_persistence()
        self.assertIs(app.animals[frodo], record)
        self.assertIs(record['events'], collection)
        self.assertIs(next(e for e in collection if e['event_id'] == row['event_id']), row)
        self.assertTrue(row['recorded_role'])
        self.assertEqual(len(audit), 1)
        self.assertIn(f'animal={adamanta}; parameter=gewicht;', audit[0][2])
        self.assertEqual(app.backend.load_core_data()['animals'][frodo], baseline['animals'][frodo])
        after_rows = self._316_sql_state(app)
        self.assertEqual(after_rows['events'], rows['events'])
        self.assertEqual([r for r in after_rows['animals'] if r[0] != adamanta],
                         [r for r in rows['animals'] if r[0] != adamanta])
        old_measurements = {r[0]: r for r in rows['measurements']}
        for r in after_rows['measurements']:
            if r[0] in old_measurements:
                self.assertEqual(r, old_measurements[r[0]])
        audit.clear()
        app._save_persistence()
        self.assertFalse(audit)
        self.assertEqual(self._316_sql_state(app), after_rows)
        app.backend.close()
        app.backend = ProgTrackBackend(app.backend.paths, acquire_process_lock=False)
        self.addCleanup(app.backend.close)
        MAIN.ProgTrackApp._load_persistence(app)
        self._315_add_ui_events(app, '_dlg_versuchstier', frodo, count=0)
        self.assertFalse(audit)
        self.assertEqual(self._316_sql_state(app), after_rows)

    def test_316_save_error_boundaries_do_not_retry_or_misattribute(self):
        app = self._307_sqlite_host(population=2)
        # This test exercises failures *after* a permitted source command.
        # Source-denial behavior is covered independently by #321 tests.
        checker = app._master_can
        app._master_can = lambda action, **kwargs: (
            True if action in ('core.edit_animal_core', 'core.edit_animal_measurements')
            else checker(action, **kwargs))
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app)
        # Establish fixture preparation once, separately from native fresh-seed tests.
        app._save_persistence()
        key = next(iter(app.animals))
        event = {'datum': MAIN.parse_stored_date('2026-01-02T14:15:16.123456+02:00'),
                 'event_type': 'measurement', 'event_id': 'ack-316',
                 'unknown_list': [{'event_id': 'metadata', 'value': 1},
                                  {'event_id': 'metadata', 'value': 2}]}
        app.animals[key]['events'].append(event)
        before = self._316_sql_state(app)
        with patch.object(app.backend, 'save_core_data', side_effect=TypeError('internal failure')) as save:
            with self.assertRaisesRegex(TypeError, 'internal failure'):
                app._save_persistence()
            self.assertEqual(save.call_count, 1)
        self.assertEqual(self._316_sql_state(app), before)
        self.assertNotIn('recorded_role', event)

        with patch.object(app, '_audit_data_snapshot_diff_when_safe', side_effect=RuntimeError('audit failed')), \
             patch.object(app.backend, 'save_core_data', wraps=app.backend.save_core_data) as save:
            with self.assertRaisesRegex(RuntimeError, 'audit failed'):
                app._save_persistence()
            self.assertEqual(save.call_count, 1)
        self.assertEqual(event['recorded_role'], app._animal_role_value(app.animals[key]))
        self.assertEqual(app._serialize_core_record(key, app.animals[key]),
                         app.backend.load_core_data()['animals'][key])
        # A switched actor must not receive another actor's accepted records/audit.
        event.pop('recorded_role')
        save = app.backend.save_core_data
        def switch_actor(*args, **kwargs):
            result = save(*args, **kwargs)
            app.master_track._current_username = 'different-actor'
            return result
        with patch.object(app.backend, 'save_core_data', side_effect=switch_actor), \
             patch.object(app, '_audit_data_snapshot_diff_when_safe') as audit:
            with self.assertRaisesRegex(RuntimeError, 'actor/backend changed'):
                app._save_persistence()
            audit.assert_not_called()
        self.assertNotIn('recorded_role', event)

    def test_316_preparation_ignores_only_event_enumeration_and_keeps_numeric_input(self):
        for old in (27, 27.5, 0):
            value = parse_measurement(str(old), original=old)
            self.assertEqual(type(value), type(old))
            self.assertEqual(value, old)
        self.assertEqual(parse_measurement('27.2', original=27), 27.2)
        self.assertIs(type(parse_measurement('0', original=False)), float)
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        for records in (app.animals, app.archived):
            for record in records.values():
                record['events'].reverse()
        # No schema/default preparation for stable-ID enumeration alone.
        with patch.object(MAIN, 'ensure_project_periods', side_effect=AssertionError('false dirty')):
            app._save_persistence()

    def test_316_fresh_load_noop_reopen_and_actor_reads_leave_rows_unchanged(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        baseline, rows = app.backend.load_core_data(), self._316_sql_state(app)
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        for actor in ('Admin', 'Vet', 'Admin'):
            app.master_track._current_username = actor
            app.master_track._current_role = app.master_track.user_db.get_user(actor)['role']
            MAIN.ProgTrackApp._load_persistence(app)
            self.assertEqual(self._316_sql_state(app), rows)
            self.assertEqual(app.backend.load_core_data(), baseline)
        app._save_persistence()
        self.assertEqual(self._316_sql_state(app), rows)
        self.assertEqual(app.backend.load_core_data(), baseline)
        self.assertFalse(audit)
        # Real backend reconstruction, not a warm-up write.
        app.backend.close()
        app.backend = ProgTrackBackend(app.backend.paths, acquire_process_lock=False)
        self.addCleanup(app.backend.close)
        MAIN.ProgTrackApp._load_persistence(app)
        app._save_persistence()
        self.assertEqual(self._316_sql_state(app), rows)
        self.assertFalse(audit)

    def test_316_seeded_role_editors_unchanged_saves_preserve_all_records_and_rows(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        baseline, rows = app.backend.load_core_data(), self._316_sql_state(app)
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        cases = [
            (MAIN.Role.SPENDER.value, '_dlg_female_animal', {'default_role': MAIN.Role.SPENDER.value}),
            (MAIN.Role.AMME.value, '_dlg_female_animal', {'default_role': MAIN.Role.AMME.value}),
            (MAIN.Role.SAMENSP.value, '_dlg_samenspender', {}),
            (MAIN.Role.OFFSPRING.value, '_dlg_offspring', {}),
            (MAIN.Role.PARTNER.value, '_dlg_partner', {}),
            (MAIN.Role.ZUCHTTIER.value, '_dlg_zuchttier', {}),
            (MAIN.Role.EXPERIMENTAL.value, '_dlg_versuchstier', {}),
            (MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING, '_dlg_versuchstier',
             {'role_value': MAIN.ROLE_VALUE_EXPERIMENTAL_OFFSPRING}),
        ]
        for role, method, kwargs in cases:
            with self.subTest(role=role):
                key = next(k for k, r in app.animals.items() if app._animal_role_value(r) == role)
                dlg = self.open_dialog(app, method, name=key, **kwargs)
                errors = []
                with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                    next(b for b in dlg.findChildren(QPushButton) if b.text() == app.messages['button.save']).click()
                self.assertFalse(errors or app.errors, (errors, app.errors))
                self.assertEqual(app.backend.load_core_data(), baseline)
                self.assertEqual(self._316_sql_state(app), rows)
                self.assertFalse(audit)

    def test_316_runtime_partial_record_is_not_normalized_by_another_animal_save(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        raw = app.backend.load_core_data()
        untouched, edited = list(raw['animals'])[:2]
        original = raw['animals'][untouched]
        for field in ('chip_nr', 'special_status', 'sick_times', 'project_period_id'):
            original.pop(field, None)
        for collection in ('events', 'daten', 'gewicht', 'sperm'):
            for row in original.get(collection, []):
                for field in list(row):
                    if field.startswith('project_period_'):
                        row.pop(field)
        app.backend.save_core_data(raw)  # Explicit disposable producer fixture, not warm-up UI Save.
        MAIN.ProgTrackApp._load_persistence(app)
        baseline = app.backend.load_core_data()
        before_rows = self._316_sql_state(app)
        app._events_in_current_project_period(app.animals[untouched])
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        app.animals[edited]['genotype'] = 'test-316-real-edit'
        app._save_persistence()
        self.assertEqual(len(audit), 1)
        self.assertIn('parameter=genotype;', audit[0][2])
        self.assertEqual(app.backend.load_core_data()['animals'][untouched],
                         baseline['animals'][untouched])
        after_rows = self._316_sql_state(app)
        self.assertEqual(after_rows['measurements'], before_rows['measurements'])
        self.assertEqual(after_rows['events'], before_rows['events'])
        self.assertEqual([r for r in after_rows['animals'] if r[0] != edited],
                         [r for r in before_rows['animals'] if r[0] != edited])

    def test_316_lossless_codec_metadata_dates_duplicates_and_invalid_rows(self):
        app = self.host()
        for literal in ('2024-01-02', '02.01.2024', '2024-01-02T09:12:13.123456+02:00'):
            for field in ('daten', 'gewicht', 'pdg', 'sperm', 'events'):
                with self.subTest(date=literal, field=field):
                    row = {'datum': literal, 'wert': 27, 'unknown': {'flag': False, 'zero': 0, 'v': None},
                           'probennummer': '', 'project_period_id': 'original-period'}
                    if field == 'events':
                        row.update(event_id='event-316', event_type='surgery', recorded_role='experimental')
                    record = self.record()
                    record[field] = [copy.deepcopy(row), copy.deepcopy(row)]
                    expected = copy.deepcopy(record[field])
                    for r in record[field]:
                        r['datum'] = MAIN.parse_stored_date(r['datum'])
                    encoded = app._serialize_core_record('Subject', record)
                    self.assertEqual(encoded[field], expected)
                    restored = pickle.loads(pickle.dumps(record))
                    self.assertEqual(app._serialize_core_record('Subject', restored)[field], expected)
                    self.assertIs(type(encoded[field][0]['unknown']['flag']), bool)
        view = MAIN.parse_stored_date('2024-01-02T09:12:13.123456+02:00')
        updated = app._measurement_row_with_inputs({'datum': view, 'wert': 27, 'meta': [1, None]},
                                                   datum=datetime(2024, 1, 2), wert=27.0)
        self.assertIs(updated['datum'], view)
        self.assertIs(type(updated['wert']), int)
        self.assertEqual(updated['meta'], [1, None])
        for form in ({}, {'count': None}, {'count': ''}):
            source = {'datum': view, **form}
            self.assertEqual(app._measurement_row_with_inputs(source, count=None), source)
        self.assertIsNone(app._measurement_row_with_inputs({'datum': view, 'count': 1}, count=None)['count'])
        changed = app._measurement_row_with_inputs(updated, datum=datetime(2024, 1, 3), wert=27.1)
        self.assertEqual(changed['datum'], datetime(2024, 1, 3))
        for field in ('events', 'gewicht'):
            record = self.record()
            record[field] = [{'datum': 'invalid', 'event_type': 'surgery'}]
            with self.assertRaisesRegex(ValueError, 'Invalid .* date'):
                app._serialize_core_record('Subject', record)

    def test_316_read_only_period_views_do_not_mutate_source(self):
        from Plugins.core.project_periods import current_period_records, event_counts_by_period, event_period_id
        from Plugins.Animal_Reports.report_index import AnimalReportIndex
        app = self.host()
        record = self.record()
        record.update(project='Backcrossing', project_entry_date='01.01.2025',
                      events=[{'event_id': 'e316', 'event_type': 'measurement',
                               'datum': datetime(2025, 2, 1), 'recorded_role': 'experimental'}])
        original = copy.deepcopy(record)
        self.assertEqual(len(current_period_records(record, 'events')), 1)
        self.assertTrue(event_counts_by_period(record))
        self.assertTrue(event_period_id(record, record['events'][0]))
        selected = app._events_in_current_project_period(record, 'experimental')
        self.assertEqual(len(selected), 1)
        self.assertIs(selected[0], record['events'][0], 'confirmation logic retains row identity')
        AnimalReportIndex(record, lambda event_type: event_type, date_format=MAIN.DATE_FORMAT)
        self.assertEqual(record, original)

    def test_316_genuine_multi_animal_active_archived_changes_are_not_hidden(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        baseline = app.backend.load_core_data()
        a, b = next(iter(app.animals)), next(iter(app.archived))
        app.animals[a]['genotype'] = 'test-316-active'
        app.archived[b]['special_status'] = 'test-316-archived'
        audit = []
        app._master_audit = lambda *args: audit.append((app.master_track.current_username, args))
        app._save_persistence()
        self.assertEqual(len(audit), 2)
        self.assertTrue(all(actor == 'Admin' for actor, _args in audit))
        expected = copy.deepcopy(baseline)
        expected['animals'][a]['genotype'] = 'test-316-active'
        expected['archived_animals'][b]['special_status'] = 'test-316-archived'
        self.assertEqual(app.backend.load_core_data(), expected)

    def test_316_read_or_decode_failure_prevents_writes_and_success_audit(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        rows = self._316_sql_state(app)
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        with patch.object(app, '_read_json', side_effect=OSError('read failed')):
            with self.assertRaisesRegex(OSError, 'read failed'):
                app._save_persistence()
            with self.assertRaisesRegex(OSError, 'read failed'):
                app._write_json({'animals': app.animals, 'archived_animals': app.archived})
        bad = app.backend.load_core_data()
        bad['animals'][next(iter(bad['animals']))]['events'].append({'event_type': 'surgery'})
        with patch.object(app, '_read_json', return_value=bad):
            MAIN.ProgTrackApp._load_persistence(app)
        with self.assertRaisesRegex(ValueError, 'Cannot save'):
            app._save_persistence()
        self.assertEqual(self._316_sql_state(app), rows)
        self.assertFalse(audit)

    def test_316_fresh_seed_first_save_has_no_unrelated_normalization(self):
        # No preparatory Save: exact real first-operation audit is contractual.
        self._315_seeded_save_case()

    def test_315_existing_database_after_startup_normalization(self):
        self._315_seeded_save_case(settle_startup=True)

    def _315_seeded_save_case(self, *, settle_startup=False):
        app = self._307_sqlite_host(seeded=True)
        # The returned witness used Admin. Seed Manager intentionally lacks
        # measurement editing; do not bypass/grant that permission in tests.
        app.master_track._current_username = 'Admin'
        app.master_track._current_role = app.master_track.user_db.get_user('Admin')['role']
        app._resolve_animal_key = MAIN.ProgTrackApp._resolve_animal_key.__get__(app)
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app)
        app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        if settle_startup:
            # Separate witness for the reported already-used database. This
            # setup is NOT acceptance of the fresh-seed first-Save case above.
            app._save_persistence()
        aliases = app.backend.records.get('seed', 'shire-colony', {})['named_animals']
        key = aliases['Adamanta']
        other = next(k for k in app.animals if k != key)
        baseline = app.backend.load_core_data()
        baseline_rows = self._316_sql_state(app)
        for field in ('death_date', 'severity'):
            self.assertNotIn(field, baseline['animals'][key])
        original_events = copy.deepcopy(app.animals[other]['events'])
        audit = []
        app._master_audit = lambda *args: audit.append((app.master_track.current_username, args))
        before_weights = len(app.animals[key]['gewicht'])
        dialog = self.open_dialog(app, '_dlg_zuchttier', name=key)
        page = self.weight_page(app, dialog)
        fields_before = set(page.findChildren(QLineEdit))
        next(b for b in page.findChildren(QPushButton)
             if b.text() == app.messages['button.add_weight']).click()
        date_field, value_field = [x for x in page.findChildren(QLineEdit) if x not in fields_before]
        date_field.setText('17.05.2019')
        value_field.setText('27.1')
        errors = []
        with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
            next(b for b in dialog.findChildren(QPushButton)
                 if b.text() == app.messages['button.save']).click()
        self.assertFalse(errors or app.errors, (errors, app.errors))
        if len(audit) != 1:
            stored_debug = app.backend.load_core_data()
            samples = {}
            for section in ('animals', 'archived_animals'):
                for ipid, before in baseline[section].items():
                    for field in ('events', 'gewicht', 'daten', 'pdg', 'sperm'):
                        old, new = before.get(field, []), stored_debug[section][ipid].get(field, [])
                        if old != new and field not in samples:
                            samples[field] = (ipid, next(((a,b) for a,b in zip(old,new) if a != b),
                                                        (len(old), len(new))))
            print('ISSUE_316_UNEXPECTED_DELTAS', samples)
        self.assertEqual(len(audit), 1, Counter(
            a[1][2].split('parameter=', 1)[-1].split(';', 1)[0] for a in audit))
        self.assertEqual(audit[0][0], 'Admin')
        self.assertIn(f'animal={key}; parameter=gewicht;', audit[0][1][2])
        self.assertEqual(app.animals[other]['events'], original_events)
        stored = app.backend.load_core_data()['animals'][key]
        after_save = app.backend.load_core_data()
        for section in ('animals', 'archived_animals'):
            for ipid, original in baseline[section].items():
                if section == 'animals' and ipid == key:
                    self.assertEqual({k: v for k, v in after_save[section][ipid].items() if k != 'gewicht'},
                                     {k: v for k, v in original.items() if k != 'gewicht'})
                else:
                    self.assertEqual(after_save[section][ipid], original, ipid)
        committed_rows = self._316_sql_state(app)
        self.assertEqual([r for r in committed_rows['animals'] if r[0] != key],
                         [r for r in baseline_rows['animals'] if r[0] != key])
        before_target = next(r for r in baseline_rows['animals'] if r[0] == key)
        after_target = next(r for r in committed_rows['animals'] if r[0] == key)
        self.assertEqual(after_target[1], before_target[1])
        self.assertEqual(after_target[2], before_target[2] + 1)
        self.assertEqual(committed_rows['events'], baseline_rows['events'])
        self.assertTrue(set(baseline_rows['measurements']).issubset(set(committed_rows['measurements'])))
        self.assertEqual(len(committed_rows['measurements']), len(baseline_rows['measurements']) + 1)
        self.assertEqual(len(stored['gewicht']), before_weights + 1)
        self.assertEqual(stored['gewicht'][-1]['wert'], 27.1)
        self.assertNotIn('death_date', stored)
        self.assertNotIn('severity', stored)
        audit.clear()
        dialog = self.open_dialog(app, '_dlg_zuchttier', name=key)
        next(b for b in dialog.findChildren(QPushButton)
             if b.text() == app.messages['button.save']).click()
        self.assertFalse(audit)
        self.assertEqual(self._316_sql_state(app), committed_rows)
        dialog = self.open_dialog(app, '_dlg_zuchttier', name=key)
        next(x for x in self.weight_page(app, dialog).findChildren(QLineEdit)
             if x.text() == '27.1').setText('999')
        dialog.reject()
        self.assertFalse(audit)
        self.assertEqual(app.backend.load_core_data()['animals'][key]['gewicht'], stored['gewicht'])
        with patch.object(app.backend, 'save_core_data', side_effect=OSError('commit failed')):
            with self.assertRaisesRegex(OSError, 'commit failed'):
                app._write_json({'animals': app.animals, 'archived_animals': app.archived})
        self.assertFalse(audit, 'Failed writes must not produce success audit evidence')

    def test_315_all_editor_saves_preserve_optional_blank_representation(self):
        cases = list(self.weight_editor_cases()) + [('_dlg_female_animal', MAIN.Role.AMME.value)]
        for method, role in cases:
            for form in ('absent', 'null', 'empty'):
                with self.subTest(editor=method, role=role, optional=form):
                    app = self.host()
                    app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
                    app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                    record = self.record(role)
                    if form != 'absent':
                        record.update(death_date=None if form == 'null' else '',
                                      severity=None if form == 'null' else '')
                    app.animals['Subject'] = record
                    original = copy.deepcopy(record)
                    dialog = self.open_dialog(app, method)
                    page = self.weight_page(app, dialog)
                    field = next(x for x in page.findChildren(QLineEdit) if x.text() == '27.1')
                    field.setText('27.2')
                    commits = []
                    app._save_persistence = lambda **_k: commits.append(copy.deepcopy(app.animals))
                    errors = []
                    with patch('sys.excepthook', lambda _t, e, _tb: errors.append(str(e))):
                        next(b for b in dialog.findChildren(QPushButton)
                             if b.text() == app.messages['button.save']).click()
                    self.assertFalse(errors or app.errors, (errors, app.errors))
                    self.assertEqual(len(commits), 1)
                    saved = app.animals['Subject']
                    self.assertEqual(saved['gewicht'][0]['wert'], 27.2)
                    for optional in ('death_date', 'severity'):
                        self.assertEqual(optional in saved, optional in original)
                        self.assertEqual(saved.get(optional), original.get(optional))

    def test_315_snapshot_audit_keeps_real_changes_to_both_animals(self):
        from tests.test_audit_logging import _315_events
        app = self.host()
        calls = []
        app._master_audit = lambda *args: calls.append(args)
        before = {'animals': {'test-1': {'gewicht': [{'wert': 27}]},
                              'test-2': {'events': _315_events()}}, 'archived_animals': {}}
        after = copy.deepcopy(before)
        after['animals']['test-1']['gewicht'].append({'wert': 27.1})
        after['animals']['test-2']['events'] = [after['animals']['test-2']['events'][i]
                                              for i in (0, 3, 2, 1, 5, 4)]
        MAIN.ProgTrackApp._audit_data_snapshot_diff(app, before, after)
        self.assertEqual(len(calls), 1)
        self.assertIn('animal=test-1; parameter=gewicht;', calls[0][2])
        calls.clear()
        after['animals']['test-2']['events'][0]['ordinal'] = 1
        MAIN.ProgTrackApp._audit_data_snapshot_diff(app, before, after)
        self.assertEqual(len(calls), 2)
        self.assertIn('animal=test-2; parameter=events;', calls[1][2])
        calls.clear()
        # Scalar-only records and identical payloads with unknown/empty fields
        # must still produce zero noise, without dropping other-animal edits.
        MAIN.ProgTrackApp._audit_data_snapshot_diff(app, before, before)
        self.assertFalse(calls)
        old = {'animals': {'test-2': {'events': [{'event_id': 'x', 'confirmed': False}]}}}
        new = {'animals': {'test-2': {'events': [{'event_id': 'x', 'confirmed': 0}]}}}
        MAIN.ProgTrackApp._audit_data_snapshot_diff(app, old, new)
        self.assertEqual(len(calls), 1)

    def test_315_real_death_clear_and_death_reason_validation_remain(self):
        app = self.host()
        capture = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        previous = dict(self.record(), death_date='01.01.2026', severity='SV1')
        cleared = dict(previous, death_date='', severity='')
        self.assertTrue(capture(cleared, previous))
        self.assertEqual((cleared['death_date'], cleared['severity']), ('', ''))
        previous = self.record()
        candidate = dict(previous, death_date='02.01.2026', severity='SV1')
        app._load_identity_conventions = lambda: {'death_causes': ['Natural death']}
        with patch.object(MAIN.QInputDialog, 'getItem', return_value=('', False)):
            self.assertFalse(capture(candidate, previous))
        with patch.object(MAIN.QInputDialog, 'getItem', return_value=('Natural death', True)):
            self.assertTrue(capture(candidate, previous))
        self.assertEqual(candidate['death_cause'], 'Natural death')
        self.assertEqual(candidate['severity'], '')
        self.assertTrue(candidate['lifecycle_events'])

    def flow_313(self):
        flow, source = self.flow_314()
        flow._cancel_canvas_gesture()
        opened = []
        flow._show_animal_info = lambda name: opened.append(('animal', name))
        flow._open_freezer_embryo_list = lambda: opened.append(('freezer',))
        flow._edit_embryo = lambda *args: opened.append(('embryo', *args))
        for name, callback in [('pick_event', flow._on_pick),
                               ('button_press_event', flow._on_mouse_press),
                               ('button_release_event', flow._on_mouse_release),
                               ('motion_notify_event', flow._on_hover),
                               ('motion_notify_event', flow._on_mouse_move),
                               ('figure_leave_event', flow._on_hover_leave),
                               ('close_event', flow._on_hover_leave)]:
            flow.canvas.mpl_connect(name, callback)
        return flow, source, opened

    def mouse_313(self, flow, name, pixel, button=1, *, buttons=None, dblclick=False):
        from matplotlib.backend_bases import MouseEvent
        if name == 'motion_notify_event' and buttons is None:
            buttons = {button} if button else set()
        event = MouseEvent(name, flow.canvas, *pixel, button=button,
                           buttons=buttons, dblclick=dblclick)
        flow.canvas.callbacks.process(name, event)
        return event

    def test_313_left_middle_pan_fixed_origin_zoom_dpi_and_no_writes(self):
        for dpi in (100, 200):
            for span in (.25, 1, 4):
                results = []
                for button in (1, 2):
                    with self.subTest(dpi=dpi, span=span, button=button):
                        flow, _source, opened = self.flow_313()
                        flow.figure.set_dpi(dpi)
                        xlim, ylim = flow._apply_aspect_fill((0, span), (0, span))
                        flow.ax.set(xlim=xlim, ylim=ylim)
                        flow.canvas.draw()
                        old = copy.deepcopy((flow.manual_data, flow.temp_positions,
                                             flow.settings, flow.embryo_visibility))
                        save = []
                        flow._save_flow_track_data = lambda: save.append('data')
                        flow._save_settings = lambda: save.append('settings')
                        x, y = flow.ax.bbox.x0 + 80, flow.ax.bbox.y0 + 80
                        self.mouse_313(flow, 'button_press_event', (x, y), button)
                        start = copy.deepcopy(flow._pan_gesture)
                        for delta in ((20, 10), (40, 30), (40, 30)):
                            self.mouse_313(flow, 'motion_notify_event', (x + delta[0], y + delta[1]), button)
                        expected = (xlim[0] - 40 * span / start['height'],
                                    ylim[0] - 30 * span / start['height'])
                        self.assertAlmostEqual(flow.ax.get_xlim()[0], expected[0])
                        self.assertAlmostEqual(flow.ax.get_ylim()[0], expected[1])
                        self.assertAlmostEqual(np.diff(flow.ax.get_xlim())[0], np.diff(xlim)[0])
                        self.assertAlmostEqual(np.diff(flow.ax.get_ylim())[0], np.diff(ylim)[0])
                        results.append((flow.ax.get_xlim(), flow.ax.get_ylim()))
                        self.mouse_313(flow, 'motion_notify_event', (x, y), button)
                        np.testing.assert_allclose(flow.ax.get_xlim(), xlim)
                        np.testing.assert_allclose(flow.ax.get_ylim(), ylim)
                        self.mouse_313(flow, 'button_release_event', (x + 40, y + 30), button)
                        self.assertFalse(flow.pan_active)
                        self.assertIsNone(flow._pan_gesture)
                        self.assertEqual((flow.manual_data, flow.temp_positions,
                                          flow.settings, flow.embryo_visibility), old)
                        self.assertFalse(save or opened)
                self.assertEqual(results[0], results[1])

    def test_313_threshold_empty_double_click_outside_and_reverse_path(self):
        flow, _source, opened = self.flow_313()
        original = (tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim()))
        for button in (1, 2):
            with self.subTest(button=button):
                self.mouse_313(flow, 'button_press_event', (100, 100), button)
                self.mouse_313(flow, 'motion_notify_event', (102, 102), button)
                self.assertEqual((tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim())), original)
                self.mouse_313(flow, 'button_release_event', (-20, -20), button)
                self.assertFalse(flow.pan_active)
        self.mouse_313(flow, 'button_press_event', (100, 100), dblclick=True)
        self.assertFalse(flow.pan_active)
        self.mouse_313(flow, 'button_press_event', (-20, -20))
        self.assertFalse(flow.pan_active)
        self.mouse_313(flow, 'button_press_event', (100, 100))
        self.mouse_313(flow, 'motion_notify_event', (150, 140))
        shifted = (tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim()))
        self.mouse_313(flow, 'button_release_event', (-20, -20))
        self.mouse_313(flow, 'motion_notify_event', (200, 200), button=None)
        self.assertEqual((tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim())), shifted)
        self.assertFalse(opened)

    def test_313_node_embryo_label_lifebar_and_pick_priority(self):
        flow, source, opened = self.flow_313()
        embryo, = flow.ax.plot(.8, .6, 'o', picker=8)
        flow.embryo_artists[embryo] = ('transfer-1', 'test-E1')
        # Isolate passive labels from the marker's larger pick radius.
        flow.node_label_artists['test-1'].set_position((0, -40))
        flow.lifebar_data['test-1'][0]['annotation'].set_position((0, -25))
        flow.canvas.draw()
        for artist in (source, embryo, flow.node_label_artists['test-1'],
                       flow.lifebar_data['test-1'][0]['annotation']):
            with self.subTest(kind=type(artist).__name__):
                box = artist.get_window_extent(flow.canvas.get_renderer())
                point = ((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2)
                old = (tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim()))
                self.mouse_313(flow, 'button_press_event', point)
                self.assertFalse(flow.pan_active)
                if artist in (source, embryo):
                    self.assertIs(flow.drag_artist, artist)  # Figure.pick preceded press.
                self.mouse_313(flow, 'button_release_event', point)
                self.assertEqual((tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim())), old)
        self.assertEqual(opened, [('animal', 'test-1'), ('embryo', 'transfer-1', 'test-E1')])
        for artist, key in ((source, ('animal', 'test-1')), (embryo, ('embryo', 'test-E1'))):
            with self.subTest(drag=key):
                point = flow.ax.transData.transform(tuple(float(v[0]) for v in artist.get_data()))
                self.mouse_313(flow, 'button_press_event', point)
                self.mouse_313(flow, 'motion_notify_event', (point[0] + 60, point[1] + 60))
                self.assertIn(key, flow.temp_positions)
                self.assertFalse(flow.pan_active)
                self.mouse_313(flow, 'button_release_event', (point[0] + 60, point[1] + 60))
        self.assertEqual(len(opened), 2)

    def test_313_permission_identity_and_lost_button_abort(self):
        for change in ('denied', 'actor', 'backend', 'epoch', 'released', 'chord'):
            with self.subTest(change=change):
                flow, _source, opened = self.flow_313()
                old = (tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim()))
                self.mouse_313(flow, 'button_press_event', (100, 100))
                if change == 'denied':
                    flow._can = lambda *_: False
                elif change == 'actor':
                    flow.parent_app.master_track.current_username = 'Manager'
                elif change == 'backend':
                    flow.parent_app.backend = object()
                elif change == 'epoch':
                    flow.parent_app._view_session_epoch = 99
                buttons = set() if change == 'released' else {1, 2} if change == 'chord' else {1}
                self.mouse_313(flow, 'motion_notify_event', (160, 160), buttons=buttons)
                self.assertFalse(flow.pan_active)
                self.assertEqual((tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim())), old)
                self.assertFalse(opened or flow.temp_positions)
        flow, _source, _opened = self.flow_313()
        flow._can = lambda *_: False
        self.mouse_313(flow, 'button_press_event', (100, 100))
        self.assertFalse(flow.pan_active)

    def test_313_leave_redraw_resize_scroll_and_qt_focus_cleanup(self):
        from matplotlib.backend_bases import Event, MouseEvent
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        for action in ('leave', 'close', 'redraw', 'resize', 'scroll', 'hide', 'focus', 'deactivate'):
            with self.subTest(action=action):
                flow, _source, opened = self.flow_313()
                flow._install_hover_cleanup()
                self.mouse_313(flow, 'button_press_event', (100, 100))
                if action in ('leave', 'close'):
                    name = 'figure_leave_event' if action == 'leave' else 'close_event'
                    flow.canvas.callbacks.process(name, Event(name, flow.canvas))
                elif action == 'redraw':
                    flow.parent_app.animals = {}
                    flow.stack = SimpleNamespace(setCurrentWidget=lambda *_: None)
                    flow.splash_widget = object()
                    FlowTrackWidget._redraw_canvas(flow)
                elif action == 'resize':
                    flow._on_resize(Event('resize_event', flow.canvas))
                elif action == 'scroll':
                    flow._on_scroll(MouseEvent('scroll_event', flow.canvas, 100, 100, button='up', step=1))
                else:
                    kind = {'hide': QEvent.Type.Hide, 'focus': QEvent.Type.FocusOut,
                            'deactivate': QEvent.Type.WindowDeactivate}[action]
                    self.qt.sendEvent(flow.widget, QEvent(kind))
                self.assertFalse(flow.pan_active)
                self.mouse_313(flow, 'button_release_event', (-20, -20))
                self.assertFalse(opened)

    def test_313_mixed_buttons_do_not_toggle_or_open_and_callbacks_stay_fixed(self):
        flow, source, opened = self.flow_313()
        flow.node_artists[source] = 'test-3'
        flow._get_transfers_for_surrogate = lambda *_: [{'id': 'transfer-1'}]
        count = {key: len(values) for key, values in flow.canvas.callbacks.callbacks.items()}
        for repeat in range(10):
            self.mouse_313(flow, 'button_press_event', (100, 100))
            point = flow.ax.transData.transform((.5, .65))
            self.mouse_313(flow, 'button_press_event', point, 3)
            self.mouse_313(flow, 'button_release_event', point, 3)
            self.mouse_313(flow, 'button_release_event', point, 1)
            self.assertFalse(flow.pan_active or flow.drag_active or flow.embryo_visibility)
        self.assertFalse(opened)
        self.assertEqual({key: len(values) for key, values in flow.canvas.callbacks.callbacks.items()}, count)

    def test_313_actual_qt_canvas_mouse_events_and_focus_lifecycle(self):
        from PyQt6 import QtGui
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        fixture = self.flow_312()
        parent = fixture.parent_app
        parent.QtGui, parent.matplotlib = QtGui, __import__('matplotlib')
        parent.backend = SimpleNamespace(records=Records())
        parent.backend.records.put('reproduction', 'flow', {'version': '3.0', 'manual_data': fixture.manual_data})
        flow = FlowTrackWidget(parent, fixture.messages)
        flow.widget.setParent(self.hosts[-1])
        flow.widget.resize(800, 600)
        flow.widget.show()
        self.qt.processEvents()
        flow.canvas.draw()
        before = copy.deepcopy((parent.backend.records.data, flow.manual_data, flow.temp_positions))
        bbox = flow.ax.bbox
        ratio = flow.canvas.device_pixel_ratio
        x = (bbox.x0 + bbox.width * .15) / ratio
        y = (flow.figure.bbox.height - bbox.y0 - bbox.height * .25) / ratio
        def send(kind, px, py, button, held):
            point = MAIN.QtCore.QPointF(px, py)
            event = QtGui.QMouseEvent(kind, point, point, button, held, Qt.KeyboardModifier.NoModifier)
            self.qt.sendEvent(flow.canvas, event)
        old = (tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim()))
        send(QEvent.Type.MouseButtonPress, x, y, Qt.MouseButton.LeftButton, Qt.MouseButton.LeftButton)
        self.assertTrue(flow.pan_active)
        send(QEvent.Type.MouseMove, x + 35, y + 25, Qt.MouseButton.NoButton, Qt.MouseButton.LeftButton)
        self.assertNotEqual((tuple(flow.ax.get_xlim()), tuple(flow.ax.get_ylim())), old)
        send(QEvent.Type.MouseButtonRelease, -20, -20, Qt.MouseButton.LeftButton, Qt.MouseButton.NoButton)
        self.assertFalse(flow.pan_active)
        send(QEvent.Type.MouseButtonPress, x, y, Qt.MouseButton.LeftButton, Qt.MouseButton.LeftButton)
        self.qt.sendEvent(flow.canvas, QEvent(QEvent.Type.FocusOut))
        self.assertFalse(flow.pan_active)
        send(QEvent.Type.MouseButtonRelease, x, y, Qt.MouseButton.LeftButton, Qt.MouseButton.NoButton)
        self.assertEqual((parent.backend.records.data, flow.manual_data, flow.temp_positions), before)

    def test_313_item_release_context_and_overlapping_pick_ownership(self):
        for change in ('denied', 'actor', 'backend', 'epoch'):
            with self.subTest(change=change):
                flow, source, opened = self.flow_313()
                point = flow.ax.transData.transform((.5, .65))
                self.mouse_313(flow, 'button_press_event', point)
                self.assertIs(flow.drag_artist, source)
                if change == 'denied':
                    flow._can = lambda *_: False
                elif change == 'actor':
                    flow.parent_app.master_track.current_username = 'Manager'
                elif change == 'backend':
                    flow.parent_app.backend = object()
                else:
                    flow.parent_app._view_session_epoch = 99
                self.mouse_313(flow, 'button_release_event', point)
                self.assertFalse(opened or flow.temp_positions or flow.drag_active)
        flow, source, opened = self.flow_313()
        overlap, = flow.ax.plot(.5, .65, 'o', picker=10)
        flow.embryo_artists[overlap] = ('transfer-1', 'test-E1')
        flow.canvas.draw()
        point = flow.ax.transData.transform((.5, .65))
        self.mouse_313(flow, 'button_press_event', point)
        self.assertIs(flow.drag_artist, source)
        self.mouse_313(flow, 'button_release_event', point)
        self.assertEqual(opened, [('animal', 'test-1')])
        self.mouse_313(flow, 'button_press_event', point)
        self.mouse_313(flow, 'button_release_event', (-20, -20))
        self.assertEqual(len(opened), 1)
        self.assertFalse(flow.drag_active or flow.pan_active)

    def flow_314(self, point=(.5, .65), figsize=(8, 6), dpi=100):
        flow = self.flow_312()
        flow.figure = flow.canvas.figure
        flow.figure.set_size_inches(*figsize)
        flow.figure.set_dpi(dpi)
        flow.figure.subplots_adjust(left=0, right=1, bottom=0, top=1)
        flow._hover_annotation = None
        flow.node_label_artists = {}
        flow.pan_active, flow.pan_start = False, None
        flow.drag_active, flow.drag_artist = False, None
        flow.drag_offset, flow.drag_threshold = (0, 0), .01
        flow.is_dragging, flow.pending_menu_action = False, None
        flow.click_start_pos = None
        flow.current_xlim, flow.current_ylim = (0, 1), (0, 1)
        flow.temp_positions, flow.embryo_visibility = {}, {}
        flow.node_artists, flow.embryo_artists = {}, {}
        flow.parent_app.matplotlib = __import__('matplotlib')
        flow.parent_app.master_track = SimpleNamespace(current_username='Admin')
        flow._save_settings = lambda: None
        flow.ax.set(xlim=(0, 1), ylim=(0, 1))
        flow.ax.axis('off')
        artist, = flow.ax.plot(*point, 'o', markersize=12, picker=10, zorder=20)
        flow.node_artists[artist] = 'test-1'
        flow._render_lifebars_for_animal('test-1', *point, 'egg_cell_donor')
        flow.node_label_artists['test-1'] = flow.ax.annotate(
            'test-1', xy=point, xytext=(0, -20), textcoords='offset points',
            ha='center', va='top', clip_on=False)
        flow.canvas.draw()
        return flow, artist

    def hover_314(self, flow, artist):
        from matplotlib.backend_bases import MouseEvent
        bbox = artist.get_window_extent(flow.canvas.get_renderer())
        event = MouseEvent('motion_notify_event', flow.canvas, (bbox.x0 + bbox.x1) / 2,
                           (bbox.y0 + bbox.y1) / 2)
        flow._on_hover(event)
        flow.canvas.draw()
        return flow._hover_annotation

    def render_314(self, flow, name):
        output = os.environ.get('PROGTRACK_314_EVIDENCE_DIR')
        if output:
            folder = Path(output)
            folder.mkdir(parents=True, exist_ok=True)
            flow.figure.savefig(folder / name, dpi=150)

    def assert_hover_314(self, flow, source, side):
        ann = flow._hover_annotation
        self.assertIsNotNone(ann)
        self.assertTrue(ann.get_visible())
        renderer = flow.canvas.get_renderer()
        box = ann.get_bbox_patch().get_window_extent(renderer)
        source_box = flow._hover_marker_bounds(source, renderer)
        self.assertFalse(box.overlaps(source_box))
        canvas = flow.figure.bbox
        self.assertGreaterEqual(box.x0, canvas.x0)
        self.assertLessEqual(box.x1, canvas.x1)
        self.assertGreaterEqual(box.y0, canvas.y0)
        self.assertLessEqual(box.y1, canvas.y1)
        self.assertEqual(ann._flow_side, side)
        from matplotlib.transforms import Bbox
        owned = [source_box]
        owner = ann._flow_owner
        label = flow.node_label_artists.get(owner)
        texts = ([label] if label is not None else []) + [
            bar['annotation'] for bar in flow.lifebar_data.get(owner, []) if 'annotation' in bar]
        for text in texts:
            if text.get_visible() and text.axes is flow.ax:
                patch = text.get_bbox_patch()
                owned.append(patch.get_window_extent(renderer) if patch else text.get_window_extent(renderer))
        stack = Bbox.union(owned)
        gap = stack.y0 - box.y1 if side == 'below' else box.y0 - stack.y1
        self.assertAlmostEqual(gap, renderer.points_to_pixels(2), places=6,
                               msg='Exactly two points from the visible owning stack, not its picker halo')
        return box

    def test_314_actual_node_lifebar_owner_below_stack_and_no_hover_churn(self):
        flow, source = self.flow_314()
        snapshot = copy.deepcopy((flow.manual_data, flow.settings, flow.temp_positions))
        ann = self.hover_314(flow, source)
        box = self.assert_hover_314(flow, source, 'below')
        label = flow.node_label_artists['test-1'].get_window_extent(flow.canvas.get_renderer())
        self.assertLess(box.y1, label.y0)
        self.assertFalse(ann.get_picker())
        self.assertFalse(ann.get_mouseover())
        bar = flow.lifebar_data['test-1'][0]['annotation']
        ann = self.hover_314(flow, bar)
        self.assertIn('0.50', ann._flow_text)
        self.assertIs(ann._flow_source, source)
        first = self.assert_hover_314(flow, source, 'below').bounds
        from matplotlib.backend_bases import MouseEvent
        bbox = bar.get_window_extent(flow.canvas.get_renderer())
        with patch.object(flow.canvas, 'draw_idle', wraps=flow.canvas.draw_idle) as draw:
            for i in range(40):
                x = bbox.x0 + bbox.width * (.3 if i % 2 else .7)
                event = MouseEvent('motion_notify_event', flow.canvas, x, (bbox.y0 + bbox.y1) / 2)
                flow._on_hover(event)
                self.assertIs(flow._hover_annotation, ann)
            self.assertEqual(draw.call_count, 0)
        self.assertEqual(self.assert_hover_314(flow, source, 'below').bounds, first)
        self.assertEqual(sum(a.get_zorder() == 20000 for a in flow.figure.artists), 1)
        self.assertEqual((flow.manual_data, flow.settings, flow.temp_positions), snapshot)
        self.render_314(flow, 'lifebar-below.png')

    def test_314_bottom_edge_above_and_horizontal_clamping(self):
        for point, side in [((.02, .7), 'below'), ((.98, .7), 'below'), ((.5, .035), 'above')]:
            with self.subTest(point=point):
                flow, source = self.flow_314(point)
                self.hover_314(flow, source)
                self.assert_hover_314(flow, source, side)
                if side == 'above':
                    self.render_314(flow, 'bottom-edge-above.png')

    def test_314_zoom_pan_resize_dpi_and_dragged_source_follow_rendered_marker(self):
        flow, source = self.flow_314()
        self.hover_314(flow, source)
        for dpi, size, limits, point in [(100, (6, 5), (0, 1), (.4, .7)),
                                         (150, (9, 6), (.1, .9), (.6, .65)),
                                         (200, (5, 5), (.2, .8), (.5, .6))]:
            with self.subTest(dpi=dpi, size=size):
                flow.figure.set_dpi(dpi)
                flow.figure.set_size_inches(*size)
                flow.ax.set(xlim=limits, ylim=limits)
                source.set_data([point[0]], [point[1]])
                flow.canvas.draw()
                self.assert_hover_314(flow, source, 'below')
        flow.ax.set_xlim(.7, .9)
        flow.canvas.draw()
        self.assertIsNone(flow._hover_annotation)

    def test_314_long_localized_text_wraps_without_losing_content_or_font_size(self):
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                flow, source = self.flow_314(point=(.5, .85), figsize=(5, 6), dpi=150)
                messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                text = ('test-1-' + 'X' * 90 + '\n' + messages['flow_track.settings.metric.usable_from_fert']
                        + '\n' + messages['flow_track.settings.metric.fert_rate_in_vivo'])
                flow._show_hover_annotation(text, .5, .85, 'test-1')
                flow.canvas.draw()
                ann = flow._hover_annotation
                self.assert_hover_314(flow, source, 'below')
                self.assertEqual(''.join(ann.get_text().split()), ''.join(text.split()))
                self.assertEqual(ann.get_fontsize(), 10)
                if language in ('de', 'ru'):
                    self.render_314(flow, f'long-content-{language}.png')

    def test_314_unrelated_overlap_never_displaces_flips_or_hides_tooltip(self):
        for point, side in [((.5, .65), 'below'), ((.5, .035), 'above')]:
            with self.subTest(side=side):
                flow, source = self.flow_314(point)
                flow._show_hover_annotation('test-1\nline two\nline three', *point, 'test-1')
                flow.canvas.draw()
                baseline = self.assert_hover_314(flow, source, side).frozen()
                center = flow.ax.transData.inverted().transform(baseline.get_points().mean(axis=0))
                other, = flow.ax.plot(*center, 'o', markersize=12, picker=10)
                flow.node_artists[other] = 'test-2'
                embryo, = flow.ax.plot(*center, 'o', markersize=10, picker=10)
                flow.embryo_artists[embryo] = ('transfer-test', 'test-embryo-1')
                flow.node_label_artists['test-2'] = flow.ax.annotate(
                    'unrelated label', center, ha='center', va='center', clip_on=False)
                bar = flow.ax.annotate('unrelated bar', center, ha='center', va='center',
                                       bbox={'facecolor': 'white'}, clip_on=False)
                flow.lifebar_data['test-2'] = [{'annotation': bar}]
                flow.ax.axhline(center[1])
                flow.ax.axvline(center[0])
                before = copy.deepcopy((flow.manual_data, flow.settings, flow.temp_positions))
                node_positions = (tuple(source.get_xdata()), tuple(source.get_ydata()),
                                  tuple(other.get_xdata()), tuple(other.get_ydata()))
                flow.canvas.draw()
                box = self.assert_hover_314(flow, source, side)
                self.assertEqual(box.bounds, baseline.bounds)
                self.assertTrue(box.overlaps(other.get_window_extent(flow.canvas.get_renderer())))
                self.assertTrue(box.overlaps(embryo.get_window_extent(flow.canvas.get_renderer())))
                self.assertFalse(flow._hover_annotation.get_picker())
                self.assertEqual((flow.manual_data, flow.settings, flow.temp_positions), before)
                self.assertEqual((tuple(source.get_xdata()), tuple(source.get_ydata()),
                                  tuple(other.get_xdata()), tuple(other.get_ydata())), node_positions)
                self.render_314(flow, f'neighbor-overlap-{side}.png')

    def test_314_close_gap_ignores_pick_halo_and_hidden_own_labels_bars(self):
        for dpi in (100, 150, 200):
            for labels, bars in ((True, True), (False, True), (True, False), (False, False)):
                with self.subTest(dpi=dpi, labels=labels, bars=bars):
                    flow, source = self.flow_314(dpi=dpi)
                    source.set_pickradius(1000)
                    flow.node_label_artists['test-1'].set_visible(labels)
                    for bar in flow.lifebar_data['test-1']:
                        bar['annotation'].set_visible(bars)
                    flow._show_hover_annotation('test-1\nsource statistics', .5, .65, 'test-1')
                    flow.canvas.draw()
                    box = self.assert_hover_314(flow, source, 'below')
                    self.assertAlmostEqual(box.get_points().mean(axis=0)[0],
                                           source.get_window_extent(flow.canvas.get_renderer()).x0
                                           + source.get_window_extent(flow.canvas.get_renderer()).width / 2)
                    if not labels and not bars:
                        expected = flow.canvas.get_renderer().points_to_pixels(
                            source.get_markersize() / 2 + source.get_markeredgewidth() / 2 + 2)
                        self.assertAlmostEqual(.65 * flow.figure.bbox.height - box.y1, expected)

    def test_314_impossible_canvas_recovers_without_changing_full_text(self):
        flow, source = self.flow_314()
        flow.figure.set_size_inches(2, 1)
        text = 'test-1\n' + '\n'.join(['long line'] * 12)
        flow._show_hover_annotation(text, .5, .65, 'test-1')
        annotation = flow._hover_annotation
        flow.canvas.draw()
        self.assertFalse(flow._hover_annotation.get_visible())
        flow.figure.set_size_inches(8, 9)
        flow.canvas.draw()
        self.assert_hover_314(flow, source, 'below')
        self.assertIs(flow._hover_annotation, annotation)
        self.assertEqual(''.join(annotation.get_text().split()), ''.join(text.split()))

    def test_314_real_mouse_right_click_double_click_and_drag_remain_available(self):
        from matplotlib.backend_bases import MouseEvent
        flow, source = self.flow_314()
        flow.node_artists[source] = 'test-3'
        flow._get_transfers_for_surrogate = lambda name: [{'id': 'transfer-1'}] if name == 'test-3' else []
        opened = []
        flow._show_animal_info = lambda name: opened.append(name)
        flow._open_freezer_embryo_list = lambda: opened.append('FREEZER')
        for event, callback in [('pick_event', flow._on_pick), ('motion_notify_event', flow._on_hover),
                                ('button_press_event', flow._on_mouse_press),
                                ('button_release_event', flow._on_mouse_release),
                                ('motion_notify_event', flow._on_mouse_move)]:
            flow.canvas.mpl_connect(event, callback)
        counts = {name: len(callbacks) for name, callbacks in flow.canvas.callbacks.callbacks.items()}
        self.hover_314(flow, source)
        def mouse(name, button, *, dblclick=False, point=(.5, .65)):
            x, y = flow.ax.transData.transform(point)
            event = MouseEvent(name, flow.canvas, x, y, button=button, dblclick=dblclick,
                               buttons={button} if name == 'motion_notify_event' else None)
            flow.canvas.callbacks.process(name, event)
        mouse('button_press_event', 3)
        self.assertTrue(flow.embryo_visibility['test-3'])
        self.assertIsNone(flow._hover_annotation)
        mouse('button_release_event', 3)
        self.hover_314(flow, source)
        mouse('button_press_event', 3)
        self.assertFalse(flow.embryo_visibility['test-3'])
        mouse('button_release_event', 3)
        # Close placement may occupy the source's invisible picker halo, but the
        # passive Text must not steal right-clicks from that still-valid target.
        source.set_pickradius(20)
        flow._show_hover_annotation('test-3', .5, .65, 'test-3')
        flow.canvas.draw()
        box = self.assert_hover_314(flow, source, 'below')
        pick_pixel = (flow.ax.transData.transform((.5, .65))[0], box.y1 - 1)
        self.assertTrue(box.contains(*pick_pixel))
        hit = MouseEvent('button_press_event', flow.canvas, *pick_pixel, button=3)
        self.assertTrue(source.contains(hit)[0])
        pick_data = flow.ax.transData.inverted().transform(pick_pixel)
        mouse('button_press_event', 3, point=pick_data)
        mouse('button_release_event', 3, point=pick_data)
        self.assertTrue(flow.embryo_visibility['test-3'])
        self.assertIsNone(flow._hover_annotation)
        mouse('button_press_event', 3)
        mouse('button_release_event', 3)
        self.assertFalse(flow.embryo_visibility['test-3'])
        self.hover_314(flow, source)
        mouse('button_press_event', 1, dblclick=True)
        mouse('button_release_event', 1)
        self.assertEqual(opened, ['test-3'])
        self.hover_314(flow, source)
        mouse('button_press_event', 1)
        mouse('motion_notify_event', 1, point=(.6, .7))
        mouse('button_release_event', 1, point=(.6, .7))
        self.assertEqual(opened, ['test-3'])  # Drag is not another details click.
        self.assertIn(('animal', 'test-3'), flow.temp_positions)
        self.assertFalse(flow.embryo_visibility['test-3'])  # No left-click visibility toggle.
        mouse('button_press_event', 2, point=(.8, .3))
        mouse('motion_notify_event', 2, point=(.85, .35))
        mouse('button_release_event', 2, point=(.85, .35))
        self.assertFalse(flow.pan_active)
        self.assertEqual(counts, {name: len(callbacks) for name, callbacks in flow.canvas.callbacks.callbacks.items()})
        flow.node_artists[source] = 'FREEZER'
        current_point = tuple(float(values[0]) for values in source.get_data())
        self.hover_314(flow, source)
        mouse('button_press_event', 3, point=current_point)
        self.assertTrue(flow.embryo_visibility['FREEZER'])
        mouse('button_release_event', 3, point=current_point)
        self.hover_314(flow, source)
        mouse('button_press_event', 1, dblclick=True, point=current_point)
        mouse('button_release_event', 1, point=current_point)
        self.assertEqual(opened[-1], 'FREEZER')

    def test_314_cleanup_on_leave_redraw_actor_backend_hide_and_destruction(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        from matplotlib.backend_bases import MouseEvent
        flow, source = self.flow_314()
        self.hover_314(flow, source)
        flow._on_hover(MouseEvent('motion_notify_event', flow.canvas, -1, -1))
        self.assertIsNone(flow._hover_annotation)
        for change in ('actor', 'backend', 'epoch'):
            with self.subTest(change=change):
                self.hover_314(flow, source)
                if change == 'actor':
                    flow.parent_app.master_track.current_username = 'Manager'
                elif change == 'backend':
                    flow.parent_app.backend = object()
                else:
                    flow.parent_app._view_session_epoch = 1
                flow.canvas.draw()
                self.assertIsNone(flow._hover_annotation)
        flow._install_hover_cleanup()
        self.hover_314(flow, source)
        self.qt.sendEvent(flow.widget, QEvent(QEvent.Type.Hide))
        self.assertIsNone(flow._hover_annotation)
        self.hover_314(flow, source)
        flow.parent_app.animals = {}
        flow.stack = SimpleNamespace(setCurrentWidget=lambda *_: None)
        flow.splash_widget = object()
        FlowTrackWidget._redraw_canvas(flow)
        self.assertIsNone(flow._hover_annotation)
        flow._show_hover_annotation('test-1', .5, .65, 'test-1')
        flow.widget.deleteLater()
        self.qt.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.assertIsNone(flow._hover_annotation)

    def test_314_real_qt_widget_initialization_leave_and_hide_lifecycle(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        from matplotlib.backend_bases import Event
        fixture = self.flow_312()
        parent = fixture.parent_app
        from PyQt6 import QtGui
        parent.QtGui = QtGui
        parent.backend = SimpleNamespace(records=Records())
        parent.matplotlib = __import__('matplotlib')
        parent.backend.records.put('reproduction', 'flow', {'version': '3.0', 'manual_data': fixture.manual_data})
        flow = FlowTrackWidget(parent, fixture.messages)
        flow.widget.setParent(self.hosts[-1])
        flow.widget.resize(800, 600)
        flow.widget.show()
        self.qt.processEvents()
        flow.canvas.draw()
        before = copy.deepcopy(parent.backend.records.data)
        source = next(artist for artist, name in flow.node_artists.items() if name == 'test-1')
        self.hover_314(flow, source)
        self.assert_hover_314(flow, source, 'below')
        flow.canvas.callbacks.process('figure_leave_event', Event('figure_leave_event', flow.canvas))
        self.assertIsNone(flow._hover_annotation)
        self.hover_314(flow, source)
        flow.widget.hide()
        self.qt.processEvents()
        self.assertIsNone(flow._hover_annotation)
        self.assertEqual(parent.backend.records.data, before)

    def flow_312(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        flow = FlowTrackWidget.__new__(FlowTrackWidget)
        flow.messages = self.host().messages
        roles = SimpleNamespace(SPENDER=SimpleNamespace(value='egg_cell_donor'),
                                SAMENSP=SimpleNamespace(value='sperm_donor'),
                                AMME=SimpleNamespace(value='surrogate'))
        flow.parent_app = SimpleNamespace(Role=roles, QtWidgets=MAIN.QtWidgets,
                                          QtCore=MAIN.QtCore, animals={
            'test-1': {'rolle': 'egg_cell_donor', 'name': 'test-1', 'events': []},
            'test-2': {'rolle': 'sperm_donor', 'name': 'test-2', 'events': [], 'sperm': []},
            'test-3': {'rolle': 'surrogate', 'name': 'test-3', 'events': []}})
        flow.widget = QWidget(self.hosts[-1])
        flow._can = lambda *_: True
        flow._animal_export_name = lambda name: name
        flow._save_manual_data = lambda: None
        flow._redraw_canvas = lambda: None
        flow.settings, flow.ui_preferences = flow._get_default_settings(), {}
        flow.egg_surgeries, flow.sperm_donations, flow.embryo_transfers = [], [], []
        flow.lifebar_data = {}
        figure = Figure(figsize=(8, 6))
        flow.canvas = FigureCanvasAgg(figure)
        flow.ax = figure.add_subplot(111)
        flow.manual_data = {'egg_donors': {'test-1': {'surgeries': {'egg-1': {
            'date': '2026-01-01', 'gv_isolated': 12, 'm1_to_m2': 10,
            'in_vitro_m2_fertilized': 8, 'embryos_from_in_vitro_m2': 6}}}},
            'sperm_donors': {'test-2': {'donations': {'sperm-1': {
                'date': '2026-01-02', 'applied_to_in_vitro_m2': 10,
                'fertilized_in_vitro_m2': 8, 'applied_to_in_vivo_m2': 2,
                'fertilized_in_vivo_m2': 1, 'embryos_from_in_vitro_m2': 6}}}},
            'transfers_by_id': {'transfer-1': {'surrogate_name': 'test-3',
                'transfer_date': '2026-01-09', 'embryos': []},
                'freezer_main': {'surrogate_name': 'FREEZER', 'transfer_date': None, 'embryos': []}}}
        for i in range(6):
            embryo = {'embryo_id': f'test-E{i + 1}', 'stage': 'in_vitro_m2',
                      'egg_donor_name': 'test-1', 'egg_donation_date': 'egg-1',
                      'sperm_donor_name': 'test-2', 'sperm_donation_id': 'sperm-1',
                      'implanted': i == 0, 'cryopreserved': i >= 4, 'freeze_date': '2026-01-03'}
            flow.manual_data['transfers_by_id']['transfer-1' if i < 4 else 'freezer_main']['embryos'].append(embryo)
        return flow

    def test_312_all_twelve_metric_contracts_and_actual_lifebars(self):
        flow = self.flow_312()
        cases = [
            ('egg_donor', 'test-1', 'SPENDER', flow._calculate_egg_donor_efficiency, 'per_surgery', {
                'ivm_rate': 100 * 10 / 12, 'fert_rate_after_ivm': 80,
                'usable_from_fert': 75, 'implantation_rate': 25,
                'usable_per_gv': .5, 'fert_per_gv': 8 / 12}),
            ('sperm_donor', 'test-2', 'SAMENSP', flow._calculate_sperm_donor_efficiency, 'per_donation', {
                'fert_rate_ivm': 80, 'usable_per_ivm': .6,
                'implantation_rate': 25, 'fert_rate_in_vivo': 50}),
            ('surrogate', 'test-3', 'AMME', flow._calculate_surrogate_efficiency, 'per_transfer', {
                'total_implantation': 25, 'transfer_success_rate': 100})]
        snapshot = copy.deepcopy(flow.manual_data)
        for prefix, animal, role, calculate, per, metrics in cases:
            result = calculate(animal)
            self.assertFalse(result['issues'])
            self.assertEqual(result[per][0]['date'].date().isoformat(),
                             {'test-1': '2026-01-01', 'test-2': '2026-01-02', 'test-3': '2026-01-09'}[animal])
            for metric, expected in metrics.items():
                for aggregation in ('total', per):
                    with self.subTest(role=prefix, metric=metric, aggregation=aggregation):
                        flow.settings[prefix + '_metric'] = metric
                        flow.settings[prefix + '_aggregation'] = aggregation
                        self.assertEqual(flow._render_lifebars_for_animal(animal, 0, 0, getattr(flow.parent_app.Role, role).value), 1)
                        actual = flow.lifebar_data[animal][0]
                        ratio = metric in {'usable_per_gv', 'fert_per_gv', 'usable_per_ivm'}
                        self.assertAlmostEqual(actual['efficiency_pct'], expected * (100 if ratio else 1))
                        self.assertEqual(actual['metric'], metric)
                        self.assertTrue(actual['annotation'].get_text())
                        self.assertNotIn('None', actual['annotation'].get_text())
        self.assertEqual(flow.manual_data, snapshot)

    def test_312_production_gap_unassigned_move_and_unconfirmed_implantation(self):
        flow = self.flow_312()
        transfers = flow.manual_data['transfers_by_id']
        unassigned = transfers['freezer_main']['embryos'].pop()
        transfers['unassigned'] = {'surrogate_name': None, 'transfer_date': None, 'embryos': [unassigned]}
        total = flow._calculate_egg_donor_efficiency('test-1')['total']
        self.assertEqual((total['total_produced_ivm'], total['total_usable_ivm']), (6, 5))
        self.assertFalse(flow._calculate_egg_donor_efficiency('test-1')['issues'])
        transfers.pop('unassigned')  # Legitimate produced > individually recorded gap.
        self.assertFalse(flow._calculate_egg_donor_efficiency('test-1')['issues'])
        frozen = transfers['freezer_main']['embryos'].pop()
        transfers['transfer-1']['embryos'].append(frozen)  # Keep historical cryopreserved=True.
        total = flow._calculate_sperm_donor_efficiency('test-2')['total']
        self.assertEqual((total['total_produced_ivm'], total['total_usable_ivm'], total['total_transferred_ivm'], total['total_frozen_ivm']), (6, 5, 5, 0))
        transfers['transfer-1']['embryos'] = transfers['transfer-1']['embryos'][:3]
        self.assertAlmostEqual(flow._calculate_surrogate_efficiency('test-3')['total']['total_implantation_pct'], 100 / 3)
        transfers['transfer-1']['embryos'][0]['implanted'] = False
        self.assertEqual(flow._calculate_egg_donor_efficiency('test-1')['total']['implantation_rate'], 0)
        transfers['transfer-1']['embryos'] = []
        self.assertIsNone(flow._calculate_surrogate_efficiency('test-3')['total']['total_implantation_pct'])

    def test_312_zero_missing_pooled_and_in_vivo_separation(self):
        flow = self.flow_312()
        rows = flow.manual_data['egg_donors']['test-1']['surgeries']
        rows['egg-2'] = {'date': '2026-02-01', 'gv_isolated': 100, 'm1_to_m2': 1,
                         'in_vitro_m2_fertilized': 0, 'embryos_from_in_vitro_m2': 0}
        total = flow._calculate_egg_donor_efficiency('test-1')['total']
        self.assertAlmostEqual(total['ivm_rate'], 11 / 112 * 100)
        self.assertAlmostEqual(total['usable_from_fert'], 75)
        self.assertEqual(flow._calculate_egg_donor_efficiency('test-1')['per_surgery'][1]['fert_rate_after_ivm'], 0)
        for bad in (None, -1, True, '12'):
            with self.subTest(missing_or_invalid=bad):
                rows['egg-2']['gv_isolated'] = bad
                result = flow._calculate_egg_donor_efficiency('test-1')
                self.assertIsNone(result['total']['ivm_rate'])
                self.assertIsNone(result['total']['usable_per_gv'])
                self.assertTrue(result['issues'])
                self.assertEqual(flow._editable_count(bad), 0)
        rows['egg-2']['gv_isolated'] = 0
        self.assertIsNone(flow._calculate_egg_donor_efficiency('test-1')['per_surgery'][1]['usable_per_gv'])
        flow.manual_data['transfers_by_id']['transfer-1']['embryos'][0]['stage'] = 'in_vivo_m2'
        sperm = flow._calculate_sperm_donor_efficiency('test-2')['total']
        self.assertEqual(sperm['total_transferred_ivm'], 3)
        self.assertEqual(sperm['fert_rate_in_vivo'], 50)
        self.assertEqual(flow._calculate_surrogate_efficiency('test-3')['total']['total_embryos'], 4)

    def test_312_duplicate_conflict_unlinked_and_overallocation_diagnostics(self):
        flow = self.flow_312()
        transfers = flow.manual_data['transfers_by_id']
        embryos = transfers['transfer-1']['embryos']
        embryos.append(copy.deepcopy(embryos[0]))
        result = flow._calculate_egg_donor_efficiency('test-1')
        self.assertEqual(result['total']['total_transferred_ivm'], 4)
        self.assertTrue(any('Duplicate' in text for text in result['issues']))
        transfers['freezer_main']['embryos'].append(embryos.pop())
        result = flow._calculate_egg_donor_efficiency('test-1')
        self.assertEqual(result['total']['total_transferred_ivm'], 3)
        self.assertTrue(any('Conflicting' in text for text in result['issues']))
        embryos[1]['sperm_donation_id'] = 'missing'
        self.assertTrue(any('Unresolved' in text for text in flow._calculate_sperm_donor_efficiency('test-2')['issues']))
        flow.manual_data['egg_donors']['test-1']['surgeries']['egg-1']['embryos_from_in_vitro_m2'] = 1
        self.assertTrue(any('exceed' in text for text in flow._calculate_egg_donor_efficiency('test-1')['issues']))
        self.assertEqual(flow._calculate_egg_donor_efficiency('test-1')['total']['total_produced_ivm'], 1)

    def test_312_settings_supported_defaults_persisted_choices_and_localization(self):
        flow = self.flow_312()
        stored = {}
        flow._settings_store = SimpleNamespace(load=lambda default: copy.deepcopy(stored), save=lambda value: stored.update(copy.deepcopy(value)))
        flow._load_settings()
        self.assertEqual(flow.settings['egg_donor_metric'], 'usable_per_gv')
        self.assertEqual(flow.settings['sperm_donor_metric'], 'usable_per_ivm')
        flow.settings.update(egg_donor_metric='fert_per_gv', egg_donor_aggregation='per_surgery')
        flow._save_settings()
        flow.settings = {}
        flow._load_settings()
        self.assertEqual(flow.settings['egg_donor_metric'], 'fert_per_gv')
        self.assertEqual(flow.settings['egg_donor_aggregation'], 'per_surgery')
        stored['settings'].update(egg_donor_metric='unsupported', sperm_donor_aggregation='per_surgery')
        flow._load_settings()
        self.assertEqual(flow.settings['egg_donor_metric'], 'usable_per_gv')
        self.assertEqual(flow.settings['sperm_donor_aggregation'], 'total')
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                flow.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                self.assertEqual(flow._format_efficiency(None), flow.messages['flow_track.no_data'])
                self.assertEqual(flow._format_ratio(None), flow.messages['flow_track.no_data'])
                self.assertIn('test-E1', flow._statistics_warning([flow._stat_issue('conflict', 'test-E1')]))

    def test_312_actual_settings_controls_and_hover(self):
        from matplotlib.backend_bases import MouseEvent
        flow = self.flow_312()
        flow.settings.update(surrogate_aggregation='per_transfer',
                             sperm_donor_aggregation='per_donation')
        flow._toggle_timeline_visibility = lambda *_: None
        saved = []
        flow._save_settings = lambda: saved.append(copy.deepcopy(flow.settings))
        def accept_settings(dialog):
            combos = dialog.findChildren(QComboBox)
            self.assertEqual([combo.count() for combo in combos], [6, 4, 2])
            self.assertEqual([combo.currentData() for combo in combos],
                             ['usable_per_gv', 'usable_per_ivm', 'total_implantation'])
            checked = [rb.text() for rb in dialog.findChildren(QRadioButton) if rb.isChecked()]
            self.assertIn(flow.messages['flow_track.settings.aggregation.per_transfer'], checked)
            self.assertIn(flow.messages['flow_track.settings.aggregation.per_donation'], checked)
            combos[0].setCurrentIndex(0)
            return QDialog.DialogCode.Accepted
        with patch.object(QDialog, 'exec', accept_settings):
            flow._open_settings_dialog()
        self.assertEqual(saved[-1]['egg_donor_metric'], 'ivm_rate')
        self.assertEqual(saved[-1]['surrogate_aggregation'], 'per_transfer')
        flow.settings.update(egg_donor_metric='usable_per_gv', egg_donor_aggregation='total')
        flow._hover_annotation = None
        flow._render_lifebars_for_animal('test-1', .5, .5, 'egg_cell_donor')
        flow.canvas.draw()
        ann = flow.lifebar_data['test-1'][0]['annotation']
        bbox = ann.get_window_extent(flow.canvas.get_renderer())
        event = MouseEvent('motion_notify_event', flow.canvas, (bbox.x0 + bbox.x1) / 2, (bbox.y0 + bbox.y1) / 2)
        shown = []
        flow._show_hover_annotation = lambda text, *_: shown.append(text)
        flow._on_hover(event)
        self.assertEqual(len(shown), 1)
        self.assertIn('0.50', shown[0])
        self.assertNotIn('50.0%', shown[0])
        flow.lifebar_data['test-1'][0]['efficiency_pct'] = None
        flow._on_hover(event)
        self.assertIn(flow.messages['flow_track.no_data'], shown[-1])

    def test_312_embryo_deletion_preserves_production_and_cancelled_dialogs_are_read_only(self):
        flow = self.flow_312()
        flow._save_flow_track_data = lambda: None
        before = copy.deepcopy(flow.manual_data)
        with patch.object(MAIN.QtWidgets.QMessageBox, 'information'):
            flow._delete_embryo('freezer_main', 'test-E6')
        self.assertEqual(flow.manual_data['egg_donors'], before['egg_donors'])
        self.assertEqual(flow.manual_data['sperm_donors'], before['sperm_donors'])
        self.assertEqual(flow._calculate_egg_donor_efficiency('test-1')['total']['total_usable_ivm'], 5)
        def reject(dialog):
            self.assertTrue(dialog.findChildren(MAIN.QtWidgets.QLabel))
            return QDialog.DialogCode.Rejected
        # No source procedures/events need adding to inspect calculated metrics.
        snapshot = copy.deepcopy(flow.manual_data)
        with patch.object(QDialog, 'exec', reject):
            flow._show_egg_donor_dialog('test-1', flow.parent_app.animals['test-1'])
            flow._show_sperm_donor_dialog('test-2', flow.parent_app.animals['test-2'])
            flow._show_surrogate_dialog('test-3', flow.parent_app.animals['test-3'])
        self.assertEqual(flow.manual_data, snapshot)

    def test_312_sqlite_close_reopen_preserves_settings_allocation_and_counts(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        from Plugins.core.backend_store import BackendJsonStore
        with tempfile.TemporaryDirectory(prefix='flow-312-') as temporary:
            paths = runtime_paths(Path(temporary))
            backend = ProgTrackBackend(paths, acquire_process_lock=False, bootstrap_seed=False)
            flow = self.flow_312()
            flow.manual_data['total_embryo_count'] = 6
            flow._settings_store = BackendJsonStore(backend, 'configuration', 'flow-track')
            flow._data_store = BackendJsonStore(backend, 'reproduction', 'flow')
            flow.settings.update(surrogate_aggregation='per_transfer', egg_donor_metric='fert_per_gv')
            moved = flow.manual_data['transfers_by_id']['freezer_main']['embryos'].pop()
            flow.manual_data['transfers_by_id']['transfer-1']['embryos'].append(moved)
            try:
                flow._save_settings()
                FlowTrackWidget._save_flow_track_data(flow)
                expected = copy.deepcopy(flow.manual_data)
            finally:
                backend.close()
            backend = ProgTrackBackend(paths, acquire_process_lock=False, bootstrap_seed=False)
            try:
                reopened = self.flow_312()
                reopened._settings_store = BackendJsonStore(backend, 'configuration', 'flow-track')
                reopened._data_store = BackendJsonStore(backend, 'reproduction', 'flow')
                reopened._load_settings()
                reopened._load_flow_track_data()
                self.assertEqual(reopened.manual_data, expected)
                self.assertEqual(reopened.settings['egg_donor_metric'], 'fert_per_gv')
                self.assertEqual(reopened.settings['surrogate_aggregation'], 'per_transfer')
                total = reopened._calculate_egg_donor_efficiency('test-1')['total']
                self.assertEqual((total['total_produced_ivm'], total['total_transferred_ivm'],
                                  total['total_frozen_ivm'], total['total_usable_ivm']), (6, 5, 1, 6))
            finally:
                backend.close()

    def test_312_donor_editor_cancel_unknown_counts_and_metadata_preserved(self):
        flow = self.flow_312()
        rows = flow.manual_data['egg_donors']['test-1']['surgeries']
        rows['2026-01-01'] = rows.pop('egg-1')
        rows['2026-01-01'].update(gv_isolated=None, m1_isolated=1, m2_isolated=2,
                                  in_vivo_m2_fertilized=1, embryos_from_in_vivo_m2=0,
                                  frozen_from_in_vivo_m2=0, frozen_from_in_vitro_m2=2)
        for transfer in flow.manual_data['transfers_by_id'].values():
            for embryo in transfer['embryos']:
                embryo['egg_donation_date'] = '2026-01-01'
        animal = flow.parent_app.animals['test-1']
        animal['events'] = [{'event_type': 'oocyte_retrieval', 'datum': datetime(2026, 1, 1)}]
        snapshot = copy.deepcopy(flow.manual_data)
        def cancel_after_typing(dialog):
            spins = dialog.findChildren(MAIN.QtWidgets.QSpinBox)
            self.assertEqual(spins[0].text(), flow.messages['flow_track.no_data'])
            spins[2].setValue(9)
            return QDialog.DialogCode.Rejected
        with patch.object(QDialog, 'exec', cancel_after_typing):
            flow._show_egg_donor_dialog('test-1', animal)
        self.assertEqual(flow.manual_data, snapshot)
        flow._save_flow_track_data = lambda: None
        with patch.object(QDialog, 'exec', return_value=QDialog.DialogCode.Accepted), patch.object(
                MAIN.QtWidgets.QMessageBox, 'question', return_value=MAIN.QtWidgets.QMessageBox.StandardButton.Yes):
            flow._show_egg_donor_dialog('test-1', animal)
        row = rows['2026-01-01']
        self.assertIsNone(row['gv_isolated'])
        self.assertEqual((row['m1_isolated'], row['date'], row['embryos_from_in_vitro_m2']), (1, '2026-01-01', 6))
        self.assertEqual(flow.manual_data['transfers_by_id'], snapshot['transfers_by_id'])
        spin = MAIN.QtWidgets.QSpinBox()
        flow._set_count_input(spin, None)
        spin.lineEdit().setText('0')
        spin.lineEdit().textEdited.emit('0')
        spin.interpretText()
        self.assertEqual(flow._entered_count(spin), 0)  # An explicitly entered zero is known.

    def test_312_real_freezer_move_and_delete_survive_redraw_without_recount(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        flow = self.flow_312()
        flow.parent_app.matplotlib = __import__('matplotlib')
        flow.manual_data['test-3'] = {'transfers': {'transfer-1': {'embryo_count': 4}}}
        flow.embryo_transfers = [{'id': 'transfer-1', 'animal_name': 'test-3', 'date': datetime(2026, 1, 9)}]
        flow._populate_events_from_progtrack = lambda: None
        flow.stack = SimpleNamespace(setCurrentWidget=lambda *_: None)
        flow.content_widget = object()
        flow.current_xlim, flow.current_ylim = (0, 1), (0, 1)
        flow.temp_positions, flow.embryo_visibility = {}, {'FREEZER': True}
        flow.node_artists, flow.embryo_artists = {}, {}
        flow._save_flow_track_data = lambda: None
        flow._redraw_canvas = FlowTrackWidget._redraw_canvas.__get__(flow)
        listing = QListWidget()
        listing.addItem('test-E5')
        listing.setCurrentRow(0)
        parent = QDialog(flow.widget)
        def choose(dialog):
            dialog.findChildren(QListWidget)[0].setCurrentRow(0)
            return QDialog.DialogCode.Accepted
        with patch.object(QDialog, 'exec', choose):
            flow._move_freezer_embryo_to_surrogate(listing,
                flow.manual_data['transfers_by_id']['freezer_main']['embryos'], parent)
        total = flow._calculate_egg_donor_efficiency('test-1')['total']
        self.assertEqual((total['total_transferred_ivm'], total['total_frozen_ivm'], total['total_usable_ivm']), (5, 1, 6))
        self.assertEqual(flow.manual_data['test-3']['transfers']['transfer-1']['embryo_count'], 5)
        with patch.object(MAIN.QtWidgets.QMessageBox, 'information'):
            flow._delete_embryo('transfer-1', 'test-E5')
        total = flow._calculate_egg_donor_efficiency('test-1')['total']
        self.assertEqual((total['total_produced_ivm'], total['total_transferred_ivm'], total['total_frozen_ivm']), (6, 4, 1))
        self.assertEqual(flow.manual_data['test-3']['transfers']['transfer-1']['embryo_count'], 4)

    def test_312_exports_share_all_metrics_and_undefined_values(self):
        from openpyxl import Workbook, load_workbook
        from io import BytesIO
        flow = self.flow_312()
        wb = Workbook()
        flow._create_all_animals_overview(wb)
        flow._create_egg_donors_overview(wb)
        flow._create_sperm_donors_overview(wb)
        flow._create_surrogates_overview(wb)
        flow._create_egg_donor_sheet(wb, 'test-1')
        flow._create_sperm_donor_sheet(wb, 'test-2')
        flow._create_surrogate_sheet(wb, 'test-3')
        flow._create_all_embryos_sheet(wb)
        flow._create_freezer_sheet(wb)
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        reopened = load_workbook(stream, data_only=True)
        egg, sperm, surrogate = [reopened[name] for name in ('Egg Donors Overview', 'Sperm Donors Overview', 'Surrogates Overview')]
        self.assertEqual((egg['M2'].value, egg['N2'].value, egg['P2'].value), (75, .5, 6))
        self.assertEqual((sperm['K2'].value, sperm['L2'].value, sperm['M2'].value), (.6, 50, 6))
        self.assertEqual((surrogate['F2'].value, surrogate['G2'].value), (25, 100))
        overview = reopened['All Animals Overview']
        self.assertEqual([overview.cell(2, c).value for c in range(4, 9)], [6, 4, 1, 2, 'IVM MII'])
        self.assertEqual(reopened['Freezer'].max_row, 3)
        self.assertEqual(reopened['All Embryos'].max_row, 7)
        for ws in reopened.worksheets:
            self.assertFalse(any(isinstance(cell.value, str) and cell.value.startswith('=') for row in ws for cell in row))
        flow.manual_data['transfers_by_id']['transfer-1']['embryos'] = []
        other = Workbook()
        flow._create_egg_donors_overview(other)
        self.assertEqual(other['Egg Donors Overview']['L2'].value, flow.messages['flow_track.no_data'])
        frozen = flow.manual_data['transfers_by_id']['freezer_main']['embryos']
        frozen.append(copy.deepcopy(frozen[0]))
        flow._create_freezer_sheet(other)
        self.assertEqual(other['Freezer'].max_row, 3)  # Duplicate ID is never a third stored embryo.

    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])

    def setUp(self):
        self.hosts = []

    def tearDown(self):
        for host in self.hosts:
            host.hide()
            host.deleteLater()
        self.qt.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.qt.processEvents()

    def host(self):
        app = MAIN.ProgTrackApp.__new__(MAIN.ProgTrackApp)
        QMainWindow.__init__(app)
        self.hosts.append(app)
        app.messages = json.loads((ROOT / 'lang/messages_en.json').read_text(encoding='utf-8'))
        app.lang = 'en'
        app._setup_menus = lambda: None
        app.backend = SimpleNamespace(records=Records())
        app._test_audit_entries = []
        app.master_track = SimpleNamespace(
            current_username='tester',
            audit=lambda *args: app._test_audit_entries.append(args))
        app.animals, app.archived = {}, {}
        app.lst = QListWidget(app)
        app.lst.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        app.selected_animals = []
        app._disabled_plugins = set()
        app._master_can = lambda *_args, **_kwargs: True
        app._project_visibility_scope = lambda **_kwargs: (True, set())
        app._is_projects_track_active = lambda: False
        app._is_steroid_track_active = lambda: True
        app.has_cage_track_plugin = app.has_heritage_plugin = app.has_medi_track_plugin = False
        app.has_projects_plugin = app.has_pdg_plugin = False
        app.reports_enabled = app.flow_track_enabled = app.read_only_mode = False
        app.cage_track_plugin = app.heritage_plugin = app.projects_plugin = None
        app._install_animal_section_preferences = lambda _dlg: None
        app._wire_status_checkboxes = lambda *_a, **_k: None
        app._connect_project_severity_reset = lambda *_a, **_k: None
        app._save_trace = lambda *_a, **_k: None
        app._save_persistence = lambda *_a, **_k: None
        app._schedule_post_animal_save_project_updates = lambda *_a, **_k: None
        app._coerce_in_experiment_for_project = lambda value, *_a, **_k: value
        app._capture_death_lifecycle = lambda *_a, **_k: True
        app._validate_identity_species_for_save = lambda *_a, **_k: True
        app._validate_existing_identity_for_save = lambda *_a, **_k: True
        app._name_species_conflict = lambda *_a, **_k: False
        app._confirm_species_change_once = lambda *_a, **_k: True
        app._resolve_animal_key = lambda name, *_a, **_k: name
        app._load_core_parentage_records = lambda **_k: copy.deepcopy(app.animals)
        app._refresh_list = lambda **_k: None
        app._on_select = lambda: None
        app.errors = []
        def default_message(*args, **kwargs):
            if args and args[0] == app.messages.get('confirm.event_save.title'):
                return MAIN.QMessageBox.StandardButton.Yes
            return app.errors.append(args)
        app._show_message_raw = default_message
        app._show_message = lambda *args, **kwargs: app.errors.append(args)
        app._show_permission_denied = lambda: app.errors.append(('denied',))
        return app

    def record(self, role='experimental'):
        return {
            'name': 'Subject', 'rolle': role, 'species': 'Callithrix jacchus',
            'birth_date': '01.01.2020', 'id': 'TEST-ID', 'origin': 'Institute',
            'sex': 'Female', 'genotype': '', 'project': '', 'ref_weight': 488.1,
            'gewicht': [{'datum': datetime(2026, 1, 1), 'wert': 27.1}],
            'daten': [], 'pdg': [], 'events': [], 'sperm': [],
        }

    def _307_sqlite_host(self, population=10, *, seeded=False):
        """Real read/permission services; only disposable test databases."""
        from Plugins.Master_Track.plugin import MasterTrackPlugin
        from Plugins.core.authorization import AuthorizationService
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        app = self.host()
        app.backend = ProgTrackBackend(
            runtime_paths(root, application_root=ROOT if seeded else root),
            acquire_process_lock=False)
        self.addCleanup(app.backend.close)
        if not seeded:
            app.backend.records.put('security', 'users', [{
                'username': 'tester', 'role': 'user', 'jobs': ['manager'],
                'unit_id': 'unit-a', 'permissions': {'granted': [], 'revoked': []},
            }])
            app.backend.records.put('security', 'organization-units', {'units': [
                {'unit_id': key, 'display_name': key, 'active': True, 'archived': False}
                for key in ('unit-a', 'unit-b')
            ]})
            app.backend.records.put('projects', 'catalog', {'projects': {
                'Allowed': {'summary': {'contact1_login': 'tester'}},
                'Hidden': {'summary': {'contact1_login': 'another-user'}},
            }})
            active, archived = {}, {}
            for index in range(population):
                key = f'test-{index + 1}'
                record = self.record()
                record.update(name=key, ipid=f'IPID-307-{index}', id=f'ID-307-{index}',
                              species='Callitrix jacchus',
                              sex='Female' if index % 2 == 0 else 'Male',
                              birth_date='01.01.2010',
                              project='Hidden' if index % 5 == 1 else 'Allowed')
                if index % 5 in (0, 1):
                    record['organization_unit_id'] = 'unit-a'
                elif index % 5 == 2:
                    record['workgroup_id'] = 'unit-b'
                elif index % 5 == 3:
                    record['project'] = ''
                else:
                    record['unit_id'] = 'unit-b'  # Housing, NOT workgroup ownership.
                (active if index < population * 3 // 4 else archived)[key] = record
            app.backend.save_core_data({'animals': active, 'archived_animals': archived})
        app.master_track = MasterTrackPlugin(app)
        app.master_track.audit_logs_dir = lambda: str(root / 'test-audit')
        username = 'Manager' if seeded else 'tester'
        user = app.master_track.user_db.get_user(username)
        self.assertIsNotNone(user)
        app.master_track._current_username = username
        app.master_track._current_role = user['role']
        app.authorization = AuthorizationService(app.master_track, backend=app.backend)
        for name in ('_project_visibility_scope', '_master_can', '_load_core_parentage_records'):
            setattr(app, name, getattr(MAIN.ProgTrackApp, name).__get__(app, MAIN.ProgTrackApp))
        if seeded:
            # The application's loader performs the actual measurement/event
            # date conversion and startup normalization before any opening.
            MAIN.ProgTrackApp._load_persistence(app)
        else:
            snapshot = app.backend.load_core_data()
            app.animals, app.archived = snapshot['animals'], snapshot['archived_animals']
            for record in list(app.animals.values()) + list(app.archived.values()):
                for field in ('daten', 'gewicht', 'pdg', 'events', 'sperm'):
                    for row in record.get(field, []):
                        row['datum'] = MAIN.parse_date(row['datum'])
                record.setdefault('sick', False)
                record.setdefault('recovery_time', MAIN.DEFAULT_RECOVERY_TIME)
                MAIN.synchronize_experimental_limits(record, record.get('rolle'))
        return app

    def _306_sidebar_shell(self, app):
        """Real sidebar/state publication; no optional plugin initialization."""
        app.category_tab = QTabBar(app)
        app._category_custom_role_values = []
        for label in ('female', 'male', 'offspring', 'partner', 'breeding', 'experimental', 'all'):
            app.category_tab.addTab(label)
        app.category_tab.setCurrentIndex(6)
        app.cmb_arch = QComboBox(app)
        app.phase_widget = QWidget(app)
        app.phase_buttons = {None: QPushButton(app)}
        for name in ('box_chk', 'box_rad'):
            setattr(app, name, QGroupBox(app))
        for name in ('chk_prog', 'chk_weight', 'chk_events', 'chk_mode_blood',
                     'rb_weight_on', 'rb_weight_off', 'chk_events_offspring',
                     'chk_events_breeding', 'chk_events_experimental'):
            setattr(app, name, QCheckBox(app))
        for name in ('btn_new', 'btn_edit', 'btn_load_blood', 'btn_load_weights',
                     'btn_archive', 'btn_restore', 'btn_delete'):
            setattr(app, name, QPushButton(app))
        app.main_tabs = QTabWidget(app)
        app.plots_tab = QWidget(app)
        app.main_tabs.addTab(app.plots_tab, 'Plots')
        for name in ('_sync_plot_action_button_widths', '_update_animal_sex_filter_texts',
                     '_set_prog_event_plot_controls_visible', '_set_sperm_controls_visible',
                     '_apply_master_button_states'):
            setattr(app, name, lambda *_a, **_k: None)
        for name in ('_refresh_list', '_on_select'):
            setattr(app, name, getattr(MAIN.ProgTrackApp, name).__get__(app, MAIN.ProgTrackApp))
        # Controls-only part is covered by #270; keep real publication dispatch.
        app._on_category_selected = lambda _index: app._refresh_list()
        app._plot_selected = lambda: None

    def test_306_sidebar_scope_catalogue_and_projection_counts(self):
        for population in (10, 695):
            app = self._307_sqlite_host(population)
            self._306_sidebar_shell(app)
            for rec in app.animals.values():
                rec['in_experiment'] = True  # Tooltip catalogue reads count too.
            records = app.backend.records
            with patch.object(records, 'get', wraps=records.get) as reads, \
                 patch.object(app, '_update_category_tab_visibility', wraps=app._update_category_tab_visibility) as counts, \
                 patch.object(app, '_populate_animal_list', wraps=app._populate_animal_list) as rows:
                app._refresh_list(update_tab_visibility=True)
            queries = Counter(tuple(call.args[:2]) for call in reads.call_args_list)
            self.assertEqual(queries['projects', 'catalog'], 1)
            self.assertLessEqual(queries['security', 'users'], 2)
            self.assertEqual((counts.call_count, rows.call_count), (1, 1))
            self.assertIsNone(app._sidebar_read_context)
            self.assertEqual({app.lst.item(i).data(Qt.ItemDataRole.UserRole)
                              for i in range(app.lst.count())}, set(app.animals))

    def test_306_ui_batch_publishes_once_and_releases_signal_blockers(self):
        app = self._307_sqlite_host()
        self._306_sidebar_shell(app)
        received = []
        app.chk_events.toggled.connect(lambda checked: received.append(checked))
        with patch.object(app, '_populate_animal_list', wraps=app._populate_animal_list) as rows, \
             patch.object(app, '_update_category_tab_visibility', wraps=app._update_category_tab_visibility) as counts, \
             patch.object(app, '_plot_selected') as plots:
            with app._batch_ui_updates():
                app.chk_events.setChecked(True)
                app._refresh_list(update_tab_visibility=True)
                app._refresh_list(update_tab_visibility=True)
                app._update_category_tab_visibility()
                app._on_select()
                app._on_select()
            self.assertEqual(rows.call_count, 1)
            # Two requests were deferred; exactly one executes outside the batch.
            self.assertEqual(counts.call_count, 2)
            self.assertEqual(plots.call_count, 1)
        self.assertEqual(received, [])
        self.assertFalse(app.chk_events.signalsBlocked())
        with self.assertRaisesRegex(ValueError, 'injected'):
            with app._batch_ui_updates():
                raise ValueError('injected')
        self.assertEqual(app._ui_update_depth, 0)
        self.assertFalse(app.category_tab.signalsBlocked())

    def test_306_origin_dialog_coalescing_identity_and_teardown(self):
        app = self.host()
        calls = []
        origin = QDialog(app)
        origin.setModal(True)
        origin.show()
        self.qt.processEvents()
        self.assertIs(QApplication.activeModalWidget(), origin)
        for index in range(1000):
            app._queue_view_update('view', lambda i=index: calls.append(i))
        updates = app._after_dialog_updates
        self.assertEqual(len(updates._pending), 1)
        self.assertEqual(len(updates._origins), 1)
        self.assertEqual(len(updates.findChildren(MAIN.QTimer)), 1)
        self.assertFalse(updates._timer.isActive())
        self.qt.processEvents()
        self.assertEqual(calls, [])
        origin.accept()
        unrelated = QDialog(app)
        unrelated.setModal(True)
        unrelated.show()
        self.qt.processEvents()
        self.assertEqual(calls, [999], 'Unrelated modal must not block completed origin')
        for switch in ('epoch', 'backend', 'actor'):
            with self.subTest(switch=switch):
                app._queue_view_update('stale', lambda: calls.append('stale'))
                if switch == 'epoch':
                    app._invalidate_pending_view_updates()
                elif switch == 'backend':
                    app.backend = SimpleNamespace()
                else:
                    app.master_track.current_username = 'another-user'
                unrelated.accept()
                self.qt.processEvents()
                self.assertEqual(calls, [999])
                unrelated.show()
                self.qt.processEvents()
        unrelated.reject()
        app._queue_view_update('closed', lambda: calls.append('closed'))
        app._invalidate_pending_view_updates()
        self.qt.processEvents()
        self.assertEqual(calls, [999])

    def test_306_save_reads_definitions_once_preserves_historical_snapshots(self):
        app = self._307_sqlite_host()
        definition = normalize_event_definition({'event_type': 'custom_test', 'name': 'Test', 'active': False})
        app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [definition]})
        keys = list(app.animals)
        historical = {'label': 'Original label', 'revision': 7, 'color': '#123456'}
        app.animals[keys[0]]['events'] = [
            {'event_type': 'CUSTOM_TEST', 'datum': datetime(2026, 1, 1)},
            {'event_type': 'custom_test', 'datum': datetime(2026, 1, 2),
             'custom_event_snapshot': copy.deepcopy(historical)},
            {'event_type': 'fsh', 'datum': datetime(2026, 1, 3)},
        ] * 50
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app, MAIN.ProgTrackApp)
        records = app.backend.records
        with patch.object(records, 'get', wraps=records.get) as reads, \
             patch.object(app, '_write_json', return_value=None) as write:
            app._save_persistence(defer_post_save_work=True)
        queries = Counter(tuple(call.args[:2]) for call in reads.call_args_list)
        self.assertEqual(queries['configuration', 'role-block-presets'], 1)
        self.assertTrue(write.call_args.kwargs['audit_after_save'])
        events = app.animals[keys[0]]['events']
        self.assertEqual(events[0]['custom_event_snapshot']['label'], 'Test')
        self.assertEqual(events[1]['custom_event_snapshot'], historical)
        self.assertNotIn('custom_event_snapshot', events[2])

    def test_306_audits_are_synchronous_even_with_modal_and_failure_is_visible(self):
        app = self.host()
        origin = QDialog(app)
        origin.setModal(True)
        origin.show()
        self.qt.processEvents()
        calls = []
        app._master_audit = lambda *args: calls.append((app.master_track.current_username, args))
        app._audit_data_snapshot_diff_when_safe(
            {'animals': {'test-1': {'ref_weight': 1}}},
            {'animals': {'test-1': {'ref_weight': 2}}})
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0], 'tester')
        app.master_track.current_username = 'Guest'
        origin.accept()
        self.qt.processEvents()
        self.assertEqual(len(calls), 1)
        with patch.object(app, '_audit_data_snapshot_diff', side_effect=OSError('audit unavailable')):
            with self.assertRaisesRegex(OSError, 'audit unavailable'):
                app._audit_data_snapshot_diff_when_safe({}, {})

    def test_306_required_followups_run_only_after_commit_in_order(self):
        app = self.host()
        app.animals = {'test-1': self.record()}
        app._read_json = lambda: {'animals': copy.deepcopy(app.animals), 'archived_animals': {}}
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app, MAIN.ProgTrackApp)
        calls = []
        identity = MAIN._actor_backend_read_identity(app)
        def committed(_data, **_kwargs):
            calls.append('commit')
        app._write_json = committed
        app._pending_project_save_updates = [(identity, lambda: calls.append('project')),
                                             (identity, lambda: calls.append('clinical'))]
        app._save_persistence(True)
        self.assertEqual(calls, ['commit', 'project', 'clinical'])
        app._pending_project_save_updates = [(identity, lambda: calls.append('lost-commit'))]
        app._write_json = lambda *_a, **_k: (_ for _ in ()).throw(OSError('commit failed'))
        with self.assertRaisesRegex(OSError, 'commit failed'):
            app._save_persistence(True)
        self.assertNotIn('lost-commit', calls)
        self.assertNotIn('_pending_project_save_updates', app.__dict__)

    def test_306_trace_is_opt_in_and_logging_failures_do_not_commit(self):
        app = self.host()
        trace = MAIN.ProgTrackApp._save_trace.__get__(app, MAIN.ProgTrackApp)
        with patch.dict(os.environ, {'PROGTRACK_SAVE_TRACE': '0'}), \
             patch('builtins.open') as opened:
            trace('normal')
            opened.assert_not_called()
        with tempfile.TemporaryDirectory() as directory, patch.object(MAIN, 'LOG_DIR', Path(directory)), \
             patch.dict(os.environ, {'PROGTRACK_SAVE_TRACE': '1', 'PROGTRACK_SAVE_TRACE_FSYNC': '0'}), \
             patch.object(MAIN.os, 'fsync') as fsync:
            trace('diagnostic', count=1)
            self.assertIn('diagnostic', (Path(directory) / 'save_trace.log').read_text())
            fsync.assert_not_called()
            with patch('builtins.open', side_effect=OSError('log unavailable')):
                trace('failed diagnostic')  # Technical log failure is harmless.

    def _306_conventions(self):
        from Plugins.core.animal_roles import AnimalRoleRegistry
        from Plugins.core.identity_conventions import DEFAULT_CONVENTIONS
        app = self._307_sqlite_host()
        app.animal_role_registry = AnimalRoleRegistry(app.backend)
        app._load_identity_conventions = lambda: copy.deepcopy(DEFAULT_CONVENTIONS)
        app._can_configure_animal_roles = lambda: True
        app._can_manage_institution_settings = lambda: False
        app._can_manage_institution_branding = lambda: False
        app._save_role_label_overrides = lambda *_a: None
        app._merge_user_messages = lambda: None
        return app

    def test_306_conventions_noop_all_tabs_and_style_only_have_no_shared_writes(self):
        app = self._306_conventions()
        records = app.backend.records
        for tab_index in range(4):
            with self.subTest(tab=tab_index):
                dlg = MAIN.StyleSettingsDialog(app, app.messages)
                dlg.tabs.setCurrentIndex(tab_index)
                with patch.object(records, 'put', wraps=records.put) as put, \
                     patch.object(records, 'put_many', wraps=records.put_many) as many, \
                     patch.object(app, '_save_user_style_settings') as style, \
                     patch.object(app, '_refresh_list') as rows, \
                     patch.object(app, '_plot_selected') as plot:
                    dlg.accept()
                    self.assertEqual(dlg.result(), QDialog.DialogCode.Accepted)
                    app._apply_accepted_conventions(dlg)
                    put.assert_not_called()
                    many.assert_not_called()
                    style.assert_not_called()
                    rows.assert_not_called()
                    plot.assert_not_called()
                self.qt.processEvents()
                dlg.deleteLater()
        dlg = MAIN.StyleSettingsDialog(app, app.messages)
        dlg.color_buttons['weight'].setProperty('color', '#123456')
        with patch.object(records, 'put', wraps=records.put) as put, \
             patch.object(records, 'put_many', wraps=records.put_many) as many, \
             patch.object(app, '_save_user_style_settings') as style:
            dlg.accept()
            app._apply_accepted_conventions(dlg)
            put.assert_not_called()
            many.assert_not_called()
            style.assert_called_once()
        self.qt.processEvents()
        dlg.deleteLater()

    def test_306_conventions_identity_only_and_real_role_CAS(self):
        app = self._306_conventions()
        dlg = MAIN.StyleSettingsDialog(app, app.messages)
        editor = next(iter(dlg._lifecycle_catalog_edits.values()))
        editor.set_values(editor.values() + ['test-entry'])
        with patch.object(app, '_save_identity_conventions', wraps=app._save_identity_conventions) as identity, \
             patch.object(app.backend.records, 'put_many', wraps=app.backend.records.put_many) as roles:
            dlg.accept()
            app._apply_accepted_conventions(dlg)
            identity.assert_called_once()
            roles.assert_not_called()
        self.qt.processEvents()
        dlg.deleteLater()
        for concurrent in (False, True):
            with self.subTest(concurrent=concurrent):
                dlg = MAIN.StyleSettingsDialog(app, app.messages)
                dlg.role_table.item(0, 4).setText('test-role-label-2' if concurrent else 'test-role-label-1')
                if concurrent:
                    key = ('configuration', 'role-block-presets')
                    payload = app.backend.records.get(*key, default={})
                    app.backend.records.put(*key, payload)
                    prior = copy.deepcopy(app.animal_role_registry.roles())
                dlg.accept()
                self.assertIsNotNone(dlg.get_role_definitions())
                app.errors.clear()
                app._apply_accepted_conventions(dlg)
                if concurrent:
                    self.assertTrue(app.errors, 'Actual changed payload must retain revision conflicts')
                    self.assertEqual(app.animal_role_registry.roles(), prior)
                else:
                    self.assertEqual(app.errors, [])
                    self.assertEqual(app.animal_role_registry.roles()[0]['label'], 'test-role-label-1')
                self.qt.processEvents()
                dlg.deleteLater()

    def test_306_seeded_session_publication_counts_and_timings(self):
        app = self._307_sqlite_host(seeded=True)
        self._306_sidebar_shell(app)
        app._apply_style_settings({})
        normalized = app._normalize_style_settings({})
        self.assertEqual(app._normalize_style_settings(normalized), normalized)
        master = app.master_track
        records = app.backend.records
        metrics = []
        def run(label, action):
            executed = Counter()
            sql = Counter()
            connect = app.backend.adapter._connect
            def traced_connect():
                sql['connections'] += 1
                connection = connect()
                connection.set_trace_callback(lambda statement: sql.update(
                    ['selects'] if statement.lstrip().upper().startswith('SELECT') else []))
                return connection
            count_method, row_method = app._update_category_tab_visibility, app._populate_animal_list
            def count():
                if not app._ui_update_depth:
                    executed['counts'] += 1
                return count_method()
            def row(*args, **kwargs):
                executed['rows'] += 1
                return row_method(*args, **kwargs)
            with patch.object(app, '_update_category_tab_visibility', side_effect=count), \
                 patch.object(app, '_populate_animal_list', side_effect=row), \
                 patch.object(master, 'save_session', wraps=master.save_session) as sessions, \
                 patch.object(records, 'get', wraps=records.get) as gets, \
                 patch.object(records, 'get_with_revision', wraps=records.get_with_revision) as revisions, \
                 patch.object(app.backend.adapter, '_connect', side_effect=traced_connect), \
                 patch.object(app, '_plot_selected') as plots:
                start = perf_counter()
                action()
                elapsed = perf_counter() - start
            queries = Counter(tuple(call.args[:2]) for call in gets.call_args_list + revisions.call_args_list)
            metrics.append({'operation': label, 'seconds': round(elapsed, 6),
                            'executed': dict(executed), 'plot_calls': plots.call_count,
                            'session_writes': sessions.call_count,
                            'session_fields': [sorted(call.args[0]) if call.args else []
                                               for call in sessions.call_args_list],
                            'catalogue_reads': queries['projects', 'catalog'],
                            'user_reads': queries['security', 'users']})
            metrics[-1]['sql'] = dict(sql)
            self.assertEqual(executed, {'counts': 1, 'rows': 1})
            self.assertEqual(plots.call_count, 1)
            self.assertLessEqual(queries['projects', 'catalog'], 2)
            self.assertEqual(app.selected_animals, [])
            self.assertFalse(app.chk_events.signalsBlocked())
            print('ISSUE_306_SESSION_ROW ' + json.dumps(metrics[-1]), flush=True)
            return sessions.call_count
        for repetition in range(2):
            self.assertEqual(run('logout-' + str(repetition + 1), app._do_master_logout), 1)
            self.assertFalse(master.is_logged_in)
            with patch.object(master, 'login_interactive', side_effect=lambda **_kwargs:
                              (master._login_as('Manager') or True)):
                # First login normalizes the seed's older user style once;
                # this is an incoming preference write, not duplicate logout.
                self.assertEqual(run('login-' + str(repetition + 1), app._do_master_login),
                                 1 if repetition == 0 else 0)
            self.assertEqual(master.current_username, 'Manager')
        print('ISSUE_306_SESSION_METRICS ' + json.dumps(metrics), flush=True)

    def test_306_session_same_different_language_cancel_and_direct_switch(self):
        app = self._307_sqlite_host()
        self._306_sidebar_shell(app)
        app._apply_style_settings({})
        master = app.master_track
        for language in ('en', 'de'):
            with self.subTest(language=language):
                master.save_session({'language': language, 'display_chk_events': False})
                with patch.object(master, 'login_interactive', return_value=True) as login, \
                     patch.object(app, '_populate_animal_list', wraps=app._populate_animal_list) as rows, \
                     patch.object(app, '_load_messages', wraps=app._load_messages) as messages, \
                     patch.object(app, '_save_master_session', wraps=app._save_master_session) as outgoing:
                    app._do_master_login()
                self.assertEqual(app.lang, language)
                self.assertEqual(messages.call_count, 0 if language == 'en' else 1)
                self.assertEqual(rows.call_count, 1)
                outgoing.assert_called_once()
                login.assert_called_once_with(save_session=False)
                self.assertFalse(app.chk_events.isChecked())
        app.selected_animals = [next(iter(app.animals))]
        original = copy.deepcopy(app.selected_animals)
        with patch.object(master, 'login_interactive', return_value=False), \
             patch.object(app, '_populate_animal_list', wraps=app._populate_animal_list) as rows:
            app._do_master_login()
            rows.assert_not_called()
        self.assertEqual(app.selected_animals, original)
        self.assertEqual(master.current_username, 'tester')

    def test_306_medi_refresh_is_view_only_and_origin_scoped(self):
        from Plugins.Medi_Track.medi_track_widget import MediTrackPlugin
        app = self._307_sqlite_host()
        plugin = MediTrackPlugin(app)
        seen = []
        plugin._widget = SimpleNamespace(_current_animal='test-1', show_animal=seen.append)
        origin = QDialog(app)
        origin.setModal(True)
        origin.show()
        self.qt.processEvents()
        for _ in range(100):
            plugin._refresh_current_animal_when_safe('test-1')
        self.assertEqual(seen, [])
        self.assertEqual(len(app._after_dialog_updates._pending), 1)
        origin.accept()
        self.qt.processEvents()
        self.assertEqual(seen, ['test-1'])
        plugin._refresh_current_animal_when_safe('test-1')
        app._prepare_actor_transition()
        self.qt.processEvents()
        self.assertEqual(seen, ['test-1'])

    def test_306_conventions_resize_timer_dies_with_dialog(self):
        app = self._306_conventions()
        dlg = MAIN.StyleSettingsDialog(app, app.messages)
        timer = dlg._tab_fit_timer
        self.assertIs(timer.parent(), dlg)
        for _ in range(1000):
            dlg._schedule_current_tab_fit()
        self.assertTrue(timer.isActive())
        sip.delete(dlg)
        self.qt.processEvents()  # No callback touches a deleted dialog.
        self.assertTrue(sip.isdeleted(timer))

    def test_306_project_followup_and_core_history_are_durable_before_return(self):
        app = self._307_sqlite_host()
        app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app, MAIN.ProgTrackApp)
        app._schedule_post_animal_save_project_updates = (
            MAIN.ProgTrackApp._schedule_post_animal_save_project_updates.__get__(app, MAIN.ProgTrackApp))
        key = next(iter(app.animals))
        record = app.animals[key]
        MAIN.ensure_project_periods(record)
        record['project'] = ''
        origin = QDialog(app)
        origin.setModal(True)
        origin.show()
        self.qt.processEvents()
        app._schedule_post_animal_save_project_updates(key, 'Allowed', '')
        with patch.object(app.backend, 'save_core_data', wraps=app.backend.save_core_data) as saves:
            app._save_persistence(True)
        self.assertEqual(saves.call_count, 2, 'The second history commit is required, not duplicate work')
        stored = app.backend.load_core_data()['animals'][key]
        self.assertEqual(stored['project'], '')
        self.assertEqual(len(stored['project_history']), 1)
        self.assertEqual(stored['project_history'][0]['project'], 'Allowed')
        self.assertTrue(stored['project_history'][0]['leave_date'])
        self.assertIn('project_note', next(iter(stored['edits'].values())))
        log = Path(app.master_track.audit_logs_dir()) / datetime.now().strftime('audit_%Y-%m.log')
        self.assertIn('[tester] [data_edit]', log.read_text(encoding='utf-8'))
        self.assertNotIn('_pending_project_save_updates', app.__dict__)
        origin.accept()
        self.qt.processEvents()

    def test_306_coordinator_keeps_reentrant_newer_invalidation(self):
        app = self.host()
        seen = []
        def first():
            seen.append('first')
            app._queue_view_update('second', lambda: seen.append('newer'))
        app._queue_view_update('first', first)
        app._queue_view_update('second', lambda: seen.append('stale'))
        self.qt.processEvents()
        self.qt.processEvents()
        self.assertEqual(seen, ['first', 'newer'])

    def _317_relationship_records(self):
        records = {}
        for key, name, sex, public_id in (
            ('g-aegnor', 'Aegnor', 'Male', 'G-ID-NOT-A-NAME'),
            ('aerin', 'Aerin', 'Female', 'G-ID-NOT-A-NAME-2'),
            ('gimli-1', 'Gimli', 'Male', 'G-1'),
            ('gimli-2', 'Gimli', 'Male', 'G-2'),
            ('galadriel-1', 'Galadriel', 'Female', 'G-3'),
            ('galadriel-2', 'Galadriel', 'Female', 'G-4'),
            ('literal-m', 'G[one]', 'Male', 'L-1'),
            ('literal-f', 'G[one]', 'Female', 'L-2'),
            ('self', 'G-child', 'Male', 'SELF'),
            ('hidden', 'GandalfHidden', 'Male', 'DENIED'),
            ('mouse', 'G-mouse', 'Male', 'WRONG-SPECIES'),
            ('unknown', 'G-unknown', 'Unknown', 'WRONG-SEX')):
            record = self.record()
            record.update(name=name, sex=sex, id=public_id, project='Allowed')
            records[key] = record
        records['hidden']['project'] = 'Hidden'
        records['mouse']['species'] = 'Mus musculus'
        return records

    def _317_completions(self, combo, prefix):
        completer = combo.completer()
        completer.setCompletionPrefix(prefix)
        model = completer.completionModel()
        return [(model.index(row, 0).data(Qt.ItemDataRole.DisplayRole),
                 model.index(row, 0).data(Qt.ItemDataRole.UserRole))
                for row in range(completer.completionCount())]

    def test_317_both_paths_use_literal_case_insensitive_name_prefix_and_filtered_ids(self):
        app = self.host()
        app.animals = self._317_relationship_records()
        app._project_visibility_scope = lambda **_kwargs: (False, {'Allowed'})
        original = copy.deepcopy(app.animals)
        authorized = app._core_parentage_records()
        self.assertNotIn('hidden', authorized)
        group, parents = app._core_parentage_group('self', app.animals['self'])
        group.setParent(app)
        for sex in ('Female', 'Male'):
            shared = MAIN.build_relationship_combo(authorized, required_sex=sex,
                                                   species='Callithrix jacchus', exclude_id='self')
            shared.setParent(app)
            fields = ('egg_donor', 'surrogate_mother') if sex == 'Female' else (
                'sperm_donor', 'surrogate_father')
            for path, combo in [('shared', shared)] + [(key, parents[key]) for key in fields]:
                expected = {key for key, record in authorized.items() if key != 'self'
                            and record['sex'] == sex and record['species'] == 'Callithrix jacchus'}
                self.assertEqual({combo.itemData(i) for i in range(1, combo.count())}, expected)
                self.assertEqual(combo.insertPolicy(), MAIN.QComboBox.InsertPolicy.NoInsert)
                self.assertEqual(combo.completer().caseSensitivity(), Qt.CaseSensitivity.CaseInsensitive)
                self.assertEqual(combo.completer().filterMode(), Qt.MatchFlag.MatchStartsWith)
                for prefix in ('g', 'G', 'ae', 'AE', 'G[', '*', 'zz', ''):
                    with self.subTest(path=path, sex=sex, prefix=prefix):
                        actual = {identity for _label, identity in self._317_completions(combo, prefix)
                                  if identity}
                        self.assertEqual(actual, {key for key in expected
                            if authorized[key]['name'].casefold().startswith(prefix.casefold())})
                        self.assertEqual(combo.currentData(), '')
        self.assertEqual(app.animals, original)

    def test_317_duplicate_labels_select_exact_completion_identity(self):
        app = self.host()
        records = self._317_relationship_records()
        app.animals = records
        group, fields = app._core_parentage_group('self', records['self'])
        group.setParent(app)
        shared = MAIN.build_relationship_combo(records, required_sex='Female',
                                               species='Callithrix jacchus', exclude_id='self')
        shared.setParent(app)
        for path, combo in (('shared', shared), ('core', fields['egg_donor'])):
            with self.subTest(path=path):
                self.assertEqual(self._317_completions(combo, 'gal'),
                                 [('Galadriel (G-3)', 'galadriel-1'), ('Galadriel (G-4)', 'galadriel-2')])
                index = combo.completer().completionModel().index(1, 0)
                combo.completer().activated[MAIN.QtCore.QModelIndex].emit(index)
                self.assertEqual(combo.currentData(), 'galadriel-2')
                self.assertEqual(combo.text(), 'galadriel-2')
                self.assertEqual(combo.currentText(), 'Galadriel (G-4)')
                combo.setEditText('no match')
                self.assertFalse(self._317_completions(combo, 'no match'))
                self.assertEqual(combo.text(), '')
                self.assertEqual(app._core_parentage_values({'egg_donor': combo})['egg_donor'], 'no match')

    def test_317_committed_legacy_reference_and_strict_new_species_refresh_are_preserved(self):
        app = self.host()
        app.animals = self._317_relationship_records()
        source = copy.deepcopy(app.animals['self'])
        source['eizellspenderin'] = 'legacy-unavailable'
        group, fields = app._core_parentage_group('self', source)
        group.setParent(app)
        combo = fields['egg_donor']
        self.assertEqual(combo.currentData(), 'legacy-unavailable')
        self.assertIn(('legacy-unavailable', 'legacy-unavailable'), self._317_completions(combo, 'legacy'))
        shared = MAIN.build_relationship_combo(app.animals, current='legacy-unavailable', required_sex='Female')
        shared.setParent(app)
        self.assertEqual(shared.currentData(), 'legacy-unavailable')
        for selector in (shared, combo):
            selector.setEditText('unmatched')
            self.assertFalse(self._317_completions(selector, 'unmatched'))
            self.assertEqual(selector.currentData(), 'legacy-unavailable')
            self.assertEqual(selector.text(), '')
        species = MAIN.QComboBox(app)
        species.addItem('Callithrix jacchus', 'Callithrix jacchus')
        species.addItem('Mus musculus', 'Mus musculus')
        group, fields = app._core_parentage_group(None, source, species_widget=species, strict_species=True)
        group.setParent(app)
        self.assertEqual(fields['egg_donor'].currentData(), '')
        self.assertNotIn('legacy-unavailable', {identity for _, identity in self._317_completions(fields['egg_donor'], '')})
        species.setCurrentIndex(1)
        self.assertEqual(self._317_completions(fields['sperm_donor'], 'g'), [('G-mouse (WRONG-SPECIES)', 'mouse')])

    def test_317_typing_uses_existing_models_without_database_reads_or_writes(self):
        from PyQt6.QtTest import QTest
        app = self._307_sqlite_host(population=50)
        group, fields = app._core_parentage_group('test-1', app.animals['test-1'])
        group.setParent(app)
        shared = MAIN.build_relationship_combo(app._core_parentage_records(), required_sex='Female')
        shared.setParent(app)
        before = copy.deepcopy(app.animals)
        with patch.object(app, '_load_core_parentage_records', side_effect=AssertionError('typing reloaded directory')), \
             patch.object(app.backend.records, 'get', side_effect=AssertionError('typing read catalogue')), \
             patch.object(app.backend.records, 'get_with_revision', side_effect=AssertionError('typing read revision')), \
             patch.object(app.backend, 'save_core_data', side_effect=AssertionError('typing wrote data')):
            for combo in [shared, *fields.values()]:
                count = combo.count()
                for prefix in ('t', 'T', 'test-', 'zz', '', '*'):
                    combo.lineEdit().clear()
                    QTest.keyClicks(combo.lineEdit(), prefix)
                    self._317_completions(combo, prefix)
                    self.assertEqual(combo.count(), count)
        self.assertEqual(app.animals, before)

    def test_317_builtin_and_configured_editors_typing_cancel_reopen_preserve_parentage(self):
        from PyQt6.QtTest import QTest
        for method, role in (('_dlg_versuchstier', 'experimental'),
                             ('_dlg_basic_animal_role', 'custom.new_role')):
            with self.subTest(editor=method, role=role):
                app = self.host()
                app.animals = self._317_relationship_records()
                source = self.record(role)
                source.update(eizellspenderin='aerin', samenspender='g-aegnor')
                app.animals['test-1'] = source
                app._role_dialog_blocks = lambda *_args, **_kwargs: ['weight', 'parenting']
                before = copy.deepcopy(app.animals)
                dialog = self.open_dialog(app, method, name='test-1', role_value=role)
                combos = dialog.findChildren(MAIN.AnimalRelationshipCombo)
                self.assertEqual(len(combos), 4)
                combo = next(widget for widget in combos if widget.currentData() == 'g-aegnor')
                combo.lineEdit().selectAll()
                QTest.keyClicks(combo.lineEdit(), 'g')
                self.assertEqual(combo.currentText(), 'g')
                self.assertNotIn('g-aegnor', {key for _, key in self._317_completions(combo, 'g')})
                combo.lineEdit().selectAll()
                QTest.keyClicks(combo.lineEdit(), 'ae')
                self.assertIn('g-aegnor', {key for _, key in self._317_completions(combo, 'ae')})
                combo.setCurrentIndex(combo.findData('gimli-2'))
                self.assertEqual(combo.text(), 'gimli-2')
                combo.setEditText('no match')
                self.assertEqual(app._core_parentage_values({'sperm_donor': combo})['sperm_donor'], 'no match')
                dialog.reject()
                self.assertEqual(app.animals, before)
                reopened = self.open_dialog(app, method, name='test-1', role_value=role)
                values = {widget.currentData() for widget in reopened.findChildren(MAIN.AnimalRelationshipCombo)}
                self.assertTrue({'aerin', 'g-aegnor'}.issubset(values))
                self.assertEqual(app.animals, before)

    def test_307_parent_read_counts_do_not_scale_with_population(self):
        counts = []
        for population in (10, 695):
            with self.subTest(population=population):
                app = self._307_sqlite_host(population)
                original = app.backend.records
                with patch.object(original, 'get', wraps=original.get) as gets, \
                     patch.object(original, 'get_with_revision', wraps=original.get_with_revision) as revisions:
                    actual = app._core_parentage_records()
                queries = Counter(tuple(call.args[:2]) for call in gets.call_args_list + revisions.call_args_list)
                self.assertEqual(set(actual), {
                    f'test-{index + 1}' for index in range(population) if index % 5 != 2})
                self.assertEqual(queries['projects', 'catalog'], 1)
                self.assertLessEqual(queries['security', 'users'], 3)
                self.assertLessEqual(queries['security', 'organization-units'], 2)
                counts.append(queries)
        self.assertEqual(counts[0], counts[1])

    def test_307_project_and_unit_candidates_match_policy_for_all_actors(self):
        cases = (
            ('Manager', 'user', ['manager'], [], [], True),
            ('Master', 'master', [], [], [], True),
            ('Lord', 'lord', [], [], [], True),
            ('Researcher', 'user', ['researcher'], [], [], False),
            ('Direct grant', 'user', [], ['project.view_all'], [], True),
            ('Revoked grant', 'user', ['manager'], [], ['project.view_all'], False),
            ('Guest', 'guest', [], [], [], False),
        )
        app = self._307_sqlite_host(15)
        for label, role, jobs, granted, revoked, unrestricted in cases:
            with self.subTest(actor=label):
                app.backend.records.put('security', 'users', [{
                    'username': 'tester', 'role': role, 'jobs': jobs, 'unit_id': 'unit-a',
                    'permissions': {'granted': granted, 'revoked': revoked},
                }])
                app.master_track._current_username = None if role == 'guest' else 'tester'
                app.master_track._current_role = role
                # Deliberately stale cache must not authorize a new picker.
                with patch.object(app.master_track, 'get_project_visibility_cache',
                                  side_effect=AssertionError('opening must read fresh scope')):
                    actual = app._core_parentage_records()
                expected = {
                    f'test-{index + 1}' for index in range(15)
                    if (index % 5 != 2 and unrestricted)
                    or (index % 5 in (0, 4) and role != 'guest')
                }
                self.assertEqual(set(actual), expected)

    def test_307_scope_is_fresh_between_opens_and_never_used_at_save(self):
        app = self._307_sqlite_host()
        users = app.backend.records.get('security', 'users')
        users[0]['jobs'] = ['researcher']
        app.backend.records.put('security', 'users', users)
        self.assertIn('test-1', app._core_parentage_records())
        projects = app.backend.records.get('projects', 'catalog')
        projects['projects']['Allowed']['summary']['contact1_login'] = 'another-user'
        app.backend.records.put('projects', 'catalog', projects)
        app.master_track.set_project_visibility_cache(['Allowed'], dirty=False)
        self.assertEqual(app._core_parentage_records(), {})
        self.assertFalse(app._core_parent_candidate_visible(app.animals['test-1']))
        users[0]['jobs'] = ['manager']
        app.backend.records.put('security', 'users', users)
        self.assertIn('test-1', app._core_parentage_records())
        units = app.backend.records.get('security', 'organization-units')
        units['units'][0]['archived'] = True
        app.backend.records.put('security', 'organization-units', units)
        self.assertNotIn('test-1', app._core_parentage_records())
        self.assertFalse(app._core_parent_candidate_visible(app.animals['test-1']))
        users[0]['unit_id'] = 'unit-b'
        app.backend.records.put('security', 'users', users)
        self.assertIn('test-3', app._core_parentage_records())

    def test_307_visibility_errors_and_mid_read_actor_backend_changes_fail_closed(self):
        app = self._307_sqlite_host()
        with self.assertLogs(MAIN.__name__, level='WARNING'), \
             patch.object(app, '_project_visibility_scope', side_effect=OSError('catalog unavailable')):
            self.assertEqual(app._core_parentage_records(), {})
        with self.assertLogs(MAIN.__name__, level='WARNING'), \
             patch.object(app.authorization, 'can_read_unit', side_effect=OSError('Unit policy unavailable')) as denied:
            self.assertEqual(set(app._core_parentage_records()),
                             {'test-4', 'test-5', 'test-9', 'test-10'})
            self.assertEqual(denied.call_count, 2, 'one failed check per distinct owner, not per animal')
        authorization = app.authorization
        app.authorization = None  # Managed actor with no policy is NOT trusted-local.
        self.assertFalse(app._record_in_unit_scope(app.animals['test-1']))
        app.authorization = authorization
        for boundary in ('actor', 'backend', 'authorization', 'master-disabled'):
            with self.subTest(boundary=boundary):
                original = (app.backend, app.master_track._current_username, app.authorization)
                def change_context(_record):
                    if boundary == 'actor':
                        app.master_track._current_username = None
                    elif boundary == 'backend':
                        app.backend = object()
                    elif boundary == 'authorization':
                        app.authorization = object()
                    else:
                        app._disabled_plugins.add('master_track')
                    return True
                try:
                    with patch.object(app, '_record_in_unit_scope', side_effect=change_context):
                        self.assertEqual(app._core_parentage_records(), {})
                finally:
                    app.backend, app.master_track._current_username, app.authorization = original
                    app._disabled_plugins.discard('master_track')

    def test_307_dialog_revision_and_candidates_share_one_authoritative_snapshot(self):
        app = self._307_sqlite_host()
        dialog = QDialog(app)
        self.hosts.append(dialog)
        with patch.object(app, '_load_core_parentage_records', wraps=app._load_core_parentage_records) as load:
            app._set_animal_dialog_context(dialog, role='experimental', editing=True)
            expected = app._core_parentage_revision(snapshot_records=dialog._core_parentage_snapshot)
            group, _fields = app._core_parentage_group('test-child', self.record(), dialog=dialog)
            group.setParent(dialog)
        self.assertEqual(load.call_count, 1)
        self.assertEqual(dialog.property('core_parentage_revision'), expected)
        self.assertFalse(hasattr(dialog, '_core_parentage_snapshot'))
        # No opening snapshot may bypass authoritative command-time reads.
        with patch.object(app, '_load_core_parentage_records', wraps=app._load_core_parentage_records) as load:
            self.assertEqual(app._core_parentage_revision(), expected)
        self.assertEqual(load.call_count, 1)

    def test_307_owner_aliases_and_trusted_local_do_not_confuse_housing_with_workgroup(self):
        app = self._307_sqlite_host()
        variants = (
            ({'organization_unit_id': 'unit-a'}, True),
            ({'organization_unit_id': '', 'organizational_unit_id': 'unit-a'}, True),
            ({'workgroup_id': 'unit-a', 'unit_id': 'unit-b'}, True),
            ({'organization_unit_id': 'unit-b', 'workgroup_id': 'unit-a'}, False),
            ({'organization_unit_id': 'invalid unit'}, False),
            ({'unit_id': 'unit-b'}, True),
        )
        for owner_fields, expected in variants:
            with self.subTest(owner=owner_fields):
                record = self.record()
                record.update(owner_fields)
                self.assertEqual(app._record_in_unit_scope(record), expected)
        app._disabled_plugins.add('master_track')
        app.authorization.trusted_local = True
        self.assertEqual(len(app._core_parentage_records()), 10)

    def test_307_all_editor_paths_reuse_one_snapshot_and_cancel_without_writes(self):
        app = self._307_sqlite_host()
        app._role_dialog_blocks = lambda *_args, **_kwargs: ['weight', 'reference_weight']
        cases = list(self.weight_editor_cases()) + [('_dlg_female_animal', MAIN.Role.AMME.value)]
        for method, role in cases:
            for creating, read_only in ((False, False), (False, True), (True, False)):
                with self.subTest(editor=method, role=role, creating=creating, read_only=read_only):
                    app.animals['test-1']['rolle'] = role
                    MAIN.synchronize_experimental_limits(app.animals['test-1'], role)
                    snapshot = copy.deepcopy(app.animals)
                    captured = []
                    kwargs = {'read_only': read_only}
                    if method in ('_dlg_basic_animal_role', '_dlg_versuchstier'):
                        kwargs['role_value'] = role
                    elif method == '_dlg_female_animal':
                        kwargs['default_role'] = role
                    def reject_dialog(dialog):
                        captured.append(dialog)
                        self.assertIs(dialog.parent(), app)
                        self.assertFalse(dialog.isVisible(), 'construction must not show a transient window')
                        self.assertFalse(hasattr(dialog, '_core_parentage_snapshot'))
                        return QDialog.DialogCode.Rejected
                    with patch.object(app, '_load_core_parentage_records', wraps=app._load_core_parentage_records) as load, \
                         patch.object(QDialog, 'exec', reject_dialog), \
                         patch.object(app.backend, 'save_core_data', wraps=app.backend.save_core_data) as save:
                        getattr(app, method)(None if creating else 'test-1', **kwargs)
                    self.assertEqual(len(captured), 1)
                    self.assertEqual(load.call_count, 1)
                    self.assertEqual(app.animals, snapshot)
                    save.assert_not_called()

    def test_307_entity_lease_releases_on_cancel_exception_and_read_only_open(self):
        app = self._307_sqlite_host()
        leases = app.backend.leases
        with patch.object(leases, 'acquire', wraps=leases.acquire) as acquire, \
             patch.object(leases, 'release', wraps=leases.release) as release:
            modes = []
            app._open_animal_editor_with_lease('test-1', lambda _key, **kw: modes.append(kw), False)
            self.assertEqual(acquire.call_count, 1)
            self.assertEqual(release.call_count, 1)
            app._open_animal_editor_with_lease('test-1', lambda _key, **kw: modes.append(kw), True)
            self.assertEqual(acquire.call_count, 1)
            def fail(_key, **_kwargs):
                raise RuntimeError('construction failed')
            with self.assertRaises(RuntimeError):
                app._open_animal_editor_with_lease('test-1', fail, False)
            self.assertEqual(acquire.call_count, 2)
            self.assertEqual(release.call_count, 2)
            self.assertEqual(modes, [{'read_only': False}, {'read_only': True}])

    def test_307_fresh_save_denies_revoked_scope_and_parentage_changes(self):
        app = self._307_sqlite_host()
        source = copy.deepcopy(app.animals['test-1'])
        source['birth_date'] = '01.01.2020'
        source['organization_unit_id'] = ''
        app._master_can = lambda *_args: True  # Isolate candidate-scope validation.
        revision = app._core_parentage_revision()
        allowed = app._core_parentage_records()
        self.assertIn('test-5', allowed)
        users = app.backend.records.get('security', 'users')
        users[0]['jobs'] = ['researcher']
        app.backend.records.put('security', 'users', users)
        projects = app.backend.records.get('projects', 'catalog')
        projects['projects']['Allowed']['summary']['contact1_login'] = 'another-user'
        app.backend.records.put('projects', 'catalog', projects)
        app.master_track.set_project_visibility_cache(['Allowed'], dirty=False)
        with self.assertRaises(ValueError):
            app._apply_core_parentage('test-child', source, {'egg_donor': 'test-5'}, expected_revision=revision)
        users[0]['jobs'] = ['manager']
        app.backend.records.put('security', 'users', users)
        app._apply_core_parentage('test-child', source, {'egg_donor': 'test-5'}, expected_revision=revision)
        self.assertEqual(source['eizellspenderin'], 'test-5')
        changed = app.backend.load_core_data()
        changed['animals']['test-5']['death_date'] = '01.01.2026'
        app.backend.save_core_data(changed)
        with self.assertRaises(ValueError):
            app._apply_core_parentage('test-child', source, {'egg_donor': 'test-5'}, expected_revision=revision)

    def test_307_seeded_manager_first_and_repeat_open_read_counts_and_timings(self):
        app = self._307_sqlite_host(seeded=True)
        self.assertEqual(len(app.animals) + len(app.archived), 695)
        expected = None
        measurements = []
        records = app.backend.records
        tiffany = next(key for key, record in app.animals.items() if record['name'] == 'Tiffany')
        for operation in ('legacy-parent-loop', 'parent-first', 'parent-repeat', 'editor-first', 'editor-repeat'):
            captured = []
            sql = Counter()
            connect = app.backend.adapter._connect
            def traced_connect():
                sql['connections'] += 1
                connection = connect()
                connection.set_trace_callback(lambda query: sql.update(
                    ['selects'] if query.lstrip().upper().startswith('SELECT') else []))
                return connection
            with patch.object(records, 'get', wraps=records.get) as gets, \
                 patch.object(records, 'get_with_revision', wraps=records.get_with_revision) as revisions, \
                 patch.object(QDialog, 'exec', lambda dialog: captured.append(dialog) or QDialog.DialogCode.Rejected), \
                 patch.object(app, '_load_core_parentage_records', wraps=app._load_core_parentage_records) as load, \
                 patch.object(app.backend.adapter, '_connect', side_effect=traced_connect):
                start = perf_counter()
                if operation == 'legacy-parent-loop':
                    # Exact previous per-record visibility strategy, in-process
                    # only. No file reversion, live DB, modal interaction or Save.
                    result = {key: record for key, record in app._load_core_parentage_records().items()
                              if app._core_parent_candidate_visible(record)}
                elif operation.startswith('parent-'):
                    result = app._core_parentage_records()
                else:
                    app._open_animal_editor_with_lease(tiffany, app._dlg_female_animal, False)
                    self.assertEqual(len(captured), 1)
                elapsed = perf_counter() - start
            queries = Counter(tuple(call.args[:2]) for call in gets.call_args_list + revisions.call_args_list)
            measurements.append({'operation': operation, 'seconds': round(elapsed, 4),
                                 'project_reads': queries['projects', 'catalog'],
                                 'user_reads': queries['security', 'users'], 'core_snapshot_reads': load.call_count,
                                 'connections': sql['connections'], 'selects': sql['selects']})
            if operation == 'legacy-parent-loop':
                expected = set(result)
                self.assertEqual(queries['projects', 'catalog'], 695)
            elif operation.startswith('parent-'):
                self.assertEqual(set(result), expected)
                self.assertEqual(queries['projects', 'catalog'], 1)
                self.assertEqual(queries['security', 'users'], 1)
            else:
                self.assertLess(queries['security', 'users'], 100,
                                'editor permissions may read per control, never per candidate')
            self.assertEqual(load.call_count, 1)
        print('ISSUE_307_TIMINGS ' + json.dumps(measurements))

    def test_307_snapshot_is_rejected_if_context_changes_before_parent_group(self):
        app = self._307_sqlite_host()
        dialog = QDialog(app)
        self.hosts.append(dialog)
        app._set_animal_dialog_context(dialog, role='experimental', editing=True)
        app.master_track._current_username = None
        with self.assertRaises(ValueError):
            app._core_parentage_group('test-child', self.record(), dialog=dialog)
        self.assertFalse(hasattr(dialog, '_core_parentage_snapshot'))

    def test_307_context_change_never_publishes_retained_refs_or_modal(self):
        app = self._307_sqlite_host()
        source = self.record()
        source['eizellspenderin'] = 'old-user-private-parent'
        def change_actor(_record):
            app.master_track._current_username = None
            return True
        with patch.object(app, '_record_in_unit_scope', side_effect=change_actor):
            with self.assertRaises(ValueError):
                app._core_parentage_group('test-child', source)
        app.master_track._current_username = 'tester'
        dialog = QDialog(app)
        self.hosts.append(dialog)
        app._set_animal_dialog_context(dialog, role='experimental', editing=True)
        app.master_track._current_username = None
        with patch.object(dialog, 'exec') as execute:
            self.assertFalse(app._exec_animal_editor_dialog(dialog))
        execute.assert_not_called()
        self.assertFalse(hasattr(dialog, '_core_parentage_snapshot'))

    def open_dialog(self, app, method, name='Subject', **kwargs):
        captured = []
        with patch.object(QDialog, 'exec', lambda dlg: captured.append(dlg) or 0):
            getattr(app, method)(name, **kwargs)
        self.assertEqual(len(captured), 1)
        dlg = captured[0]
        self.hosts.append(dlg)
        return dlg

    def test_268_plain_decimal_validator_and_lossless_format(self):
        validator = MeasurementValidator()
        for value in (27.1, 350.1, 488.1, 0.125, 1e-8, 123456.123456789):
            for text in (format_measurement(value), format_measurement(value).replace('.', ',')):
                with self.subTest(text=text):
                    self.assertEqual(parse_measurement(text), value)
                    self.assertEqual(validator.validate(text, len(text))[0], QValidator.State.Acceptable)
        for text in ('1,234.5', '1.234,5', '1 234', 'NaN', 'inf', '1e5', '-1', '1,2,3'):
            with self.subTest(text=text):
                with self.assertRaises(ValueError):
                    parse_measurement(text)
                self.assertEqual(validator.validate(text, len(text))[0], QValidator.State.Invalid)
        self.assertEqual(parse_measurement('', default=100), 100)

    def weight_editor_cases(self):
        return (
            ('_dlg_partner', 'partner'), ('_dlg_samenspender', 'sperm_donor'),
            ('_dlg_offspring', 'offspring'), ('_dlg_zuchttier', 'breeding_animal'),
            ('_dlg_versuchstier', 'experimental'),
            ('_dlg_versuchstier', 'experimental_offspring'),
            ('_dlg_female_animal', 'donor'), ('_dlg_basic_animal_role', 'custom_role'),
        )

    def weight_page(self, app, dialog):
        label = app.messages['dialog.animal.tab.weights']
        return next(tab.widget(index) for tab in dialog.findChildren(QTabWidget)
                    for index in range(tab.count()) if tab.tabText(index) == label)

    def test_294_eight_editors_four_locales_use_singular_action_and_add_rows(self):
        expected = {'en': 'Add Weight', 'de': 'Gewicht hinzufügen',
                    'it': 'Aggiungi peso', 'ru': 'Добавить вес'}
        for language, label in expected.items():
            for method, role in self.weight_editor_cases():
                with self.subTest(language=language, editor=method, role=role):
                    app = self.host()
                    app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                    self.assertEqual(app.messages['button.add_weight'], label)
                    app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                    app.animals['Subject'] = self.record(role)
                    dialog = self.open_dialog(app, method)
                    page = self.weight_page(app, dialog)
                    buttons = [button for button in page.findChildren(QPushButton) if button.text() == label]
                    self.assertEqual(len(buttons), 1)
                    self.assertTrue(buttons[0].isEnabled())
                    fields_before = page.findChildren(QLineEdit)
                    original_texts = [field.text() for field in fields_before]
                    snapshot = copy.deepcopy(app.animals)
                    buttons[0].click()
                    self.assertEqual(len(page.findChildren(QLineEdit)), len(fields_before) + 2)
                    self.assertEqual([field.text() for field in fields_before], original_texts)
                    self.assertEqual(app.animals, snapshot, 'Add only edits the unsaved dialog')
                    self.assertEqual(app.errors, [])

    def test_294_denied_weight_actions_stay_disabled_in_all_editors(self):
        for method, role in self.weight_editor_cases():
            with self.subTest(editor=method, role=role):
                app = self.host()
                app._master_can = lambda permission, **_k: permission != 'core.edit_animal_measurements'
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                app.animals['Subject'] = self.record(role)
                dialog = self.open_dialog(app, method)
                page = self.weight_page(app, dialog)
                button = next(button for button in page.findChildren(QPushButton)
                              if button.text() == 'Add Weight')
                self.assertFalse(button.isEnabled())
                count = len(page.findChildren(QLineEdit))
                snapshot = copy.deepcopy(app.animals)
                button.click()
                self.assertEqual(len(page.findChildren(QLineEdit)), count)
                self.assertEqual(app.animals, snapshot)

    def test_294_generic_nonweight_lists_keep_their_existing_action(self):
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                app = self.host()
                app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                title = app.messages.get('dialog.female_animal.tab.measurements', 'Progesterone Values')
                scroll, rows = app._build_editable_list(
                    title, [], lambda row: row, lambda _rows: ('01.01.2026', '0', 'S1'),
                    col_headers=('Date', 'Value', 'Sample'))
                self.hosts.append(scroll)
                expected = app.messages['button.new_item'].format(title=title)
                button = next(button for button in scroll.findChildren(QPushButton)
                              if button.text() == expected)
                self.assertNotEqual(button.text(), app.messages['button.add_weight'])
                button.click()
                self.assertEqual(len(rows), 1)
                self.assertEqual([field.text() for field in rows[0]], ['01.01.2026', '0', 'S1'])

    def test_294_basic_creation_retains_initial_weight_field_and_manual_labels_match(self):
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                app = self.host()
                app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                dialog = self.open_dialog(app, '_dlg_basic_animal_role', name=None)
                initial_label = app.messages.get('dialog.field.initial_weight', 'Initial weight (g):')
                self.assertIn(initial_label, [label.text() for label in dialog.findChildren(MAIN.QLabel)])
                self.assertNotIn(app.messages['button.add_weight'],
                                 [button.text() for button in dialog.findChildren(QPushButton)])
                manual = (ROOT / f'manual/ProgTrack_User_Guide - {language}.html').read_text(encoding='utf-8')
                self.assertIn(f'<strong>{app.messages["button.add_weight"]}</strong>', manual)
                self.assertNotIn('dialog.partner.button.new_weight', app.messages)

    def weight_host(self, *, role='experimental', value=434.7, reference=440):
        app = self.host()
        app._role_block_enabled = lambda _role, block, _mode='edit': block == 'offspring_weight_evaluation'
        app._normalize_report_event_type = lambda value: str(value or '').casefold()
        app._custom_event_definitions = lambda **_k: []
        app._is_steroid_track_active = lambda: False
        app._get_status_at_date = lambda *_a, **_k: ''
        app._is_in_recovery_period = lambda *_a, **_k: False
        record = self.record(role)
        record['ref_weight'] = reference
        record['gewicht'] = [{'datum': datetime(2026, 1, 10), 'wert': value}]
        app.animals['Subject'] = record
        return app, record

    def test_293_raw_arithmetic_and_two_decimal_display(self):
        for current, reference, expected in (
                (434.7, 440, '-1.20'), (350.05, 350.1, '-0.01'),
                (350.1, 350.05, '+0.01'), (440, 440, '+0.00'),
                (0, 440, '-100.00'), (0.125, 0.124, '+0.81')):
            with self.subTest(current=current, reference=reference):
                app, record = self.weight_host(value=current, reference=reference)
                day = datetime(2026, 1, 10)
                before = copy.deepcopy(record)
                text, change = app._weight_display_values(record, day, current)
                self.assertEqual(text, f'{current:.2f}')
                self.assertEqual(change, (current - reference) / reference * 100)
                tooltip = app._format_tooltip('weight', day, current, 'Subject')
                self.assertIn(f'{current:.2f} g', tooltip)
                self.assertIn(expected + '%', tooltip)
                self.assertEqual(record, before, 'A tooltip must not normalize/persist animal data')
                daily = app._generate_daily_data('Subject', record, day.date())
                self.assertIn(f'{current:.2f}g', daily)
                self.assertIn(expected + '%', daily)

    def test_293_missing_zero_invalid_reference_and_missing_weight(self):
        for reference in (None, 0, -1, '', False, float('nan'), float('inf')):
            with self.subTest(reference=reference):
                app, record = self.weight_host(reference=reference)
                day = datetime(2026, 1, 10)
                tooltip = app._format_tooltip('weight', day, 434.7, 'Subject')
                daily = app._generate_daily_data('Subject', record, day.date())
                self.assertIn('434.70', tooltip)
                self.assertIn('434.70', daily)
                self.assertNotIn('%', tooltip + daily)
        record.pop('ref_weight')
        self.assertNotIn('%', app._format_tooltip('weight', day, 434.7, 'Subject'))
        self.assertNotIn('%', app._generate_daily_data('Subject', record, day.date()))
        self.assertIn('n/a', app._format_tooltip('weight', day, None, 'Subject'))

    def test_293_offspring_predecessor_not_future_same_day_or_list_order(self):
        for role in ('offspring', 'experimental_offspring'):
            with self.subTest(role=role):
                app, record = self.weight_host(role=role, value=350.1)
                record['gewicht'] += [
                    {'datum': datetime(2026, 2, 1), 'wert': 900},
                    {'datum': datetime(2026, 1, 9, 17), 'wert': 350.05},
                    {'datum': datetime(2026, 1, 9, 8), 'wert': 300},
                    {'datum': datetime(2026, 1, 10, 1), 'wert': 360},
                    {'datum': datetime(2026, 1, 8), 'wert': 250},
                    {'datum': 'bad-date', 'wert': 200}, None,
                ]
                day = datetime(2026, 1, 10)
                before = copy.deepcopy(record)
                tooltip = app._format_tooltip('weight', day, 350.1, 'Subject')
                self.assertEqual(record, before)
                self.assertIn('350.10 g (+0.01% vs. previous)', tooltip)
                index, _ = app._report_index_for_animal('Subject', record)
                before = copy.deepcopy(record)
                text, change = app._weight_display_values(record, day, 350.1, report_index=index)
                self.assertEqual(change, (350.1 - 350.05) / 350.05 * 100)
                daily = app._generate_daily_data('Subject', record, day.date(), report_index=index)
                self.assertIn('350.10g', daily)
                self.assertIn('+0.01%', daily)
                self.assertEqual(record, before)

    def test_293_offspring_first_zero_and_invalid_previous_do_not_fall_back(self):
        for previous in (None, 0, -1, 'invalid'):
            with self.subTest(previous=previous):
                app, record = self.weight_host(role='offspring')
                if previous is not None:
                    record['gewicht'].append({'datum': datetime(2026, 1, 9), 'wert': previous})
                tooltip = app._format_tooltip('weight', datetime(2026, 1, 10), 434.7, 'Subject')
                daily = app._generate_daily_data('Subject', record, datetime(2026, 1, 10).date())
                self.assertNotIn('%', tooltip + daily)
        app._role_block_enabled = lambda *_a, **_k: False
        self.assertIn('-1.20%', app._format_tooltip('weight', datetime(2026, 1, 10), 434.7, 'Subject'))

    def test_293_all_locales_keep_units_labels_and_precision(self):
        for language in ('en', 'de', 'ru', 'it'):
            with self.subTest(language=language):
                app, record = self.weight_host()
                app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                tooltip = app._format_tooltip('weight', datetime(2026, 1, 10), 434.7, 'Subject')
                daily = app._generate_daily_data('Subject', record, datetime(2026, 1, 10).date())
                self.assertIn(app.messages['plot.series.weight'], tooltip)
                self.assertIn(app.messages.get('daily.weight', 'Weight'), daily)
                self.assertIn('434.70', tooltip + daily)
                self.assertIn('-1.20%', tooltip)
                self.assertIn('-1.20%', daily)

    def test_293_sperm_tooltip_rounding_unchanged(self):
        app, _record = self.weight_host()
        tooltip = app._format_tooltip('sperm_motility', datetime(2026, 1, 10), 123.6,
                                     'Subject', {'mot_pct': 55.55})
        self.assertIn('56%', tooltip)
        self.assertIn('124 Sperm/ml', tooltip)

    def test_293_report_cache_refreshes_after_weight_or_reference_edit(self):
        app, record = self.weight_host()
        day = datetime(2026, 1, 10).date()
        first = app._report_month_daily_data('Subject', record, 2026, 1, app.messages)
        self.assertIn('-1.20%', first['10.01.2026'])
        record['gewicht'][0]['wert'] = 350.1
        record['ref_weight'] = 350.05
        second = app._report_month_daily_data('Subject', record, 2026, 1, app.messages)
        self.assertIn('350.10g', second['10.01.2026'])
        self.assertIn('+0.01%', second['10.01.2026'])
        self.assertIn('+0.01%', app._format_tooltip('weight', day, 350.1, 'Subject'))

    def test_293_real_pdf_xlsx_exports_match_reports_and_preserve_history(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        app, record = self.weight_host()
        day = datetime(2026, 1, 10).date()
        app.report_locked_dates, app.report_edits = set(), {}
        app._load_report_data = lambda *_a: None
        app._display_name = lambda name: name
        app._permission_scope_export_ids = lambda names, **_k: list(names)
        app._get_status_localized = lambda *_a: ''
        app._get_event_statistics_localized = lambda *_a: '-'
        app._format_id_with_species = lambda *_a, **_k: 'TEST-ID'
        app._format_report_title_subject = lambda *_a, **_k: 'Subject'
        app._format_project_severity = lambda *_a: ''
        app._get_localized_role = lambda role, *_a: role
        app._report_index_for_animal('Subject', record)
        before = copy.deepcopy(record)
        with tempfile.TemporaryDirectory() as directory:
            xlsx, pdf = Path(directory) / 'weights.xlsx', Path(directory) / 'weights.pdf'
            app._export_reports_xlsx(['Subject'], day, day, str(xlsx))
            workbook = load_workbook(xlsx, read_only=True)
            try:
                exported = workbook['Subject']['B6'].value
                self.assertIn('434.70g', exported)
                self.assertIn('-1.20%', exported)
            finally:
                workbook.close()
            app._create_single_pdf_report('Subject', record, 2026, 1, str(pdf), day, day, 'en')
            pdf_text = '\n'.join(page.extract_text() for page in PdfReader(str(pdf)).pages)
            self.assertIn('434.70g', pdf_text)
            self.assertIn('-1.20%', pdf_text)
        self.assertEqual(record, before)

    def test_268_history_reader_rejects_invalid_rows_atomically(self):
        app = self.host()
        rows = [(QLineEdit('01.01.2026'), QLineEdit('27,1'), None),
                (QLineEdit('02.01.2026'), QLineEdit('bad'), None)]
        with self.assertRaises(ValueError):
            app._read_dialog_measurements(rows)
        rows[1][1].setText('488.1')
        result = app._read_dialog_measurements(rows)
        self.assertEqual([row['wert'] for row in result], [27.1, 488.1])
        rows[1][0].setText('not-a-date')
        with self.assertRaises(ValueError):
            app._read_dialog_measurements(rows)

    def test_268_seven_editors_roundtrip_fractional_history(self):
        dialogs = {
            '_dlg_partner': 'partner', '_dlg_samenspender': 'sperm_donor',
            '_dlg_offspring': 'offspring', '_dlg_zuchttier': 'breeding_animal',
            '_dlg_versuchstier': 'experimental', '_dlg_basic_animal_role': 'custom_role',
            '_dlg_female_animal': 'donor',
        }
        for method, role in dialogs.items():
            with self.subTest(dialog=method):
                app = self.host()
                app.animals['Subject'] = self.record(role)
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                dlg = self.open_dialog(app, method)
                tabs = dlg.findChildren(QTabWidget)
                weights = next(tab.widget(index) for tab in tabs for index in range(tab.count())
                               if tab.tabText(index) == app.messages.get('dialog.animal.tab.weights', 'Weights'))
                self.assertIn('27.1', [field.text() for field in weights.findChildren(QLineEdit)])
                save = next(button for button in dlg.findChildren(QPushButton)
                            if button.text() == app.messages.get('button.save', 'Save'))
                callback_errors = []
                with patch('sys.excepthook', lambda typ, exc, tb: callback_errors.append(
                        ''.join(traceback.format_exception(typ, exc, tb)))):
                    save.click()
                self.assertEqual(callback_errors, [])
                self.assertEqual(app.errors, [])
                self.assertEqual(app.animals['Subject']['gewicht'][0]['wert'], 27.1)

    def test_265_basic_weight_event_block_combinations(self):
        for weight, event in ((True, True), (True, False), (False, True), (False, False)):
            with self.subTest(weight=weight, event=event):
                app = self.host()
                app.animals['Subject'] = self.record('custom_role')
                app._role_dialog_blocks = lambda *_a, **_k: ['weight'] if weight else []
                app._active_custom_event_definitions_for_role = lambda *_a, **_k: (
                    [{'event_type': 'custom', 'name': 'Custom', 'kind': 'counting'}] if event else [])
                dlg = self.open_dialog(app, '_dlg_basic_animal_role')
                tabs = dlg.findChildren(QTabWidget)
                labels = [tab.tabText(index) for tab in tabs for index in range(tab.count())]
                self.assertEqual('Weights' in labels, weight)
                self.assertEqual('Events' in labels, event)
                self.assertEqual(len(tabs), int(weight or event))
                self.assertNotIn('New weight (g):', [label.text() for label in dlg.findChildren(MAIN.QLabel)])

    def test_268_seven_editors_reject_invalid_rows_without_saving(self):
        for method, role in (('_dlg_partner', 'partner'), ('_dlg_samenspender', 'sperm_donor'),
                             ('_dlg_offspring', 'offspring'), ('_dlg_zuchttier', 'breeding_animal'),
                             ('_dlg_versuchstier', 'experimental'), ('_dlg_basic_animal_role', 'custom_role'),
                             ('_dlg_female_animal', 'donor')):
            with self.subTest(dialog=method):
                app = self.host()
                app.animals['Subject'] = self.record(role)
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                dlg = self.open_dialog(app, method)
                weights = next(tab.widget(index) for tab in dlg.findChildren(QTabWidget)
                               for index in range(tab.count()) if tab.tabText(index) == 'Weights')
                field = next(field for field in weights.findChildren(QLineEdit) if field.text() == '27.1')
                field.setText('1,2.3')
                before = copy.deepcopy(app.animals)
                writes = []
                app._save_persistence = lambda **_k: writes.append(True)
                callback_errors = []
                save = next(button for button in dlg.findChildren(QPushButton) if button.text() == 'Save')
                with patch('sys.excepthook', lambda typ, exc, tb: callback_errors.append(str(exc))):
                    save.click()
                self.assertEqual(callback_errors, [])
                self.assertTrue(app.errors)
                self.assertEqual(writes, [])
                self.assertEqual(app.animals, before)

    def test_268_sqlite_save_and_reopen_keep_fractional_measurements(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = runtime_paths(Path(directory))
            backend = ProgTrackBackend(paths, acquire_process_lock=False)
            try:
                record = self.record()
                record['gewicht'] += [{'datum': datetime(2026, 1, 2),
                                       'wert': parse_measurement('488,1')}]
                backend.save_core_data({'animals': {'Subject': record}, 'archived': {}})
                snapshot = backend.load_core_data()
                values = next(iter(snapshot['animals'].values()))['gewicht']
                self.assertEqual([item['wert'] for item in values], [27.1, 488.1])
                backend.save_core_data(snapshot)
            finally:
                backend.close()
            reopened = ProgTrackBackend(paths, acquire_process_lock=False)
            try:
                record = next(iter(reopened.load_core_data()['animals'].values()))
                self.assertEqual([item['wert'] for item in record['gewicht']], [27.1, 488.1])
            finally:
                reopened.close()

    def test_265_denied_custom_weights_remain_scrollable(self):
        app = self.host()
        app.animals['Subject'] = self.record('custom_role')
        app._role_dialog_blocks = lambda *_a, **_k: ['weight']
        app._master_can = lambda perm: perm != 'core.edit_animal_measurements'
        dlg = self.open_dialog(app, '_dlg_basic_animal_role')
        tab = dlg.findChild(QTabWidget).widget(0)
        scroll = tab.findChild(QScrollArea)
        self.assertTrue(scroll.isEnabled())
        self.assertTrue(scroll.verticalScrollBar().isEnabled())
        self.assertTrue(all(not field.isEnabled() for field in tab.findChildren(QLineEdit)))
        self.assertTrue(all(not button.isEnabled() for button in tab.findChildren(QPushButton)))

    def test_266_definition_kind_survives_save_reload(self):
        records = Records()
        backend = SimpleNamespace(records=records)
        registry = RoleBlockPresetRegistry(backend)
        definition = {'name': 'Count', 'event_type': 'count', 'kind': 'counting',
                      'render_mode': 'symbol', 'marker': 'D', 'color': '#123456'}
        registry.save_presets([], event_definitions=[definition])
        reloaded = RoleBlockPresetRegistry(backend).event_definitions()[0]
        self.assertEqual(reloaded['kind'], 'counting')
        self.assertEqual(reloaded['limit_block'], '')
        app = self.host()
        app.backend = backend
        animal = self.record('custom_role')
        animal['events'] = [{'datum': datetime(2026, 1, 1), 'event_type': 'count'}]
        self.assertEqual(app._custom_event_statistics(animal), ['Count: 1'])
        appearance = app._custom_event_appearance('count')
        self.assertEqual(appearance['marker'], 'D')
        self.assertEqual(appearance['render_mode'], 'symbol')

    def test_269_protected_selectors_style_only_disabled_state(self):
        app = self.host()
        denied, allowed = QComboBox(), QComboBox()
        app._master_can = lambda permission: permission == 'allowed'
        app._apply_dialog_field_permissions({'denied': denied, 'allowed': allowed})
        self.assertFalse(denied.isEnabled())
        self.assertTrue(allowed.isEnabled())
        self.assertIn('QComboBox:disabled', denied.styleSheet())
        saved, new = QComboBox(), QComboBox()
        for combo in (saved, new):
            combo.addItem('Count', 'count')
        MAIN.restore_event_combo_selection(saved, 'count', persisted=True)
        MAIN.restore_event_combo_selection(new, 'count')
        self.assertFalse(saved.isEnabled())
        self.assertTrue(new.isEnabled())
        self.assertIn('color: #666', saved.styleSheet())

    def test_267_hidden_tabs_have_no_geometry_or_hit_target(self):
        app = self.host()
        app.category_tab = QTabBar(app)
        app.category_tab.resize(600, 35)
        app._category_custom_role_values = []
        app._custom_category_roles = lambda: [{'value': 'custom_role', 'name': 'Empty role'}]
        app._animal_visible_to_current_user = lambda record: record.get('visible', True)
        app._sync_heritage_category_tab_state = lambda: None
        app.animals['Subject'] = self.record('experimental')
        app._rebuild_category_tabs()
        app.category_tab.show()
        self.qt.processEvents()
        self.assertFalse(app.category_tab.expanding())
        self.assertFalse(app.category_tab.isTabVisible(6))
        self.assertTrue(app.category_tab.tabRect(6).isEmpty())
        self.assertTrue(app.category_tab.isTabVisible(7))
        app.animals['Custom'] = self.record('custom_role')
        app._update_category_tab_visibility()
        self.assertTrue(app.category_tab.isTabVisible(6))
        app.animals['Custom']['visible'] = False
        app.category_tab.setCurrentIndex(6)
        app._update_category_tab_visibility()
        self.qt.processEvents()
        self.assertFalse(app.category_tab.isTabVisible(6))
        self.assertTrue(app.category_tab.tabRect(6).isEmpty())
        self.assertTrue(app.category_tab.isTabVisible(app.category_tab.currentIndex()))

    def test_267_role_setup_save_reflows_empty_tab_before_any_tab_click(self):
        self._assert_267_role_save_geometry()

    def test_267_all_selection_does_not_scroll_past_hidden_long_role(self):
        for locale in ('en', 'de', 'it', 'ru'):
            for width in (240, 300, 520):
                with self.subTest(locale=locale, sidebar_width=width):
                    self._assert_267_role_save_geometry(
                        start_all=True, locale=locale, sidebar_width=width)

    def test_267_conventions_accept_publishes_final_sidebar_without_click(self):
        self._assert_267_role_save_geometry(start_all=True, conventions=True)

    def test_267_empty_sidebar_and_custom_reorder_preserve_visible_identity(self):
        app = self.host()
        app.category_tab = QTabBar(app)
        app.category_tab.resize(300, 35)
        app._category_custom_role_values = []
        roles = [
            {'value': 'custom.first', 'name': 'First custom category with a long label'},
            {'value': 'custom.second', 'name': 'Second custom category with a long label'},
        ]
        app._custom_category_roles = lambda: roles
        app._animal_visible_to_current_user = lambda record: record.get('visible', True)
        app._sync_heritage_category_tab_state = lambda: None
        app._rebuild_category_tabs()
        app.show()
        app.category_tab.show()
        self.qt.processEvents()
        for _ in range(2):
            app._rebuild_category_tabs()
            self.assertEqual(app.category_tab.currentIndex(), app._all_category_tab_index())
            self.assertTrue(all(app.category_tab.tabRect(index).isEmpty()
                                for index in range(app._all_category_tab_index())))
        app.animals['Custom'] = self.record('custom.first')
        app._update_category_tab_visibility()
        app.category_tab.setCurrentIndex(6)
        roles.reverse()
        app._rebuild_category_tabs()
        self.assertEqual(app.category_tab.currentIndex(), 7)
        self.assertEqual(app._category_descriptor_for_index(7), 'role:custom.first')
        self.assertTrue(app.category_tab.isTabVisible(7))
        self.assertFalse(app.category_tab.isTabVisible(6))
        self.assertTrue(app.category_tab.tabRect(6).isEmpty())
        app.animals['Custom']['visible'] = False
        app._update_category_tab_visibility()
        self.assertEqual(app.category_tab.currentIndex(), app._all_category_tab_index())

    def _assert_267_role_save_geometry(
            self, *, start_all=False, conventions=False, locale='en', sidebar_width=300):
        app = self.host()
        app.messages = json.loads((ROOT / f'lang/messages_{locale}.json').read_text(encoding='utf-8'))

        class RoleRegistry:
            def __init__(self):
                builtin_values = (
                    MAIN.Role.SPENDER.value,
                    MAIN.Role.SAMENSP.value,
                    MAIN.Role.OFFSPRING.value,
                    MAIN.Role.PARTNER.value,
                    MAIN.Role.ZUCHTTIER.value,
                    MAIN.Role.EXPERIMENTAL.value,
                )
                self._roles = [
                    {
                        'role_id': value,
                        'value': value,
                        'label': value,
                        'icon': f'role.{value}',
                        'representation_mode': 'icon',
                        'active': True,
                        'built_in': True,
                    }
                    for value in builtin_values
                ]

            def roles(self):
                return copy.deepcopy(self._roles)

            def save_roles(self, roles):
                self._roles = copy.deepcopy(roles)

            def label_for_value(self, value, _messages):
                role = next((item for item in self._roles if item.get('value') == value), None)
                return role.get('label') if role else None

            def display_for_value(self, value, _messages):
                return self.label_for_value(value, _messages) or value

            def representation_for_value(self, value, _messages):
                role = next((item for item in self._roles if item.get('value') == value), {})
                if role.get('built_in'):
                    return {'mode': 'icon', 'icon': role.get('icon')}
                return {'mode': 'text', 'text': role.get('representation_text') or role.get('label') or value}

        app.animal_role_registry = RoleRegistry()
        if conventions:
            # Real configuration CAS/registry and real Conventions acceptance;
            # no live backend, user style file, identity settings or seed writes.
            from Plugins.core.animal_roles import AnimalRoleRegistry
            from Plugins.core.identity_conventions import DEFAULT_CONVENTIONS
            temporary = tempfile.TemporaryDirectory()
            self.addCleanup(temporary.cleanup)
            root = Path(temporary.name)
            app.backend = ProgTrackBackend(
                runtime_paths(root, application_root=root), acquire_process_lock=False)
            self.addCleanup(app.backend.close)
            app.animal_role_registry = AnimalRoleRegistry(app.backend)
            app._load_identity_conventions = lambda: copy.deepcopy(DEFAULT_CONVENTIONS)
            app._save_identity_conventions = lambda *_args: None
            app._can_configure_animal_roles = lambda: True
            app._can_manage_institution_settings = lambda: False
            app._can_manage_institution_branding = lambda: False
            app._load_user_style_settings = lambda: {}
            app._save_user_style_settings = lambda *_args: None
            app._plot_selected = lambda: None
        app._category_tab_tooltips = lambda: ['category'] * 7
        app._animal_visible_to_current_user = lambda record: record.get('visible', True)
        app._is_steroid_track_active = lambda: False
        app._save_role_label_overrides = lambda *_args: None
        app._merge_user_messages = lambda: None
        app._apply_sidebar_button_visibility_for_category = lambda *_args: None
        app._apply_master_button_states = lambda: None
        app._sync_plot_action_button_widths = lambda: None
        app._update_animal_sex_filter_texts = lambda: None

        sidebar = QWidget(app)
        sidebar.setFixedWidth(sidebar_width)
        sidebar_layout = QVBoxLayout(sidebar)
        app.setCentralWidget(sidebar)
        app.category_tab = QTabBar(sidebar)
        app._category_custom_role_values = []
        sidebar_layout.addWidget(app.category_tab)
        builtin_values = (
            MAIN.Role.SPENDER.value,
            MAIN.Role.SAMENSP.value,
            MAIN.Role.OFFSPRING.value,
            MAIN.Role.PARTNER.value,
            MAIN.Role.ZUCHTTIER.value,
            MAIN.Role.EXPERIMENTAL.value,
        )
        app.animals = {
            f'Subject {index}': self.record(role_value)
            for index, role_value in enumerate(builtin_values)
        }

        app.box_chk = QGroupBox(app)
        app.box_rad = QGroupBox(app)
        for name in ('chk_prog', 'chk_weight', 'chk_events', 'chk_mode_blood',
                     'rb_weight_on', 'rb_weight_off'):
            setattr(app, name, QCheckBox(app))
        for name in ('btn_new', 'btn_edit', 'btn_load_blood', 'btn_load_weights',
                     'btn_archive', 'btn_restore', 'btn_delete'):
            setattr(app, name, QPushButton(app))
        app.phase_buttons = {
            None: QPushButton(app),
            MAIN.Phase.FOLLIKEL.value: QPushButton(app),
            MAIN.Phase.LUTEAL.value: QPushButton(app),
        }
        app.phase_widget = QWidget(app)

        app.resize(720, 300)
        pixmap = QPixmap(24, 24)
        pixmap.fill(Qt.GlobalColor.black)
        fake_role_icon = QIcon(pixmap)
        with patch.object(MAIN, 'ui_icon', return_value=fake_role_icon):
            app._rebuild_category_tabs()
            app.show()
            self.qt.processEvents()
            app.category_tab.setCurrentIndex(app._all_category_tab_index() if start_all else 0)
        self.assertEqual(
            app.category_tab.currentIndex(),
            app._all_category_tab_index() if start_all else 0,
            msg=str([
                (name, record['rolle'], app._animal_role_value(record),
                 app._category_tab_index_for_role(app._animal_role_value(record)))
                for name, record in app.animals.items()
            ]),
        )

        def visible_geometry():
            return {
                app._category_descriptor_for_index(index): (
                    app.category_tab.tabRect(index).x(),
                    app.category_tab.tabRect(index).y(),
                    app.category_tab.tabRect(index).width(),
                    app.category_tab.tabRect(index).height(),
                )
                for index in range(app.category_tab.count())
                if app.category_tab.isTabVisible(index)
            }

        def sidebar_geometry():
            return (
                sidebar.geometry().getRect(),
                app.category_tab.geometry().getRect(),
                visible_geometry(),
            )

        before = sidebar_geometry()
        role = {
            'role_id': 'custom.test_role',
            'value': 'custom.test_role',
            'label_key': 'role.custom.test_role',
            'label': 'A long descriptive empty custom role',
            'representation_mode': 'text',
            'representation_text': 'A long descriptive empty custom role',
            'order': 100,
            'active': True,
            'built_in': False,
        }
        publications = []
        if conventions:
            # Exercise the actual refresh/list projection after the modal save,
            # retaining nontrivial canonical selection/order in the All view.
            app.cmb_arch = QComboBox(app)
            app._project_visibility_scope = lambda **_kwargs: (True, set())
            app._get_status = lambda _name: ''
            app._status_icon_ids = lambda *_args: []
            app._status_text_without_icon_tokens = lambda _status: ''
            app._can_use_primary_edit_button = lambda: True
            app._set_prog_event_plot_controls_visible = lambda *_args, **_kwargs: None
            app._set_sperm_controls_visible = lambda *_args: None
            app.btn_edit_animal = QPushButton(app)
            app.selected_animals = ['Subject 2', 'Subject 0']
            app._plot_selection_order = list(app.selected_animals)
            app._edit_selection_order = list(reversed(app.selected_animals))

            def refresh(**kwargs):
                MAIN.ProgTrackApp._refresh_list(app, **kwargs)
                publications.append(sidebar_geometry())

            app._refresh_list = refresh
            app.category_tab.currentChanged.connect(app._on_category_selected)
            app._refresh_list()
            publications.clear()

            def accept_conventions(dialog):
                self.hosts.append(dialog)
                dialog._add_role_table_row(role)
                dialog.accept()
                self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted)
                return dialog.result()

        with patch.object(MAIN, 'ui_icon', return_value=fake_role_icon):
            if conventions:
                try:
                    with patch.object(MAIN.StyleSettingsDialog, 'exec', accept_conventions):
                        app._show_style_settings()
                finally:
                    app.category_tab.currentChanged.disconnect(app._on_category_selected)
                self.assertTrue(publications, 'Conventions must refresh the actual sidebar')
                self.assertTrue(all(item == before for item in publications))
                self.assertEqual(app.errors, [])
                self.assertEqual(app.selected_animals, ['Subject 2', 'Subject 0'])
                self.assertEqual(app._edit_selection_order, ['Subject 0', 'Subject 2'])
                payload = app.backend.records.get('configuration', 'animal-roles')
                self.assertIn('custom.test_role', [item['value'] for item in payload['roles']])
            else:
                self.assertTrue(app._save_animal_role_definitions(
                    app.animal_role_registry.roles() + [role]))

        # This is the first state after Role Setup's save callback. Check it
        # before and after queued GUI work; do not click a category to force it.
        custom_index = app._category_custom_role_values.index('custom.test_role') + 6
        self.assertFalse(app.category_tab.isTabVisible(custom_index))
        self.assertTrue(app.category_tab.tabRect(custom_index).isEmpty())
        self.assertEqual(sidebar_geometry(), before)
        self.qt.processEvents()
        self.assertFalse(app.category_tab.isTabVisible(custom_index))
        self.assertTrue(app.category_tab.tabRect(custom_index).isEmpty())
        self.assertEqual(sidebar_geometry(), before)

        all_index = app._all_category_tab_index()
        self.assertEqual(
            app._category_descriptor_for_index(app.category_tab.currentIndex()),
            'all' if start_all else f'role:{MAIN.Role.SPENDER.value}',
        )
        other_visible = next(
            index for index in range(app.category_tab.count())
            if index != all_index and app.category_tab.isTabVisible(index)
        )
        app.category_tab.setCurrentIndex(other_visible)
        app.category_tab.setCurrentIndex(all_index)
        self.qt.processEvents()
        self.assertEqual(
            app._category_descriptor_for_index(app.category_tab.currentIndex()), 'all'
        )
        self.assertFalse(app.category_tab.isTabVisible(custom_index))
        self.assertTrue(app.category_tab.tabRect(custom_index).isEmpty())

        # A first accessible animal shows the role exactly once; losing
        # visibility hides it again without retaining a stale tab hit target.
        app.animals['Custom'] = self.record('custom.test_role')
        app._rebuild_category_tabs()
        self.qt.processEvents()
        custom_index = app._category_custom_role_values.index('custom.test_role') + 6
        self.assertTrue(app.category_tab.isTabVisible(custom_index))
        app.animals['Custom']['visible'] = False
        app._update_category_tab_visibility()
        self.qt.processEvents()
        self.assertFalse(app.category_tab.isTabVisible(custom_index))
        self.assertTrue(app.category_tab.tabRect(custom_index).isEmpty())
        app.animals['Custom']['visible'] = True
        app._update_category_tab_visibility()
        self.qt.processEvents()
        self.assertTrue(app.category_tab.isTabVisible(custom_index))
        del app.animals['Custom']
        app._update_category_tab_visibility()
        self.qt.processEvents()
        self.assertFalse(app.category_tab.isTabVisible(custom_index))
        self.assertTrue(app.category_tab.tabRect(custom_index).isEmpty())

    def plot_host(self, event_mode='symbol', measured=False, multiple=False):
        app = self.host()
        app._is_steroid_track_active = lambda: False
        app.current_figure = app.current_canvas = None
        app.detail_widget = QWidget(app)
        app.dlay = QVBoxLayout(app.detail_widget)
        app.dlay.addWidget(QWidget())
        app.scroll = QScrollArea(app)
        for name in ('chk_mode_blood', 'chk_prog', 'chk_weight', 'chk_events'):
            widget = QCheckBox(app)
            widget.setChecked(True)
            setattr(app, name, widget)
        for name in ('rb_blood_on', 'rb_blood_off', 'rb_weight_on', 'rb_weight_off',
                     'rb_sperm_on', 'rb_sperm_off'):
            setattr(app, name, QRadioButton(app))
        app._update_plot_display_control_visibility = lambda _availability: None
        app._apply_mode = lambda: None
        app._apply_plot_display_state = lambda: None
        app._apply_urine_scale = lambda: None
        app._get_current_events_checkbox = lambda: app.chk_events
        app._plot_x_viewport = None
        definition = {'name': 'Count', 'event_type': 'count', 'kind': 'counting',
                      'render_mode': event_mode, 'marker': 'D', 'color': '#123456'}
        app.backend.records.put('configuration', 'role-block-presets', {
            'event_definitions': [normalize_event_definition(definition)]})
        record = self.record('custom_role')
        record['events'] = [{'datum': datetime(2026, 1, 1), 'event_type': 'count'}]
        if not measured:
            record['gewicht'] = []
        app.animals = {'Subject': record}
        app.selected_animals = ['Subject']
        if multiple:
            app.animals['Second'] = copy.deepcopy(record)
            app.selected_animals.append('Second')
        return app

    def event_control_host(self, event_mode='symbol', measured=False, category=6):
        app = self.plot_host(event_mode, measured)
        for name in ('_update_plot_display_control_visibility',
                     '_apply_plot_display_state', '_get_current_events_checkbox'):
            delattr(app, name)  # Use the production hooks, not the old renderer stubs.
        app.category_tab = QTabBar(app)
        for index in range(8):
            app.category_tab.addTab(str(index))
        app.category_tab.setCurrentIndex(category)
        app._category_custom_role_values = ['custom_role']
        app._role_dialog_blocks = lambda role, *_a, **_k: ['weight', 'count'] if role == 'custom_role' else ['weight']
        for name in ('chk_ref_weight', 'chk_sperm_measurements',
                     'chk_events_offspring', 'chk_events_breeding', 'chk_events_experimental'):
            checkbox = QCheckBox(app)
            checkbox.setChecked(True)
            setattr(app, name, checkbox)
        for name in ('chk_events', 'chk_events_offspring', 'chk_events_breeding', 'chk_events_experimental'):
            checkbox = getattr(app, name)
            checkbox.setText(app.messages['checkbox.events'])
            checkbox.toggled.connect(app._apply_plot_display_state)
        app._update_toggle_controls = lambda: None
        app.phase_widget = QWidget(app)
        app.phase_buttons = {None: QRadioButton(app)}
        app.btn_edit = QPushButton(app)
        app._apply_sidebar_button_visibility_for_category = lambda _index: None
        app._apply_master_button_states = lambda: None
        app._refresh_list = lambda **_k: app._plot_selected()
        return app

    def test_295_real_custom_symbol_line_controls_toggle_without_changing_weights_or_data(self):
        for mode in ('symbol', 'line'):
            for measured in (False, True):
                for steroid in (False, True):
                    for category in (6, 7):  # Custom, then dynamic All.
                        with self.subTest(mode=mode, measured=measured, steroid=steroid, category=category):
                            app = self.event_control_host(mode, measured, category)
                            app._is_steroid_track_active = lambda: steroid
                            app._on_category_selected(category)
                            checkbox = app._get_current_events_checkbox()
                            self.assertIs(checkbox, app.chk_events)
                            self.assertFalse(checkbox.isHidden())
                            self.assertTrue(checkbox.isEnabled())
                            self.assertEqual(len(app.ev_lines), 1)
                            self.assertTrue(all(artist.get_visible() for artist in app.ev_lines + app.ev_texts))
                            snapshot = copy.deepcopy(app.animals)
                            weight_visibility = [artist.get_visible() for artist in app.weight_lines]
                            checkbox.click()
                            self.assertFalse(checkbox.isChecked())
                            self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))
                            self.assertEqual([artist.get_visible() for artist in app.weight_lines], weight_visibility)
                            app._plot_selected()  # Recalculation must keep Events OFF.
                            self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))
                            checkbox.click()
                            self.assertTrue(all(artist.get_visible() for artist in app.ev_lines + app.ev_texts))
                            self.assertEqual(app.animals, snapshot)

    def test_295_empty_undated_and_no_capability_states_keep_preferences(self):
        for configured, history in ((True, []), (True, [{'event_type': 'count'}]),
                                    (False, []), (False, [{'event_type': 'count'}])):
            with self.subTest(configured=configured, history=history):
                app = self.event_control_host()
                app.animals['Subject']['events'] = copy.deepcopy(history)
                if not configured:
                    app._role_dialog_blocks = lambda *_a, **_k: ['weight']
                app.chk_events.setChecked(False)
                snapshot = copy.deepcopy(app.animals)
                app._set_prog_event_plot_controls_visible(False, events_only=True)
                self.assertEqual(not app.chk_events.isHidden(), configured)
                self.assertFalse(app.chk_events.isEnabled())
                self.assertFalse(app.chk_events.isChecked())
                self.assertEqual(app.animals, snapshot)
                app.selected_animals = []
                app._update_plot_display_control_visibility()
                self.assertTrue(app.chk_events.isHidden())

    def test_295_retired_historical_and_mixed_all_events_remain_toggleable(self):
        app = self.event_control_host('line', measured=True, category=7)
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'][0]['active'] = False
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight']
        app.animals['Subject']['events'][0]['custom_event_snapshot'] = {
            'name': 'Retired Count', 'render_mode': 'line', 'color': '#123456'}
        app.animals['NoEvents'] = self.record('partner_animal')
        app.selected_animals.append('NoEvents')
        app._plot_selected()
        self.assertFalse(app.chk_events.isHidden())
        self.assertTrue(app.chk_events.isEnabled())
        self.assertEqual(len(app.ev_lines), 1)
        self.assertIn('Retired Count', app.ev_texts[0].get_text())
        app.chk_events.click()
        self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))
        app.selected_animals = ['NoEvents']
        app._plot_selected()
        self.assertTrue(app.chk_events.isHidden())
        self.assertEqual(app.ev_lines, [])
        app.selected_animals = ['Subject']
        app._plot_selected()
        self.assertFalse(app.chk_events.isHidden())
        self.assertFalse(app.chk_events.isChecked())
        self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))

    def test_295_builtin_category_uses_only_its_existing_preference_control(self):
        for category, role, key in (
                (0, 'egg_cell_donor', 'chk_events'),
                (2, 'offspring', 'chk_events_offspring'),
                (4, 'breeding_animal', 'chk_events_breeding'),
                (5, 'experimental_animal', 'chk_events_experimental'),
                (7, 'experimental_offspring', 'chk_events')):
            with self.subTest(category=category, role=role):
                app = self.event_control_host(category=category)
                app.animals['Subject']['rolle'] = role
                app._set_prog_event_plot_controls_visible(False, events_only=True)
                self.assertIs(app._get_current_events_checkbox(), getattr(app, key))
                for candidate in ('chk_events', 'chk_events_offspring',
                                  'chk_events_breeding', 'chk_events_experimental'):
                    control = getattr(app, candidate)
                    self.assertEqual(not control.isHidden(), candidate == key)
                    self.assertEqual(control.isEnabled(), candidate == key)
                    self.assertTrue(control.isChecked())

    def test_295_session_restore_keeps_custom_events_off_and_actor_reset_clears_artists(self):
        app = self.event_control_host()
        app.main_tabs = None
        app.category_tab = None  # Session restoration can precede tab reconstruction.
        app._restore_session_ui({'display_chk_events': False})
        app.category_tab = QTabBar(app)
        for index in range(8):
            app.category_tab.addTab(str(index))
        app.category_tab.setCurrentIndex(6)
        app._plot_selected()
        self.assertFalse(app.chk_events.isChecked())
        self.assertFalse(app.chk_events.isHidden())
        self.assertTrue(app.chk_events.isEnabled())
        self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))
        app._clear_session_animal_selection()
        app._plot_selected()
        self.assertEqual(app.selected_animals, [])
        self.assertEqual(app.ev_lines, [])
        self.assertEqual(app.hover_data, [])
        self.assertIsNone(app.current_canvas)
        self.assertTrue(app.chk_events.isHidden())
        # A new actor's restored preference is not the outgoing actor's OFF.
        app.category_tab = None
        app._restore_session_ui({'display_chk_events': True})
        self.assertTrue(app.chk_events.isChecked())

    def test_295_nonselected_denied_project_history_does_not_enable_events(self):
        app = self.event_control_host(category=7)
        denied = app.animals.pop('Subject')
        denied['project'] = 'DeniedProject'
        app.animals['Denied'] = denied
        allowed = self.record('custom_role')
        allowed.update(project='AllowedProject', events=[], gewicht=[])
        app.animals['Allowed'] = allowed
        app._project_visibility_scope = lambda **_kwargs: (False, {'AllowedProject'})
        self.assertFalse(app._animal_visible_to_current_user(denied))
        self.assertTrue(app._animal_visible_to_current_user(allowed))
        app.selected_animals = [name for name, record in app.animals.items()
                                if app._animal_visible_to_current_user(record)]
        snapshot = copy.deepcopy(denied)
        app._plot_selected()
        self.assertFalse(app.chk_events.isHidden())  # Role has a configured block.
        self.assertFalse(app.chk_events.isEnabled())  # Only denied animal has data.
        self.assertEqual(app.ev_lines, [])
        self.assertEqual(app.hover_data, [])
        self.assertEqual(denied, snapshot)

    def test_295_event_controls_relabel_in_all_locales_without_preference_changes(self):
        app = self.event_control_host()
        app.chk_events.setChecked(False)
        for language in ('en', 'de', 'it', 'ru'):
            with self.subTest(language=language):
                app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                app._update_plot_display_control_visibility()
                for key in ('chk_events', 'chk_events_offspring',
                            'chk_events_breeding', 'chk_events_experimental'):
                    self.assertEqual(getattr(app, key).text(), app.messages['checkbox.events'])
                self.assertFalse(app.chk_events.isChecked())

    def test_295_real_steroid_state_hook_preserves_custom_event_controls(self):
        app = self.event_control_host()
        app._update_category_tab_visibility = lambda: None
        app._on_select = lambda: app._plot_selected()
        app.chk_events.setChecked(False)
        for steroid in (True, False, True):
            with self.subTest(steroid=steroid):
                app._is_steroid_track_active = lambda: steroid
                app._apply_steroid_track_state()
                self.assertFalse(app.chk_events.isHidden())
                self.assertTrue(app.chk_events.isEnabled())
                self.assertFalse(app.chk_events.isChecked())
                self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines + app.ev_texts))

    def tooltip_host(self, kind='counting', mode='symbol', measured=False):
        app = self.event_control_host(mode, measured)
        definition = normalize_event_definition({
            'name': 'Custom reading', 'event_type': 'custom_reading', 'kind': kind,
            'render_mode': mode, 'marker': 'D', 'color': '#123456'})
        app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [definition]})
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom_reading']
        app.animals['Subject']['events'] = [{'event_id': 'E1', 'event_type': 'custom_reading',
            'datum': datetime(2026, 1, 1), 'recorded_role': 'custom_role'}]
        app.animals['Subject']['experimental_limits'] = {definition['block_id']: 3}
        return app

    def tooltip_mouse(self, app, event_name, artist=None, point=None):
        from matplotlib.backend_bases import MouseEvent
        app.current_canvas.draw()
        if point is None:
            if hasattr(artist, 'get_offsets'):
                point = artist.get_offset_transform().transform(artist.get_offsets()[0])
            else:
                point = artist.get_transform().transform((MAIN.mdates.date2num(datetime(2026, 1, 1)), 0.7))
        event = MouseEvent(event_name, app.current_canvas, *point,
                           button=1 if event_name == 'button_press_event' else None)
        app.current_canvas.callbacks.process(event_name, event)
        return event

    def assert_296_hover(self, app, artist, *, point=None):
        previous = app._annotation
        self.tooltip_mouse(app, 'motion_notify_event', artist, point)
        self.assertIs(app._annotation, previous, 'Hover cannot open/replace clicked text')
        highlight = app._highlight_scatter
        self.assertIsNotNone(highlight)
        if hasattr(artist, 'get_paths'):
            expected_path = artist.get_paths()[0]
            source_size = artist.get_sizes()[0]
            expected_point = artist.get_offset_transform().transform(artist.get_offsets()[0])
        else:
            marker = MarkerStyle(artist.get_marker())
            expected_path = marker.get_path().transformed(marker.get_transform())
            source_size = artist.get_markersize() ** 2
            expected_point = point
        np.testing.assert_allclose(highlight.get_paths()[0].vertices, expected_path.vertices)
        np.testing.assert_array_equal(highlight.get_paths()[0].codes, expected_path.codes)
        np.testing.assert_allclose(highlight.get_facecolors()[0][:3], to_rgba('yellow')[:3])
        self.assertGreater(highlight.get_sizes()[0], source_size)
        if expected_point is not None:
            actual_point = highlight.get_offset_transform().transform(highlight.get_offsets()[0])
            np.testing.assert_allclose(actual_point, expected_point)
        return highlight

    def test_296_real_symbol_hover_click_zoom_resize_and_no_record_mutation(self):
        for kind in ('counting', 'limiting'):
            for measured in (False, True):
                with self.subTest(kind=kind, measured=measured):
                    app = self.tooltip_host(kind, measured=measured)
                    app._plot_selected()
                    artist = app.ev_lines[0]
                    snapshot = copy.deepcopy(app.animals)
                    self.assert_296_hover(app, artist)
                    self.assertIsNone(app._annotation)
                    self.tooltip_mouse(app, 'button_press_event', artist)
                    text = app._annotation.get_text()
                    self.assertEqual(text, '01.01.2026 - Custom reading(1/3)' if kind == 'limiting'
                                     else '01.01.2026 - Custom reading(1)')
                    self.assertNotIn('?', text)
                    directory = os.environ.get('PROGTRACK_296_EVIDENCE_DIR')
                    if directory and kind == 'limiting' and not measured:
                        app.current_figure.savefig(Path(directory) / 'custom-click-compact.png')
                    axis = artist.axes
                    center = MAIN.mdates.date2num(datetime(2026, 1, 1))
                    axis.set_xlim(center - 1, center + 1)
                    app.current_figure.set_size_inches(8, 4)
                    axis.set_ylim(10, 1000)  # Symbol is at axes fraction, NOT numeric y=.02.
                    self.assert_296_hover(app, artist)
                    self.tooltip_mouse(app, 'button_press_event', artist)
                    self.assertIn('Custom reading', app._annotation.get_text())
                    expected = artist.get_offset_transform().transform(artist.get_offsets()[0])
                    actual = app._plot_event_display_point(artist, datetime(2026, 1, 1), MAIN.TRI_Y)
                    self.assertAlmostEqual(actual[0], expected[0])
                    self.assertAlmostEqual(actual[1], expected[1])
                    bounds = app._annotation.get_window_extent(app.current_canvas.get_renderer())
                    self.assertLessEqual(bounds.x1, app.current_figure.bbox.x1)
                    self.assertGreaterEqual(bounds.x0, app.current_figure.bbox.x0)
                    self.assertLessEqual(bounds.y1, app.current_figure.bbox.y1)
                    self.assertGreaterEqual(bounds.y0, app.current_figure.bbox.y0)
                    self.assertEqual(app.animals, snapshot)

    def test_296_coincident_symbols_include_both_counting_and_limiting_contexts(self):
        app = self.tooltip_host()
        definition = normalize_event_definition({'name': 'Limit', 'event_type': 'limit',
            'kind': 'limiting', 'render_mode': 'symbol', 'marker': 'D'})
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'].append(definition)
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom_reading', 'limit']
        app.animals['Subject']['experimental_limits'][definition['block_id']] = 3
        app.animals['Subject']['events'].append({'event_id': 'E2', 'event_type': 'limit',
            'datum': datetime(2026, 1, 1), 'recorded_role': 'custom_role'})
        app._plot_selected()
        self.assert_296_hover(app, app.ev_lines[0])
        self.assertIsNone(app._annotation)
        self.tooltip_mouse(app, 'button_press_event', app.ev_lines[0])
        self.assertIn('01.01.2026 - Custom reading(1)', app._annotation.get_text())
        self.assertIn('01.01.2026 - Limit(1/3)', app._annotation.get_text())

    def test_305_coincident_custom_symbols_use_first_style_red_outline_and_ordered_tooltip(self):
        app = self.tooltip_host()
        definitions = [
            normalize_event_definition({'name': 'Limit', 'event_type': 'limit',
                'kind': 'limiting', 'render_mode': 'symbol', 'marker': 's', 'color': '#abcdef'}),
            normalize_event_definition({'name': 'Third', 'event_type': 'third',
                'kind': 'counting', 'render_mode': 'symbol', 'marker': '^', 'color': '#654321'}),
        ]
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'].extend(definitions)
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: [
            'weight', 'custom_reading', 'limit', 'third'
        ]
        for definition in definitions:
            if definition['kind'] == 'limiting':
                app.animals['Subject']['experimental_limits'][definition['block_id']] = 3

        records = {
            'custom_reading': {'event_id': 'E1', 'event_type': 'custom_reading',
                'datum': datetime(2026, 1, 1),
                'recorded_role': 'custom_role'},
            'limit': {'event_id': 'E2', 'event_type': 'limit', 'datum': datetime(2026, 1, 1),
                'recorded_role': 'custom_role'},
            'third': {'event_id': 'E3', 'event_type': 'third', 'datum': datetime(2026, 1, 1),
                'recorded_role': 'custom_role'},
        }
        expected = {
            'custom_reading': ('D', '#123456'),
            'limit': ('s', '#abcdef'),
            'third': ('^', '#654321'),
        }
        for order in (
            ('custom_reading', 'limit'),
            ('limit', 'custom_reading'),
            ('custom_reading', 'limit', 'third'),
            ('third', 'limit', 'custom_reading'),
        ):
            with self.subTest(order=order):
                app.animals['Subject']['events'] = [copy.deepcopy(records[key]) for key in order]
                before = copy.deepcopy(app.animals)
                app._plot_selected()
                self.assertEqual(len(app.ev_lines), 1)
                artist = app.ev_lines[0]
                marker, fill = expected[order[0]]
                marker_style = MarkerStyle(marker)
                expected_path = marker_style.get_path().transformed(marker_style.get_transform())
                np.testing.assert_allclose(artist.get_paths()[0].vertices, expected_path.vertices)
                np.testing.assert_allclose(artist.get_facecolors()[0], to_rgba(fill))
                np.testing.assert_allclose(artist.get_edgecolors()[0], to_rgba('#ff0000'))
                self.assertGreater(float(artist.get_linewidths()[0]), 0)

                self.assert_296_hover(app, artist)
                self.assertIsNone(app._annotation)
                self.tooltip_mouse(app, 'button_press_event', artist)
                text = app._annotation.get_text()
                labels = {'custom_reading': 'Custom reading', 'limit': 'Limit', 'third': 'Third'}
                positions = [text.index(labels[event_type]) for event_type in order]
                self.assertEqual(positions, sorted(positions), msg=text)
                self.assertEqual(app.animals, before)

                app.chk_events.setChecked(False)
                self.assertFalse(artist.get_visible())
                self.tooltip_mouse(app, 'motion_notify_event', artist)
                self.assertIsNone(app._annotation)
                app.chk_events.setChecked(True)
                self.assertTrue(artist.get_visible())

    def test_305_builtin_fsh_and_symbol_configured_pgf_keep_combined_tooltip(self):
        app = self.tooltip_host()
        app.animals['Subject']['events'] = [
            {'event_id': 'FSH-1', 'event_type': 'fsh', 'datum': datetime(2026, 1, 1)},
            {'event_id': 'PGF-1', 'event_type': 'pgf', 'datum': datetime(2026, 1, 1)},
        ]
        base_appearance = app._custom_event_appearance

        def appearance(event_type):
            result = base_appearance(event_type)
            if event_type == 'pgf':
                result = dict(result)
                result.update(render_mode='symbol', marker='D', color='#123456')
            return result

        app._custom_event_appearance = appearance
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 1)
        artist = app.ev_lines[0]
        self.assertIsNone(artist._progtrack_event_context)
        np.testing.assert_allclose(artist.get_edgecolors()[0], to_rgba('#ff0000'))
        self.assert_296_hover(app, artist)
        self.assertIsNone(app._annotation)
        self.tooltip_mouse(app, 'button_press_event', artist)
        text = app._annotation.get_text()
        self.assertLess(text.index('FSH'), text.index('PGF'), msg=text)

    def test_305_single_symbol_control_remains_unoutlined(self):
        app = self.tooltip_host()
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 1)
        artist = app.ev_lines[0]
        self.assertFalse(hasattr(artist, '_progtrack_coincident_event_items'))
        np.testing.assert_allclose(artist.get_edgecolors()[0], artist.get_facecolors()[0])

    def test_305_does_not_group_different_dates_or_different_animals(self):
        app = self.tooltip_host()
        definition = normalize_event_definition({'name': 'Limit', 'event_type': 'limit',
            'kind': 'counting', 'render_mode': 'symbol', 'marker': 's', 'color': '#abcdef'})
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'].append(definition)
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom_reading', 'limit']
        app.animals['Subject']['events'] = [
            {'event_id': 'E1', 'event_type': 'custom_reading', 'datum': datetime(2026, 1, 1),
             'recorded_role': 'custom_role'},
            {'event_id': 'E2', 'event_type': 'limit', 'datum': datetime(2026, 1, 2),
             'recorded_role': 'custom_role'},
        ]
        second = copy.deepcopy(app.animals['Subject'])
        second['events'] = [{'event_id': 'E3', 'event_type': 'limit',
            'datum': datetime(2026, 1, 1), 'recorded_role': 'custom_role'}]
        app.animals['Other'] = second
        app.selected_animals.append('Other')
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 3)
        self.assertTrue(all(not hasattr(artist, '_progtrack_coincident_event_items')
                            for artist in app.ev_lines))

    def test_296_prefix_counts_use_date_period_role_and_stable_id_projection(self):
        app = self.tooltip_host('limiting')
        record = app.animals['Subject']
        record['events'] = [
            {'event_id': identity, 'event_type': 'custom_reading', 'datum': datetime(2026, 1, day),
             'recorded_role': role, 'project_period_id': period}
            for identity, day, role, period in (
                ('current2', 7, 'custom_role', 'new'), ('old', 1, 'custom_role', 'old'),
                ('other', 6, 'other_role', 'new'), ('current1', 5, 'custom_role', 'new'),
                ('future', 8, 'custom_role', 'new'))]
        record['events'].append(copy.deepcopy(record['events'][0]))
        snapshot = copy.deepcopy(record)
        contexts = app._plot_event_tooltip_contexts(record, MAIN.plot_event_entries(record),
                                                   scoped=True, active_period='new')
        by_id = {event['event_id']: contexts[id(event)] for event in record['events'][:-1]}
        self.assertEqual([by_id[key][1] for key in ('current1', 'current2', 'future')], [1, 2, 3])
        self.assertEqual(by_id['current2'][2], '3')
        for key in ('old', 'other'):
            self.assertEqual(by_id[key][1], 1)
            self.assertIsNone(by_id[key][2])
            self.assertTrue(by_id[key][3])
        self.assertEqual(record, snapshot)

    def test_296_retired_unknown_snapshot_tooltips_do_not_borrow_current_quotas(self):
        for state in ('retired', 'removed', 'detached'):
            with self.subTest(state=state):
                app = self.tooltip_host('limiting')
                event = app.animals['Subject']['events'][0]
                event['custom_event_snapshot'] = {'name': 'Old custom name', 'kind': 'limiting',
                                                  'render_mode': 'symbol', 'marker': 'D'}
                payload = app.backend.records.get('configuration', 'role-block-presets')
                if state == 'removed':
                    payload['event_definitions'] = []
                elif state == 'retired':
                    payload['event_definitions'][0]['active'] = False
                else:
                    app._role_dialog_blocks = lambda *_a, **_k: ['weight']
                app.backend.records.put('configuration', 'role-block-presets', payload)
                before = copy.deepcopy(app.animals['Subject'])
                app._plot_selected()
                self.assert_296_hover(app, app.ev_lines[0])
                self.assertIsNone(app._annotation)
                self.tooltip_mouse(app, 'button_press_event', app.ev_lines[0])
                self.assertIn('01.01.2026 - Old custom name(1)', app._annotation.get_text())
                self.assertIn('historical', app._annotation.get_text())
                self.assertNotIn('/3', app._annotation.get_text())
                self.assertNotIn('Recorded role', app._annotation.get_text())
                self.assertNotIn(event['recorded_role'], app._annotation.get_text())
                self.assertEqual(app.animals['Subject'], before)

    def test_296_hidden_clipped_symbols_and_hover_leave_do_not_expose_tooltips(self):
        app = self.tooltip_host()
        app._plot_selected()
        artist = app.ev_lines[0]
        self.assert_296_hover(app, artist)
        self.assertIsNone(app._annotation)
        self.tooltip_mouse(app, 'motion_notify_event', point=(-10, -10))
        self.assertIsNone(app._highlight_scatter)
        self.assertIsNone(app._annotation)
        self.tooltip_mouse(app, 'button_press_event', artist)
        app.chk_events.setChecked(False)
        self.assertIsNone(app._annotation)
        for event_name in ('motion_notify_event', 'button_press_event'):
            self.tooltip_mouse(app, event_name, artist)
            self.assertIsNone(app._annotation)
        app.chk_events.setChecked(True)
        artist.axes.set_xlim(MAIN.mdates.date2num(datetime(2026, 2, 1)),
                             MAIN.mdates.date2num(datetime(2026, 2, 2)))
        app.current_canvas.draw()
        self.assertIsNone(app._plot_event_display_point(artist, datetime(2026, 1, 1), MAIN.TRI_Y))

    def test_296_event_lines_never_spawn_hover_or_click_tooltips(self):
        for kind in ('counting', 'limiting'):
            for measured in (False, True):
                with self.subTest(kind=kind, measured=measured):
                    app = self.tooltip_host(kind, mode='line', measured=measured)
                    app._plot_selected()
                    line = app.ev_lines[0]
                    self.assertFalse(any(len(entry) == 6 and entry[5] is line for entry in app.hover_data))
                    self.assertIn('Custom reading (1/3)' if kind == 'limiting' else 'Custom reading (1)',
                                  app.ev_texts[0].get_text())
                    for event_name in ('motion_notify_event', 'button_press_event'):
                        self.tooltip_mouse(app, event_name, line)
                        self.assertIsNone(app._annotation)

    def _303_plot_session_host(self, backend):
        from Plugins.Master_Track.plugin import MasterTrackPlugin
        app = self.event_control_host('line', measured=False)
        app.backend = backend
        app.master_track = MasterTrackPlugin(app)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'birth', 'embryo_transfer', 'surgery']
        app.animals['Subject']['events'] = [
            {'event_id': f'303-{i}', 'event_type': typ,
             'datum': datetime(2026, 1, 1)+timedelta(days=i),
             'custom_event_snapshot': {'name': typ, 'color': '#000000', 'render_mode': 'line'}}
            for i, typ in enumerate(('birth', 'embryo_transfer', 'surgery'))]
        return app

    def _303_until_first_ui(self, app, username, disabled=False):
        """Execute real startup authentication ordering up to UI construction."""
        class FirstUI(Exception):
            pass
        app.has_master_track = True
        app._init_master_track_plugin = lambda: None
        app.plugin_manager = SimpleNamespace(validate_all=lambda **_k: {}, required_failures=lambda: [])
        app._load_disabled_plugins = lambda: {'master_track'} if disabled else set()
        app._init_all_plugins_deferred = lambda *_a, **_k: None
        def authenticate():
            app.master_track._current_username = username
            app.master_track._current_role = 'user' if username else 'guest'
        app.master_track.startup = authenticate
        app.master_track._current_username = None
        app._load_settings()  # Exactly the real pre-authentication initial load.
        with patch.object(MAIN.QTimer, 'singleShot'), \
             patch.object(MAIN, 'QMenuBar', side_effect=FirstUI):
            with self.assertRaises(FirstUI):
                app._init_application_state()

    def _303_artist_colors(self, app):
        if not app.selected_animals:
            app.selected_animals = ['Subject']  # Select after legitimate Guest clearing.
        app._plot_selected()
        app.current_canvas.draw()
        return [tuple(line.get_facecolors()[0]) if hasattr(line, 'get_facecolors')
                else to_rgba(line.get_color()) for line in app.ev_lines]

    def test_303_persisted_sqlite_reconstructed_widget_first_render_uses_authenticated_styles(self):
        from Plugins.Master_Track.session import SessionManager
        with tempfile.TemporaryDirectory() as temporary:
            paths = runtime_paths(Path(temporary), application_root=Path(temporary))
            colors = {'birth': '#123456', 'embryo_transfer': '#654321', 'surgery': '#aa22bb'}
            backend = ProgTrackBackend(paths, acquire_process_lock=False)
            writer = self._303_plot_session_host(backend)
            writer.master_track._current_username = 'style-a'
            snapshot = writer._normalize_style_settings({'custom_events': {
                key: {'color': value} for key, value in colors.items()}})
            writer._save_user_style_settings(snapshot)
            saved = SessionManager('', backend).load('style-a')
            backend.close()
            reopened = ProgTrackBackend(paths, acquire_process_lock=False)
            try:
                app = self._303_plot_session_host(reopened)
                animals = copy.deepcopy(app.animals)
                self._303_until_first_ui(app, 'style-a')
                self.assertEqual(self._303_artist_colors(app), [to_rgba(value) for value in colors.values()])
                self.assertEqual(SessionManager('', reopened).load('style-a'), saved)
                self.assertEqual(app.animals, animals)
            finally:
                reopened.close()

    def test_303_persisted_symbol_color_and_marker_override_win_over_event_snapshot(self):
        app = self._307_sqlite_host()
        renderer = self._303_plot_session_host(app.backend)
        renderer.animals['Subject']['events'][2]['event_type'] = 'fsh'
        renderer.master_track._current_username = 'style-a'
        renderer._save_user_style_settings({'custom_events': {
            'birth': {'color': '#123456'}, 'embryo_transfer': {'color': '#654321'},
            'fsh': {'color': '#008833', 'marker': 'D'}}})
        self._303_until_first_ui(renderer, 'style-a')
        self.assertEqual(self._303_artist_colors(renderer),
                         [to_rgba('#123456'), to_rgba('#654321'), to_rgba('#008833')])
        symbol = renderer.ev_lines[2]
        marker = MarkerStyle('D')
        np.testing.assert_allclose(symbol.get_paths()[0].vertices,
                                   marker.get_path().transformed(marker.get_transform()).vertices)
        directory = os.environ.get('PROGTRACK_303_EVIDENCE_DIR')
        if directory:
            day = MAIN.mdates.date2num(datetime(2026, 1, 1))
            renderer._apply_shared_plot_viewport((day-2, day+4))
            renderer.current_figure.savefig(Path(directory) / 'first-render-user-colours.png')

    def test_303_cold_guest_disabled_and_standalone_do_not_borrow_user_styles(self):
        with tempfile.TemporaryDirectory() as temporary:
            backend = ProgTrackBackend(runtime_paths(Path(temporary), application_root=Path(temporary)),
                                       acquire_process_lock=False)
            try:
                for disabled in (False, True):
                    with self.subTest(disabled=disabled):
                        app = self._303_plot_session_host(backend)
                        app._apply_style_settings({'custom_events': {'birth': {'color': '#123456'}}})
                        self._303_until_first_ui(app, None, disabled)
                        self.assertEqual(self._303_artist_colors(app),
                                         [to_rgba('#ead700'), to_rgba('#2feade'), to_rgba('#0000ff')])
                app = self._303_plot_session_host(backend)
                app.master_track = None
                app._save_user_style_settings({'custom_events': {'birth': {'color': '#ff00ff'}}})
                app._load_settings()
                self.assertEqual(self._303_artist_colors(app)[0], to_rgba('#ff00ff'))
            finally:
                backend.close()

    def test_303_direct_login_applies_style_before_plugin_and_session_callbacks(self):
        app = self.event_control_host('line', measured=False)
        outgoing = {'custom_events': {'birth': {'color': '#111111'}}}
        incoming = {'custom_events': {'birth': {'color': '#abcdef'}}}
        app._apply_style_settings(outgoing)
        app.master_track.is_logged_in = True
        app.master_track.load_session = lambda: {'style_settings': incoming}
        app.master_track.save_session = lambda *_a, **_k: None
        app.master_track.login_interactive = lambda **_k: True
        received = []
        def inspect(*_a, **_k):
            received.append(app._custom_event_appearance('birth')['color'])
        for method in ('_apply_all_plugin_states', '_restore_session_ui', '_setup_menus',
                       '_refresh_list', '_on_select', '_refresh_master_menu_states',
                       '_update_master_status_bar', '_apply_master_button_states'):
            setattr(app, method, inspect)
        app._save_master_session = lambda: None
        app._prepare_reports_for_user_context_change = lambda: None
        app._load_disabled_plugins = lambda: set()
        app._do_master_login()
        self.assertTrue(received)
        self.assertTrue(all(color == '#abcdef' for color in received), received)

    def test_303_real_sessions_switch_cancel_and_logout_isolate_first_render(self):
        app = self._307_sqlite_host()
        renderer = self._303_plot_session_host(app.backend)
        master = renderer.master_track
        expected = {
            'style-a': ['#123456', '#654321', '#aa22bb'],
            'style-b': ['#abcdef', '#fedcba', '#224466'],
        }
        for username, colors in expected.items():
            master._current_username = username
            renderer._save_user_style_settings({'custom_events': {
                key: {'color': color} for key, color in zip(('birth', 'embryo_transfer', 'surgery'), colors)}})
        sessions = {key: master.session_mgr.load(key) for key in expected}
        renderer._save_master_session = lambda: None
        renderer._prepare_reports_for_user_context_change = lambda: None
        renderer._load_disabled_plugins = lambda: set()
        renderer._get_global_language_fallback = lambda: renderer.lang
        for method in ('_apply_all_plugin_states', '_setup_menus', '_refresh_list',
                       '_on_select', '_refresh_master_menu_states', '_update_master_status_bar',
                       '_apply_master_button_states'):
            setattr(renderer, method, lambda *_a, **_k: None)
        frames = []
        renderer._restore_session_ui = lambda *_a: frames.append([
            to_rgba(renderer._custom_event_appearance(typ)['color'])
            for typ in ('birth', 'embryo_transfer', 'surgery')])
        for username in ('style-a', 'style-b', 'style-a'):
            with self.subTest(user=username):
                def login(**_k):
                    master._current_username = username
                    return True
                with patch.object(master, 'login_interactive', side_effect=login):
                    renderer._do_master_login()
                self.assertEqual(frames[-1], [to_rgba(value) for value in expected[username]])
                # Real rendering follows the existing batched publication boundary.
                self.assertEqual(self._303_artist_colors(renderer), frames[-1])
        style = copy.deepcopy(renderer.custom_event_styles)
        count = len(frames)
        with patch.object(master, 'login_interactive', return_value=False):
            renderer._do_master_login()
        self.assertEqual(renderer.custom_event_styles, style)
        self.assertEqual(len(frames), count)
        # Real session ended, but suppress external audit/idle services in this fixture.
        with patch.object(master, 'logout', side_effect=lambda **_k: master._set_guest()):
            renderer._do_master_logout()
        self.assertEqual(self._303_artist_colors(renderer),
                         [to_rgba('#ead700'), to_rgba('#2feade'), to_rgba('#0000ff')])
        for username, saved in sessions.items():
            self.assertEqual(master.session_mgr.load(username), saved)

    def test_303_actual_conventions_reset_persists_only_current_actor_defaults(self):
        app = self._306_conventions()
        app._apply_style_settings({'custom_events': {
            'birth': {'color': '#123456'}, 'embryo_transfer': {'color': '#654321'}}})
        app._save_user_style_settings({'custom_events': {
            'birth': {'color': '#123456'}, 'embryo_transfer': {'color': '#654321'}}})
        app.master_track.session_mgr.save('style-b', {'style_settings': {
            'custom_events': {'birth': {'color': '#abcdef'}}}})
        other = app.master_track.session_mgr.load('style-b')
        data = app.backend.load_core_data()
        dialog = MAIN.StyleSettingsDialog(app, app.messages)
        self.hosts.append(dialog)
        self.assertEqual(dialog._custom_event_style_controls['birth'][1].property('color'), '#123456')
        dialog._reset_to_defaults()
        settings = dialog.get_settings()
        self.assertNotIn('birth', settings['custom_events'])
        self.assertNotIn('embryo_transfer', settings['custom_events'])
        dialog.accept()
        app._apply_accepted_conventions(dialog)
        app._apply_style_settings({})  # Simulate fresh in-memory defaults before reload.
        app._reload_user_style_settings()
        self.assertEqual(app._custom_event_appearance('birth')['color'], '#ead700')
        self.assertEqual(app._custom_event_appearance('embryo_transfer')['color'], '#2feade')
        self.assertEqual(app.master_track.session_mgr.load('style-b'), other)
        self.assertEqual(app.backend.load_core_data(), data)

    def line_label_host(self, count=12, measured=False, same_date=False):
        app = self.tooltip_host(mode='line', measured=measured)
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'].append(normalize_event_definition({
            'name': 'Second limiting event', 'event_type': 'second_limit',
            'kind': 'limiting', 'render_mode': 'line', 'color': '#ab3456'}))
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom_reading', 'second_limit']
        app.animals['Subject']['events'] = [
            {'event_id': f'line-{i}', 'event_type': 'custom_reading' if i % 2 else 'second_limit',
             'datum': datetime(2026, 1, 1) + timedelta(days=0 if same_date else (count-i)//3),
             'recorded_role': 'custom_role'} for i in range(count)]
        return app

    def assert_311_slots(self, app):
        app.current_canvas.draw()
        layout = app.current_figure._progtrack_event_label_layout
        renderer = app.current_canvas.get_renderer()
        groups = {}
        for order, (label, day) in enumerate(layout.entries):
            if label.get_visible():
                groups.setdefault(label.axes, []).append((day, order, label))
        for axis, entries in groups.items():
            previous = None
            for day, _order, label in sorted(entries, key=lambda item: item[:2]):
                bounds = label.get_window_extent(renderer)
                anchor = label.get_transform().transform((day, label.get_position()[1]))
                self.assertGreaterEqual(bounds.x0 + 1e-6, anchor[0] + renderer.points_to_pixels(3))
                if previous is not None:
                    self.assertGreaterEqual(bounds.x0 + 1e-6, previous + renderer.points_to_pixels(3))
                previous = bounds.x1
                self.assertEqual(label.get_fontsize(), 8)
                self.assertEqual(label.get_rotation(), 90)
                self.assertTrue(label.get_clip_on())
                self.assertFalse(label.get_wrap())
            end = axis.transData.transform((layout.right_extent_by_axis[axis], 0))[0]
            self.assertGreaterEqual(end + 1e-6, previous + renderer.points_to_pixels(3))
        return layout

    def test_311_two_three_and_dense_same_near_date_labels_preserve_lines_and_data(self):
        for count in (2, 3, 12, 60):
            for same_date in (False, True):
                for measured in (False, True):
                    with self.subTest(count=count, same_date=same_date, measured=measured):
                        app = self.line_label_host(count, measured, same_date)
                        snapshot = copy.deepcopy(app.animals)
                        app._plot_selected()
                        lines = [(tuple(line.get_xdata()), line.get_color()) for line in app.ev_lines]
                        labels = [(label.get_text(), label.get_position()[1]) for label in app.ev_texts]
                        self.assertEqual(len(lines), count)
                        self.assertEqual(len(labels), count)
                        self.assert_311_slots(app)
                        self.assertEqual(lines, [(tuple(line.get_xdata()), line.get_color()) for line in app.ev_lines])
                        self.assertEqual(labels, [(label.get_text(), label.get_position()[1]) for label in app.ev_texts])
                        self.assertEqual(app.animals, snapshot)

    def test_311_right_edge_overflow_is_reachable_by_existing_pan_without_resetting_slots(self):
        app = self.line_label_host(60, same_date=True)
        app._plot_selected()
        day = MAIN.mdates.date2num(datetime(2026, 1, 1))
        app._apply_shared_plot_viewport((day-30, day))
        layout = self.assert_311_slots(app)
        axis = app.ev_texts[0].axes
        self.assertGreater(layout.right_extent_by_axis[axis], axis.get_xlim()[1])
        last = app.ev_texts[-1]
        box = last.get_window_extent(app.current_canvas.get_renderer())
        self.assertGreater(box.x0, axis.bbox.x1)
        # Same shared-viewport operation used by the existing drag gesture.
        target = axis.transData.inverted().transform(((box.x0+box.x1)/2, box.y0))[0]
        app._apply_shared_plot_viewport((target-15, target+15))
        self.assert_311_slots(app)
        box = last.get_window_extent(app.current_canvas.get_renderer())
        self.assertGreaterEqual(box.x0, axis.bbox.x0)
        self.assertLessEqual(box.x1, axis.bbox.x1)
        self.assertLess(day, axis.get_xlim()[0])
        directory = os.environ.get('PROGTRACK_311_EVIDENCE_DIR')
        if directory:
            app.current_figure.savefig(Path(directory) / 'rightward-overflow-panned.png')

    def test_311_existing_left_drag_reaches_overflow_without_new_navigation(self):
        from matplotlib.backend_bases import MouseEvent
        app = self.line_label_host(60, same_date=True)
        app._plot_selected()
        day = MAIN.mdates.date2num(datetime(2026, 1, 1))
        app._apply_shared_plot_viewport((day-30, day))
        self.assert_311_slots(app)
        axis = app.ev_texts[0].axes
        last = app.ev_texts[-1]
        for _ in range(10):
            box = last.get_window_extent(app.current_canvas.get_renderer())
            if axis.bbox.x0 <= box.x0 and box.x1 <= axis.bbox.x1:
                break
            y = axis.bbox.y0 + 5
            start = axis.bbox.x0 + .8 * axis.bbox.width
            end = axis.bbox.x0 + .2 * axis.bbox.width
            for name, x in (('button_press_event', start), ('motion_notify_event', end),
                            ('button_release_event', end)):
                event = MouseEvent(name, app.current_canvas, x, y, button=1)
                app.current_canvas.callbacks.process(name, event)
            self.assert_311_slots(app)
        box = last.get_window_extent(app.current_canvas.get_renderer())
        self.assertGreaterEqual(box.x0, axis.bbox.x0)
        self.assertLessEqual(box.x1, axis.bbox.x1)
        self.assertIsNone(app._annotation)
        self.assertEqual(len(app.ev_texts), 60)

    def test_311_first_draw_zoom_resize_dpi_and_changed_text_reflow_without_drift(self):
        app = self.line_label_host(12)
        app._plot_selected()
        day = MAIN.mdates.date2num(datetime(2026, 1, 1))
        for width, dpi, span in ((8, 100, 30), (4, 150, 90), (12, 200, 3), (8, 100, 30)):
            with self.subTest(width=width, dpi=dpi, span=span):
                app.current_figure.set_size_inches(width, 5)
                app.current_figure.set_dpi(dpi)
                app._apply_shared_plot_viewport((day-1, day+span))
                self.assert_311_slots(app)
                first = [label.get_position() for label in app.ev_texts]
                self.assert_311_slots(app)
                np.testing.assert_allclose(first, [label.get_position() for label in app.ev_texts])
        app.ev_texts[0].set_text('Long localized historical context (12/123)')
        self.assert_311_slots(app)
        self.assertEqual(app.ev_texts[0].get_text(), 'Long localized historical context (12/123)')
        with patch.object(app.current_canvas, 'draw_idle') as idle:
            app.current_canvas.draw()
            idle.assert_not_called()  # Layout happens in this paint, not a queued second one.
        directory = os.environ.get('PROGTRACK_311_EVIDENCE_DIR')
        if directory:
            app.current_figure.savefig(Path(directory) / 'near-date-line-labels.png')

    def test_311_localized_historical_labels_keep_their_full_text_and_counts(self):
        for locale in ('en', 'de', 'it', 'ru'):
            with self.subTest(locale=locale):
                app = self.line_label_host(3, same_date=True)
                app.messages = json.loads((ROOT / f'lang/messages_{locale}.json').read_text(encoding='utf-8'))
                # Detach the definition, retaining real saved history as in #296.
                app._role_dialog_blocks = lambda *_a, **_k: ['weight']
                app._plot_selected()
                original = [label.get_text() for label in app.ev_texts]
                self.assertTrue(all(app.messages['plot.event.historical'] in text for text in original))
                self.assert_311_slots(app)
                self.assertEqual(original, [label.get_text() for label in app.ev_texts])
                self.assertTrue(all('/3' not in text for text in original))

    def test_311_mixed_weight_and_event_only_animals_keep_separate_label_axes(self):
        app = self.line_label_host(3, measured=True, same_date=True)
        app.animals['Second'] = copy.deepcopy(app.animals['Subject'])
        app.animals['Second']['gewicht'] = []
        app.selected_animals.append('Second')
        snapshot = copy.deepcopy(app.animals)
        app._plot_selected()
        self.assert_311_slots(app)
        self.assertTrue(all(label.axes is app.weight_axes[0] for label in app.ev_texts[:3]))
        self.assertTrue(all(label.axes is app._plot_axes[1] for label in app.ev_texts[3:]))
        self.assertEqual(app.animals, snapshot)

    def test_311_toggle_and_line_label_hit_regions_never_create_tooltips(self):
        app = self.line_label_host(3, same_date=True)
        app._plot_selected()
        self.assert_311_slots(app)
        for label in app.ev_texts:
            box = label.get_window_extent(app.current_canvas.get_renderer())
            for name in ('motion_notify_event', 'button_press_event'):
                self.tooltip_mouse(app, name, point=((box.x0+box.x1)/2, (box.y0+box.y1)/2))
                self.assertIsNone(app._annotation)
        app.chk_events.setChecked(False)
        app.current_canvas.draw()
        self.assertTrue(all(not artist.get_visible() for artist in app.ev_lines+app.ev_texts))
        self.assertEqual(app.current_figure._progtrack_event_label_layout.right_extent_by_axis, {})
        app.chk_events.setChecked(True)
        self.assert_311_slots(app)
        self.assertTrue(all(artist.get_visible() for artist in app.ev_lines+app.ev_texts))

    def test_311_replots_own_one_layout_artist_and_no_extra_canvas_callbacks(self):
        app = self.line_label_host(12)
        previous = None
        for _ in range(3):
            app._plot_selected()
            self.assert_311_slots(app)
            canvas = app.current_canvas
            current_ids = list(app._mpl_cids)
            self.assertEqual(len(current_ids), 7)  # #296 lifecycle; layout adds none.
            self.assertEqual(sum(isinstance(item, MAIN._EventLineLabelLayout)
                                 for item in app.current_figure.artists), 1)
            if previous is not None:
                old, ids = previous
                for callbacks in old.callbacks.callbacks.values():
                    self.assertTrue(all(cid not in callbacks for cid in ids))
            previous = (canvas, current_ids)

    def test_296_real_measurement_point_tooltip_and_sperm_rounding_remain_unchanged(self):
        app = self.tooltip_host(measured=True)
        app.animals['Subject']['gewicht'][0]['wert'] = 434.7
        app.animals['Subject']['ref_weight'] = 440
        app._plot_selected()
        entry = next(entry for entry in app.hover_data if len(entry) == 6 and entry[3] == 'weight')
        point = entry[2].transData.transform((MAIN.mdates.date2num(entry[0]), entry[1]))
        self.tooltip_mouse(app, 'button_press_event', point=point)
        self.assertIn('434.70', app._annotation.get_text())
        self.assertIn('-1.20%', app._annotation.get_text())
        self.assertFalse(getattr(app._annotation, '_progtrack_event_tooltip', False))
        self.assertIn('56%', app._format_tooltip('sperm_motility', datetime(2026, 1, 1), 123.6,
            'Subject', {'mot_pct': 55.55}))
        for event_type, label in (('fsh', 'FSH'), ('progesterone', 'Progesterone')):
            with self.subTest(builtin_event=event_type):
                builtin = self.tooltip_host()
                builtin.animals['Subject']['events'][0]['event_type'] = event_type
                builtin._plot_selected()
                artist = builtin.ev_lines[0]
                builtin.current_canvas.draw()
                point = artist.get_offset_transform().transform(artist.get_offsets()[0])
                point[0] += 15  # Preserve the accepted 18px built-in click tolerance.
                self.tooltip_mouse(builtin, 'button_press_event', point=point)
                self.assertIn(f'{label} – 01.01.2026', builtin._annotation.get_text())
                self.assertNotIn('Value:', builtin._annotation.get_text())

    def test_296_missing_real_limit_and_localized_tooltip_labels_are_explicit(self):
        for language in ('en', 'de', 'it', 'ru'):
            for maximum in ('?', None, 0, 3):
                with self.subTest(language=language, maximum=maximum):
                    app = self.tooltip_host('limiting')
                    app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                    app._get_report_event_max = lambda *_a: maximum
                    record = app.animals['Subject']
                    context = app._plot_event_tooltip_contexts(record, MAIN.plot_event_entries(record),
                        scoped=False, active_period='')[id(record['events'][0])]
                    text = app._format_tooltip('custom_reading', datetime(2026, 1, 1), None,
                        'Subject', extra={'event_context': context, 'compact_event': True})
                    self.assertTrue(text.startswith('01.01.2026 - '))
                    self.assertIn('Custom reading(1/0)' if maximum == 0 else
                                  'Custom reading(1/3)' if maximum == 3 else 'Custom reading(1)', text)
                    self.assertNotIn('?', text)

    def test_296_fsh_progesterone_and_custom_actual_marker_matrix(self):
        for event_type in ('fsh', 'progesterone', 'custom_reading'):
            for marker in ('o', '^', 's', 'D', '*'):
                with self.subTest(event_type=event_type, marker=marker):
                    app = self.tooltip_host()
                    app.animals['Subject']['events'][0]['event_type'] = event_type
                    app.custom_event_styles = {event_type: {'marker': marker}}
                    app._plot_selected()
                    artist = app.ev_lines[0]
                    center = MAIN.mdates.date2num(datetime(2026, 1, 1))
                    artist.axes.set_xlim(center - 3, center + 3)
                    app.current_figure.set_size_inches(8, 6)
                    before = (copy.deepcopy(app.animals), artist.get_facecolors().copy(),
                              artist.get_edgecolors().copy(), artist.get_sizes().copy())
                    highlight = self.assert_296_hover(app, artist)
                    self.assertIsNone(app._annotation)
                    self.assertIs(highlight.get_offset_transform(), artist.get_offset_transform())
                    np.testing.assert_allclose(highlight.get_transform().get_matrix(),
                                               artist.get_transform().get_matrix())
                    self.assertEqual(app.animals, before[0])
                    np.testing.assert_array_equal(artist.get_facecolors(), before[1])
                    np.testing.assert_array_equal(artist.get_edgecolors(), before[2])
                    np.testing.assert_array_equal(artist.get_sizes(), before[3])
                    directory = os.environ.get('PROGTRACK_296_EVIDENCE_DIR')
                    if directory and (event_type, marker) == ('fsh', 'o'):
                        app.current_figure.savefig(Path(directory) / 'fsh-circle-hover.png')

    def test_296_nonstandard_marker_path_is_copied_without_renormalizing(self):
        from matplotlib.path import Path as MarkerPath
        from matplotlib.transforms import Affine2D
        app = self.tooltip_host()
        app._plot_selected()
        artist = app.ev_lines[0]
        custom = MarkerPath([[-0.12, -0.2], [0.3, 0.05], [-0.05, 0.18], [-0.12, -0.2]],
                            [MarkerPath.MOVETO, MarkerPath.LINETO, MarkerPath.LINETO,
                             MarkerPath.CLOSEPOLY])
        artist.set_paths([custom])
        artist.set_transform(Affine2D().rotate_deg(23))
        center = MAIN.mdates.date2num(datetime(2026, 1, 1))
        artist.axes.set_xlim(center - 3, center + 3)
        app.current_figure.set_size_inches(8, 6)
        self.assert_296_hover(app, artist)
        np.testing.assert_array_equal(artist.get_paths()[0].vertices, custom.vertices)
        directory = os.environ.get('PROGTRACK_296_EVIDENCE_DIR')
        if directory:
            app.current_figure.savefig(Path(directory) / 'custom-path-hover.png')

    def test_296_weight_hover_uses_own_marker_and_keeps_measurement_values(self):
        for marker in ('^', 'D', '*', 's'):
            with self.subTest(marker=marker):
                app = self.tooltip_host(measured=True)
                app.weight_marker = marker
                app.animals['Subject']['gewicht'][0]['wert'] = 434.7
                app.animals['Subject']['ref_weight'] = 440
                app._plot_selected()
                entry = next(item for item in app.hover_data if len(item) == 6 and item[3] == 'weight')
                day, value, axis, _typ, _name, artist = entry
                center = MAIN.mdates.date2num(day)
                axis.set_xlim(center - 3, center + 3)
                point = artist.get_transform().transform((MAIN.mdates.date2num(day), value))
                self.assert_296_hover(app, artist, point=point)
                self.assertIsNone(app._annotation)
                directory = os.environ.get('PROGTRACK_296_EVIDENCE_DIR')
                if directory and marker == '*':
                    app.current_figure.savefig(Path(directory) / 'weight-star-hover.png')
                self.tooltip_mouse(app, 'button_press_event', point=point)
                text = app._annotation.get_text()
                self.assertIn('434.70', text)
                self.assertIn('-1.20%', text)
                self.assertFalse(getattr(app._annotation, '_progtrack_event_tooltip', False))
                app.chk_weight.setChecked(False)
                app._apply_plot_display_state()
                self.assertIsNone(app._highlight_scatter)

    def test_296_clicked_popup_persists_on_hover_but_highlight_does_not_leak(self):
        from matplotlib.backend_bases import LocationEvent
        app = self.tooltip_host()
        app._plot_selected()
        artist = app.ev_lines[0]
        self.tooltip_mouse(app, 'button_press_event', artist)
        clicked = app._annotation
        self.assert_296_hover(app, artist)
        self.assertIs(app._annotation, clicked)
        app.current_canvas.callbacks.process('figure_leave_event',
            LocationEvent('figure_leave_event', app.current_canvas, -1, -1))
        self.assertIsNone(app._highlight_scatter)
        self.assertIs(app._annotation, clicked)
        self.assert_296_hover(app, artist)
        app.chk_events.setChecked(False)
        self.assertIsNone(app._highlight_scatter)
        self.assertIsNone(app._annotation)
        app.chk_events.setChecked(True)
        self.tooltip_mouse(app, 'button_press_event', artist)
        self.tooltip_mouse(app, 'button_press_event', point=(5, 5))
        self.assertIsNone(app._annotation, 'An ordinary empty click dismisses text')

    def test_296_historical_period_role_unassigned_contexts_and_locales(self):
        for locale in ('en', 'de', 'it', 'ru'):
            with self.subTest(locale=locale):
                app = self.tooltip_host('limiting')
                app.messages = json.loads((ROOT / f'lang/messages_{locale}.json').read_text(encoding='utf-8'))
                record = app.animals['Subject']
                record.update(project='Current', project_period_id='new', project_entry_date='01.01.2026',
                    project_history=[{'project': 'Former', 'period_id': 'old',
                                      'entry_date': '01.01.2025', 'leave_date': '31.12.2025'}])
                record['events'] = [
                    {'event_id': key, 'event_type': 'custom_reading', 'datum': day,
                     'recorded_role': role, 'project_period_id': period,
                     'custom_event_snapshot': {'name': 'Stored title', 'kind': 'limiting'}}
                    for key, day, role, period in (
                        ('former', datetime(2025, 6, 1), 'custom_role', 'old'),
                        ('role', datetime(2026, 1, 1), 'custom.former_role', 'new'),
                        ('unassigned', datetime(2024, 1, 1), 'custom_role', 'unassigned'),
                        ('current', datetime(2026, 1, 2), 'custom_role', 'new'))]
                MAIN.ensure_project_periods(record)  # Existing renderer normalization, before oracle.
                before = copy.deepcopy(record)
                app._plot_selected()
                app._apply_shared_plot_viewport((MAIN.mdates.date2num(datetime(2023, 12, 1)),
                                                 MAIN.mdates.date2num(datetime(2026, 2, 1))),
                                                app.current_figure.axes)
                for artist in app.ev_lines:
                    context = artist._progtrack_event_context
                    day = next(item[0] for item in app.hover_data if item[5] is artist)
                    center = MAIN.mdates.date2num(day)
                    artist.axes.set_xlim(center - 10, center + 10)
                    app.current_canvas.draw()  # Separate near dates beyond integer pixel rounding.
                    self.assert_296_hover(app, artist)
                    self.tooltip_mouse(app, 'button_press_event', artist)
                    text = app._annotation.get_text()
                    block = next(part for part in text.split('\n\n')
                                 if day.strftime('%d.%m.%Y') + ' - Stored title' in part)
                    self.assertIn(' - Stored title(1/3)' if not context[3] else ' - Stored title(1)', text)
                    if context[3]:
                        self.assertIn(app.messages['plot.event.historical'], block)
                        self.assertIn(app.messages['plot.tooltip.event_period'], block)
                        if context[4] == 'unassigned':
                            self.assertIn(app.messages['plot.tooltip.unassigned_period'], block)
                        self.assertNotIn('/3', block)
                    self.assertNotIn(app.messages['plot.tooltip.recorded_role'], text)
                    self.assertNotIn('Recorded role', text)
                    self.assertNotIn(context[5], text)
                self.assertEqual(record, before)

    def test_296_coincident_historical_popup_hides_role_but_preserves_scoped_provenance(self):
        for locale in ('en', 'de', 'it', 'ru'):
            with self.subTest(locale=locale):
                app = self.tooltip_host('limiting')
                app.messages = json.loads((ROOT / f'lang/messages_{locale}.json').read_text(encoding='utf-8'))
                record = app.animals['Subject']
                record['events'] = [
                    {'event_id': identity, 'event_type': 'custom_reading',
                     'datum': datetime(2026, 1, 1), 'recorded_role': role,
                     'custom_event_snapshot': {'name': title, 'kind': 'limiting',
                                                'render_mode': 'symbol', 'marker': 'D'}}
                    for identity, role, title in (
                        ('former', 'custom.new_role', 'Former title'),
                        ('current', 'custom_role', 'Current title'))]
                before = copy.deepcopy(record)
                app._plot_selected()
                self.tooltip_mouse(app, 'button_press_event', app.ev_lines[0])
                text = app._annotation.get_text()
                self.assertLess(text.index('Former title(1)'), text.index('Current title(1/3)'))
                self.assertIn(app.messages['plot.event.historical'], text)
                self.assertNotIn(app.messages['plot.tooltip.recorded_role'], text)
                self.assertNotIn('Recorded role', text)
                self.assertNotIn('custom.new_role', text)
                self.assertNotIn('custom_role', text)
                self.assertEqual(record, before)
                contexts = app._plot_event_tooltip_contexts(record, MAIN.plot_event_entries(record),
                                                           scoped=False, active_period='')
                self.assertEqual(contexts[id(record['events'][0])][1:],
                                 (1, None, True, '', 'custom.new_role'))
                self.assertEqual(contexts[id(record['events'][1])][1:],
                                 (1, '3', False, '', 'custom_role'))

    def test_296_clicked_event_popup_tracks_zoom_resize_and_clipped_source(self):
        app = self.tooltip_host()
        app._plot_selected()
        artist = app.ev_lines[0]
        self.tooltip_mouse(app, 'button_press_event', artist)
        popup = app._annotation
        app.current_figure.set_size_inches(6, 3)
        center = MAIN.mdates.date2num(datetime(2026, 1, 1))
        artist.axes.set_xlim(center - 1, center + 0.05)
        app.current_canvas.draw()
        self.assertIs(app._annotation, popup)
        bounds = popup.get_window_extent(app.current_canvas.get_renderer())
        self.assertGreaterEqual(bounds.x0, app.current_figure.bbox.x0)
        self.assertLessEqual(bounds.x1, app.current_figure.bbox.x1)
        self.assertGreaterEqual(bounds.y0, app.current_figure.bbox.y0)
        self.assertLessEqual(bounds.y1, app.current_figure.bbox.y1)
        artist.axes.set_xlim(center + 1, center + 2)
        app.current_canvas.draw()
        self.assertIsNone(app._annotation)

    def test_266_271_event_only_plots_initial_symbols_lines_and_redraw_lifecycle(self):
        for event_mode, measured, multiple in (('symbol', False, False),
                                               ('symbol', True, False),
                                               ('line', False, True)):
            with self.subTest(mode=event_mode, measured=measured, multiple=multiple):
                app = self.plot_host(event_mode, measured, multiple)
                managers = set(Gcf.figs)
                unexpected = []
                class Observer(MAIN.QtCore.QObject):
                    def eventFilter(self, widget, event):
                        if event.type() == QEvent.Type.Show and isinstance(widget, QWidget) and widget.isWindow():
                            unexpected.append(widget.windowTitle())
                        return False
                observer = Observer()
                self.qt.installEventFilter(observer)
                try:
                    for _ in range(3):
                        old_canvas = app.current_canvas
                        app._plot_selected()
                        app.current_canvas.draw()
                        self.qt.processEvents()
                        self.qt.sendPostedEvents(None, QEvent.Type.DeferredDelete)
                        if old_canvas is not None:
                            self.assertTrue(sip.isdeleted(old_canvas))
                        self.assertEqual(set(Gcf.figs), managers)
                        self.assertFalse(app.current_canvas.isWindow())
                        self.assertTrue(app.chk_events.isEnabled())
                        self.assertEqual(len(app.ev_lines), len(app.selected_animals))
                        summaries = [widget for widget in app.detail_widget.findChildren(MAIN.QLabel)
                                     if widget.property('animalKey') is not None]
                        self.assertEqual(len(summaries), len(app.selected_animals))
                        self.assertTrue(all(widget.property('projectLabel') == '' for widget in summaries))
                        self.assertTrue(all(widget.text().startswith('Subject — ') for widget in summaries))
                        if event_mode == 'line':
                            self.assertTrue(all('/' not in label.get_text() for label in app.ev_texts))
                        else:
                            self.assertTrue(all(artist.get_visible() for artist in app.ev_lines))
                    app.selected_animals = []
                    app._plot_selected()
                    self.qt.sendPostedEvents(None, QEvent.Type.DeferredDelete)
                    self.qt.processEvents()
                    self.assertIsNone(app.current_canvas)
                    self.assertEqual(app.hover_data, [])
                    self.assertEqual(unexpected, [])
                finally:
                    self.qt.removeEventFilter(observer)

    def test_271_animal_editors_do_not_show_controls_before_dialog_exec(self):
        dialogs = {
            '_dlg_partner': 'partner', '_dlg_samenspender': 'sperm_donor',
            '_dlg_offspring': 'offspring', '_dlg_zuchttier': 'breeding_animal',
            '_dlg_versuchstier': 'experimental', '_dlg_basic_animal_role': 'custom_role',
            '_dlg_female_animal': 'donor',
        }
        for method, role in dialogs.items():
            with self.subTest(dialog=method):
                app = self.host()
                app.animals = {'Subject': self.record(role)}
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
                unexpected = []

                class Observer(MAIN.QtCore.QObject):
                    def eventFilter(self, widget, event):
                        if event.type() == QEvent.Type.Show and isinstance(widget, QWidget) and widget.isWindow():
                            unexpected.append((type(widget).__name__, widget.windowTitle()))
                        return False

                observer = Observer()
                self.qt.installEventFilter(observer)
                try:
                    dlg = self.open_dialog(app, method)
                    self.qt.processEvents()
                    self.assertEqual(unexpected, [], method)
                    self.assertFalse(dlg.isVisible())
                    dlg.show()
                    self.qt.processEvents()
                    self.assertTrue(dlg.isVisible())
                    self.assertEqual(unexpected, [(type(dlg).__name__, dlg.windowTitle())])
                    dlg.reject()
                    self.qt.processEvents()
                    self.assertFalse(dlg.isVisible())
                finally:
                    self.qt.removeEventFilter(observer)

    def test_309_status_limits_and_localization_matrix(self):
        for language in ('en', 'de', 'it', 'ru'):
            for in_experiment, old_count, added, maximum, expected in (
                (False, 0, 1, 2, 1), (True, 0, 1, 2, 0),
                (True, 1, 1, 2, 0), (True, 2, 1, 2, 1),
                (False, 2, 2, 2, 2), (False, 3, 0, 2, 0),
                (True, 0, 1, 0, 1)):
                with self.subTest(language=language, status=in_experiment,
                                  old=old_count, added=added, maximum=maximum):
                    app = self.host()
                    app.messages = json.loads((ROOT / f'lang/messages_{language}.json').read_text(encoding='utf-8'))
                    previous = self.record('experimental')
                    previous['max_measurements'] = maximum
                    previous['events'] = [{'event_id': f'old-{i}', 'event_type': 'measurement',
                        'recorded_role': 'experimental_animal', 'datum': datetime(2025, 1, 1)}
                        for i in range(old_count)]
                    candidate = copy.deepcopy(previous)
                    candidate['in_experiment'] = in_experiment
                    candidate['events'] += [{'event_type': 'measurement', 'datum': datetime(2025, 1, 1)}
                                             for _ in range(added)]
                    snapshot = copy.deepcopy((previous, candidate))
                    prompts = []
                    app._show_message_raw = lambda *a, **k: prompts.append((a, k)) or MAIN.QMessageBox.StandardButton.Yes
                    self.assertTrue(app._confirm_new_animal_events(previous, candidate))
                    self.assertEqual(len(prompts), expected)
                    self.assertEqual((previous, candidate), snapshot)
                    for args, kwargs in prompts:
                        self.assertEqual(args[0], app.messages['confirm.event_save.title'])
                        self.assertNotIn('{', args[1])
                        self.assertTrue(kwargs['plain_text'])
                    if added and old_count + added > maximum:
                        text = prompts[-1][0][1]
                        self.assertIn(app.messages['stats.measurement'], text)
                        self.assertIn(str(old_count + added), text)
                    # Cancel is not cached as an approval on the next attempt.
                    if expected:
                        app._show_message_raw = lambda *_a, **_k: MAIN.QMessageBox.StandardButton.No
                        self.assertFalse(app._confirm_new_animal_events(previous, candidate))

    def test_309_all_builtin_limited_event_types_and_aggregation(self):
        fields = {'surgery': 'max_op', 'measurement': 'max_measurements',
                  'pgf': 'max_pgf', 'fsh': 'max_fsh', 'embryo_transfer': 'max_embryo',
                  'pregnancy': 'max_pregnancies', 'birth': 'max_geburten'}
        for role, types in ((MAIN.Role.SPENDER.value, ('surgery', 'pgf', 'fsh')),
                            ('surrogate', ('embryo_transfer', 'pregnancy', 'birth', 'pgf')),
                            ('breeding_animal', ('pregnancy', 'birth')),
                            ('experimental', ('surgery', 'measurement')),
                            ('experimental_offspring', ('surgery', 'measurement'))):
            with self.subTest(role=role):
                app = self.host()
                previous = self.record(role)
                for event_type in types:
                    previous[fields[event_type]] = 1
                previous['in_experiment'] = True
                candidate = copy.deepcopy(previous)
                candidate['events'] = [{'event_type': event_type, 'datum': datetime(2025, 1, 1)}
                                       for event_type in types for _ in range(2)]
                prompts = []
                app._show_message_raw = lambda *a, **k: prompts.append(a) or MAIN.QMessageBox.StandardButton.Yes
                self.assertTrue(app._confirm_new_animal_events(previous, candidate))
                self.assertEqual(len(prompts), 1)
                overruns = app._event_save_limit_overruns(previous, candidate)
                self.assertEqual(len(overruns), len(types))
                self.assertTrue(all(maximum == 1 and count == 2 for _, maximum, count in overruns))

    def test_309_projection_preserves_role_period_and_custom_applicability(self):
        from Plugins.core.project_periods import ensure_project_periods
        app = self.tooltip_host('limiting')
        previous = app.animals['Subject']
        previous.update(project='Current', project_entry_date='01.02.2025', in_experiment=True,
            project_history=[{'project': 'Former', 'entry_date': '01.01.2025', 'leave_date': '01.02.2025'}])
        previous['events'] = [{'event_id': identity, 'event_type': 'custom_reading',
            'datum': datetime(2025, month, 4), 'recorded_role': role}
            for identity, month, role in (('former', 1, 'custom_role'), ('other-role', 2, 'offspring'))]
        ensure_project_periods(previous)
        candidate = copy.deepcopy(previous)
        candidate['events'].append({'event_type': 'custom_reading', 'datum': datetime(2025, 2, 5)})
        self.assertEqual(app._event_save_limit_overruns(previous, candidate), [])  # 1/3, not 3/3
        candidate['events'] += [{'event_type': 'custom_reading', 'datum': datetime(2025, 2, 5)} for _ in range(3)]
        snapshot = copy.deepcopy((previous, candidate))
        self.assertEqual(app._event_save_limit_overruns(previous, candidate), [('Custom reading', 3, 4)])
        self.assertEqual((previous, candidate), snapshot)
        payload = app.backend.records.get('configuration', 'role-block-presets')
        second = copy.deepcopy(payload['event_definitions'][0])
        second.update(event_type='custom_second', name='Second')
        payload['event_definitions'].append(second)
        app.backend.records.put('configuration', 'role-block-presets', payload)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', second['block_id']]
        grouped = copy.deepcopy(candidate)
        grouped['events'].append({'event_type': 'custom_second', 'datum': datetime(2025, 2, 5)})
        self.assertEqual(app._event_save_limit_overruns(previous, grouped), [('Custom reading, Second', 3, 5)])
        app._role_dialog_blocks = lambda *_a, **_k: ['weight']
        self.assertEqual(app._event_save_limit_overruns(previous, candidate), [])
        app._role_dialog_blocks = lambda *_a, **_k: ['custom_reading']
        payload = app.backend.records.get('configuration', 'role-block-presets')
        payload['event_definitions'][0]['active'] = False
        app.backend.records.put('configuration', 'role-block-presets', payload)
        self.assertEqual(app._event_save_limit_overruns(previous, candidate), [])
        # Existing ID date/role/scalar changes alone are never additions.
        candidate = copy.deepcopy(previous)
        candidate['rolle'] = 'offspring'
        candidate['events'][0]['datum'] = datetime(2025, 2, 5)
        app._show_message_raw = lambda *_a, **_k: self.fail('existing events must not prompt')
        self.assertTrue(app._confirm_new_animal_events(previous, candidate))

    def test_309_five_editors_decline_retry_accept_sqlite(self):
        cases = (('_dlg_offspring', 'offspring', 'custom.audit_event', None),
                 ('_dlg_zuchttier', 'breeding_animal', 'pregnancy', 'max_pregnancies'),
                 ('_dlg_versuchstier', 'experimental', 'measurement', 'max_measurements'),
                 ('_dlg_female_animal', MAIN.Role.SPENDER.value, 'pgf', 'max_pgf'),
                 ('_dlg_basic_animal_role', 'custom_role', 'custom.audit_event', None))
        for method, role, event_type, limit_field in cases:
            with self.subTest(editor=method), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                backend = ProgTrackBackend(runtime_paths(root, application_root=root), acquire_process_lock=False)
                self.addCleanup(backend.close)
                app = self.host()
                app.backend = backend
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom.audit_event']
                app._retain_selection_after_animal_save = lambda *_a, **_k: None
                app._run_post_persistence_sync = lambda *_a, **_k: None
                app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [{
                    'event_type': 'custom.audit_event', 'block_id': 'custom.audit_event',
                    'name': 'Test-Limit', 'kind': 'limiting', 'default_maximum': 1, 'active': True}]})
                record = self.record(role)
                record.update(in_experiment=False, edits={'2025-01-01': {'keep': 'original'}},
                    experimental_limits={'custom.audit_event': 1})
                if limit_field:
                    record[limit_field] = 1
                record['events'] = [{'event_id': 'existing', 'event_type': event_type,
                    'datum': datetime(2025, 1, 1), 'recorded_role': MAIN.canonical_role_value(role)}]
                app.animals['Subject'] = record
                real_save = MAIN.ProgTrackApp._save_persistence.__get__(app)
                real_save()
                disk = copy.deepcopy(backend.load_core_data())
                memory = copy.deepcopy(app.animals)
                writes, side_effects = [], []
                app._save_persistence = lambda **kwargs: writes.append('save') or real_save(**kwargs)
                app._schedule_post_animal_save_project_updates = lambda *_a, **_k: side_effects.append('project')
                app._rewrite_animal_references_after_identity_change = lambda *_a, **_k: side_effects.append('rewrite')
                app._audit_animal_created = lambda *_a, **_k: side_effects.append('audit')
                dialog = self.open_dialog(app, method)
                combo = self.event_editor_rows_247(dialog)['existing'][1]
                scroll = combo.parentWidget().parentWidget().parentWidget()
                add = next(b for b in scroll.findChildren(QPushButton)
                           if b.text().casefold() in ('new event', 'add events'))
                for _ in range(2):
                    add.click()
                    self.qt.processEvents()
                    row = next((field, control) for field, control, _button in
                               self.event_editor_rows_247(dialog).values() if not field.text() == '02.01.2025'
                               and not getattr(control, '_progtrack_source_event', {}).get('event_id'))
                    row[0].setText('02.01.2025')
                    self.assertGreaterEqual(row[1].findData(event_type), 0)
                    row[1].setCurrentIndex(row[1].findData(event_type))
                save_button = next(b for b in dialog.findChildren(QPushButton) if b.text() == app.messages['button.save'])
                for answers, expected_prompts in (([False], 1), ([True, False], 2), ([True, True], 2)):
                    prompts, errors = [], []
                    decisions = iter(answers)
                    def answer(*args, **kwargs):
                        prompts.append(args)
                        # Reentrant Save while a modal warning is active must
                        # not start another attempt or publish any state.
                        save_button.click()
                        self.assertEqual(writes, [])
                        self.assertEqual(side_effects, [])
                        self.assertEqual(backend.load_core_data(), disk)
                        self.assertEqual(app.animals, memory)
                        return (MAIN.QMessageBox.StandardButton.Yes if next(decisions)
                                else MAIN.QMessageBox.StandardButton.No)
                    app._show_message_raw = answer
                    with patch('sys.excepthook', lambda *a: errors.append(a)):
                        save_button.click()
                        self.qt.processEvents()
                    self.assertEqual(errors, [])
                    self.assertEqual(len(prompts), expected_prompts)
                    if not all(answers):
                        self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)
                        self.assertTrue(save_button.isEnabled())
                        self.assertEqual(backend.load_core_data(), disk)
                        self.assertEqual(app.animals, memory)
                self.assertEqual(writes, ['save'])
                self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted)
                saved = backend.load_core_data()['animals']['Subject']['events']
                self.assertEqual(len(saved), 3)
                self.assertEqual(len({e['event_id'] for e in saved}), 3)
                self.assertIn('existing', {e['event_id'] for e in saved})
                save_button.click()
                self.assertEqual(writes, ['save'])
                # Existing over-limit history alone does not prompt.
                app._show_message_raw = lambda *_a, **_k: self.fail('unchanged save prompted')
                unchanged = self.open_dialog(app, method)
                next(b for b in unchanged.findChildren(QPushButton)
                     if b.text() == app.messages['button.save']).click()
                self.qt.processEvents()
                self.assertEqual(backend.load_core_data()['animals']['Subject']['events'], saved)

    def test_309_effective_status_uses_authorized_checkbox_or_coerced_result(self):
        cases = [(method, permitted, coerced, expected)
                 for method in ('_dlg_versuchstier', '_dlg_basic_animal_role')
                 for permitted, coerced, expected in ((True, True, 0), (False, False, 1), (True, False, 1))]
        for method, permitted, coerced, expected in cases:
            with self.subTest(editor=method, permitted=permitted, coerced=coerced):
                app = self.host()
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'health_flags', 'project_severity', 'measurement']
                app.backend.records.put('configuration', 'role-block-presets', {'event_definitions': [{
                    'event_type': 'measurement', 'block_id': 'measurement', 'kind': 'counting',
                    'name': 'Measurement', 'active': True}]})
                app._is_projects_track_active = lambda: True
                app._project_names_for_species = lambda _s: ['Current']
                app._master_can = lambda permission, **_k: permitted or permission != 'project.set_in_experiment'
                app._coerce_in_experiment_for_project = lambda value, *_a, **_k: value if coerced else False
                record = self.record('experimental_animal')
                record.update(project='Current', project_entry_date='01.01.2025',
                              in_experiment=False, max_measurements=10)
                app.animals['Subject'] = record
                dialog = self.open_dialog(app, method)
                dialog.show()
                self.qt.processEvents()
                checkbox = next(c for c in dialog.findChildren(QCheckBox)
                                if c.text() == app.messages['checkbox.in_experiment'])
                checkbox.setChecked(True)  # also tests forged disabled input
                add = next(b for b in dialog.findChildren(QPushButton) if b.text().casefold() == 'new event')
                add.click()
                row = self.event_editor_rows_247(dialog)['']
                row[0].setText('01.02.2025')
                row[1].setCurrentIndex(row[1].findData('measurement'))
                prompts = []
                app._show_message_raw = lambda *a, **k: prompts.append(a) or MAIN.QMessageBox.StandardButton.Yes
                next(b for b in dialog.findChildren(QPushButton)
                     if b.text() == app.messages['button.save']).click()
                self.qt.processEvents()
                self.assertEqual(len(prompts), expected, prompts)
                self.assertEqual(app.animals['Subject']['in_experiment'], coerced)

    def test_309_period_change_preview_and_retained_overrun_are_read_only(self):
        from Plugins.core.project_periods import ensure_project_periods
        app = self.host()
        before = self.record('experimental_animal')
        before.update(project='Former', project_entry_date='01.01.2020',
                      in_experiment=True, max_measurements=1)
        before['events'] = [{'event_id': f'old-{i}', 'event_type': 'measurement',
            'recorded_role': 'experimental_animal', 'datum': datetime.now()} for i in range(4)]
        ensure_project_periods(before)
        candidate = copy.deepcopy(before)
        candidate['project'] = 'New'
        candidate['events'].append({'event_type': 'measurement', 'datum': datetime.now()})
        snapshot = copy.deepcopy((before, candidate))
        self.assertEqual(app._event_save_limit_overruns(before, candidate), [])  # New period 1/1
        candidate['events'].append({'event_type': 'measurement', 'datum': datetime.now()})
        self.assertEqual(app._event_save_limit_overruns(before, candidate), [('Measurement', 1, 2)])
        candidate['events'].pop()
        self.assertEqual((before, candidate), snapshot)
        # Adding while deleting even more does not increase an old overrun.
        candidate = copy.deepcopy(before)
        candidate['events'] = candidate['events'][:2] + [{'event_type': 'measurement', 'datum': datetime.now()}]
        self.assertEqual(app._event_save_limit_overruns(before, candidate), [])
        # A changed maximum in this very save uses the effective dialog value,
        # not the old canonical-map projection waiting for persistence sync.
        candidate = copy.deepcopy(before)
        candidate['experimental_limits'] = {'max_measurements': 10}
        candidate['max_measurements'] = 0
        candidate['events'].append({'event_type': 'measurement', 'datum': datetime.now()})
        self.assertEqual(app._event_save_limit_overruns(before, candidate), [('Measurement', 0, 5)])

    def test_309_new_dialogs_without_event_additions_never_prompt(self):
        for method in ('_dlg_offspring', '_dlg_zuchttier', '_dlg_versuchstier',
                       '_dlg_female_animal', '_dlg_basic_animal_role'):
            with self.subTest(editor=method):
                source = inspect.getsource(getattr(MAIN.ProgTrackApp, method))
                self.assertIn('_confirm_new_animal_events(', source)
                self.assertIn('_connect_event_save_button(', source)
                self.assertNotIn('_custom_event_limit_additions_allowed(', source)
        app = self.host()
        app._show_message_raw = lambda *_a, **_k: self.fail('empty New Animal save prompted')
        self.assertTrue(app._confirm_new_animal_events({}, self.record('experimental')))

    def test_247_custom_quotas_require_active_attached_role_block(self):
        app = self.host()
        definition = {'event_type': 'custom.limit', 'block_id': 'custom.limit.block',
                      'name': 'Test-Limit', 'kind': 'limiting', 'active': True,
                      'max_field': 'max_custom_limit', 'default_maximum': 0}
        record = self.record('custom_role')
        record['experimental_limits'] = {'custom.limit.block': 0}
        record['max_custom_limit'] = 99  # stale flat value cannot create applicability
        for active in (True, False):
            for attached in (True, False):
                for count in (0, 1):
                    with self.subTest(active=active, attached=attached, count=count):
                        definition['active'] = active
                        app.backend.records.put('configuration', 'role-block-presets',
                                                {'event_definitions': [definition]})
                        app._role_dialog_blocks = lambda *_a, **_k: (
                            ['custom.limit.block'] if attached else ['weight'])
                        record['events'] = [{'event_type': 'custom.limit',
                                             'datum': datetime(2025, 1, 1),
                                             'recorded_role': 'custom_role'}] * count
                        applicable = active and attached
                        expected = ([f'Test-Limit: {count}/0'] if applicable else
                                    ['Test-Limit: 1'] if count else [])
                        self.assertEqual(app._custom_event_statistics(record), expected)
                        self.assertEqual(app._get_report_event_max('custom.limit', record),
                                         0 if applicable else None)
        # The same block may be attached by canonical event type, not just ID.
        definition['active'] = True
        app.backend.records.put('configuration', 'role-block-presets',
                                {'event_definitions': [definition]})
        app._role_dialog_blocks = lambda *_a, **_k: ['custom.limit']
        record['events'] = []
        self.assertEqual(app._custom_event_statistics(record), ['Test-Limit: 0/0'])
        # Missing is not explicit zero: preserve the existing omission of an
        # empty statistic when an applicable definition has no maximum at all.
        definition['default_maximum'] = None
        record['experimental_limits'] = {}
        record['max_custom_limit'] = None
        app.backend.records.put('configuration', 'role-block-presets',
                                {'event_definitions': [definition]})
        self.assertEqual(app._custom_event_statistics(record), [])
        self.assertIsNone(app._get_report_event_max('custom.limit', record))

    def event_editor_rows_247(self, dialog):
        rows = {}
        for combo in dialog.findChildren(QComboBox):
            source = getattr(combo, '_progtrack_source_event', None)
            if source is None:
                continue
            def find_row(layout):
                if layout is None:
                    return None
                widgets = [layout.itemAt(i).widget() for i in range(layout.count())]
                if combo in widgets:
                    return (next(w for w in widgets if isinstance(w, QLineEdit)), combo,
                            next(w for w in widgets if isinstance(w, QPushButton)))
                for i in range(layout.count()):
                    found = find_row(layout.itemAt(i).layout())
                    if found:
                        return found
                return None
            found = find_row(combo.parentWidget().layout())
            if found is None:
                # deleteLater children can remain discoverable until the next
                # deferred-delete pass; they no longer own an editor row.
                continue
            rows[str(source.get('event_id') or '')] = found
        return rows

    def test_247_custom_statistics_scope_zero_and_maximum_sources(self):
        from Plugins.core.project_periods import ensure_project_periods, current_period_id
        app = self.host()
        definition = {'name': 'Test-Limit', 'event_type': 'custom.limit',
                      'block_id': 'custom.block', 'max_field': 'max_custom',
                      'kind': 'limiting', 'active': True, 'default_maximum': 8}
        app.backend.records.put('configuration', 'role-block-presets',
                                {'event_definitions': [definition]})
        app._role_dialog_blocks = lambda role, *_a, **_k: (
            ['custom.block'] if role == 'custom_role' else ['weight'])
        record = self.record('custom_role')
        record.update(project='Current', project_entry_date='01.02.2025',
                      project_history=[{'project': 'Former', 'entry_date': '01.01.2025',
                                        'leave_date': '01.02.2025'}])
        record['events'] = [
            {'event_id': identity, 'event_type': 'custom.limit',
             'datum': datetime(2025, month, 3), 'recorded_role': role}
            for identity, month, role in (('old', 1, 'custom_role'),
                                          ('other', 2, 'offspring'),
                                          ('current', 2, 'custom_role'))]
        ensure_project_periods(record)
        snapshot = copy.deepcopy(record['events'])
        for limits, flat, expected in (({}, None, 8), ({}, 5, 5),
                                        ({'custom.block': 0}, 5, 0)):
            record['experimental_limits'] = limits
            record['max_custom'] = flat
            self.assertEqual(app._get_report_event_max('custom.limit', record), expected)
            self.assertEqual(app._custom_event_statistics(record), [f'Test-Limit: 1/{expected}'])
        record['rolle'] = 'offspring'
        self.assertEqual(app._custom_event_statistics(record), ['Test-Limit: 1'])
        self.assertIsNone(app._get_report_event_max('custom.limit', record))
        # Historical contexts keep their own prefix and never borrow a quota.
        record['rolle'] = 'custom_role'
        contexts = app._plot_event_tooltip_contexts(record, MAIN.plot_event_entries(record),
            scoped=True, active_period=current_period_id(record))
        for event in record['events']:
            context = contexts[id(event)]
            self.assertEqual(context[1], 1)
            self.assertEqual(context[2], '0' if event['event_id'] == 'current' else None)
        self.assertEqual(record['events'], snapshot)
        record['events'] = [e for e in record['events'] if e['event_id'] != 'current']
        self.assertEqual(app._custom_event_statistics(record), ['Test-Limit: 0/0'])
        self.assertEqual(record['events'], snapshot[:2])
        record['events'].append({'event_id': 'new-current', 'event_type': 'custom.limit',
            'datum': datetime(2025, 2, 4), 'recorded_role': 'custom_role',
            'project_period_id': current_period_id(record)})
        contexts = app._plot_event_tooltip_contexts(record, MAIN.plot_event_entries(record),
            scoped=True, active_period=current_period_id(record))
        self.assertEqual(contexts[id(record['events'][-1])][1], 1)
        definition['kind'] = 'counting'
        app.backend.records.put('configuration', 'role-block-presets',
                                {'event_definitions': [definition]})
        self.assertEqual(app._custom_event_statistics(record), ['Test-Limit: 1'])
        self.assertIsNone(app._get_report_event_max('custom.limit', record))

    def test_247_plot_rebuild_drops_deleted_same_day_context_and_last_marker(self):
        app = self.tooltip_host('counting')
        record = app.animals['Subject']
        record['events'] = [
            {'event_id': identity, 'event_type': 'custom_reading',
             'datum': datetime(2026, 1, day), 'recorded_role': 'custom_role'}
            for identity, day in (('first', 1), ('middle', 2), ('last', 2))]
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 2)
        group = app.ev_lines[1]._progtrack_coincident_event_items
        self.assertEqual([item[3][1] for item in group], [2, 3])
        record['events'] = [e for e in record['events'] if e['event_id'] != 'middle']
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 2)
        self.assertFalse(hasattr(app.ev_lines[1], '_progtrack_coincident_event_items'))
        self.assertEqual(app.ev_lines[1]._progtrack_event_context[1], 2)
        record['events'] = [e for e in record['events'] if e['event_id'] != 'last']
        app._plot_selected()
        self.assertEqual(len(app.ev_lines), 1)
        self.assertEqual([entry[0] for entry in app.hover_data], [datetime(2026, 1, 1)])
        record['events'] = []
        app._plot_selected()
        self.assertEqual(app.ev_lines, [])
        self.assertEqual(app.hover_data, [])

    def test_247_deletion_all_editors_first_middle_last_and_empty_sqlite(self):
        from Plugins.core.project_periods import ensure_project_periods
        from Plugins.Animal_Reports.report_index import ReportIndexRepository
        from datetime import date
        cases = (('_dlg_offspring', 'offspring', 'special_measurement'),
                 ('_dlg_zuchttier', 'breeding_animal', 'pregnancy'),
                 ('_dlg_versuchstier', 'experimental', 'measurement'),
                 ('_dlg_female_animal', 'donor', 'pgf'),
                 ('_dlg_basic_animal_role', 'custom_role', 'custom.audit_event'))
        for method, role, event_type in cases:
            with self.subTest(editor=method), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                backend = ProgTrackBackend(runtime_paths(root, application_root=root),
                                           acquire_process_lock=False)
                self.addCleanup(backend.close)
                app = self.host()
                app.backend = backend
                app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'custom.audit_event']
                app._retain_selection_after_animal_save = lambda *_a, **_k: None
                app._run_post_persistence_sync = lambda *_a, **_k: None
                app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(app)
                app._show_message_raw = lambda *_a, **_k: MAIN.QMessageBox.StandardButton.Yes
                app.backend.records.put('configuration', 'role-block-presets', {
                    'event_definitions': [{'event_type': 'custom.audit_event',
                        'block_id': 'custom.audit_event', 'kind': 'counting',
                        'name': 'Test-Event', 'active': True}]})
                record = self.record(role)
                record['events'] = [
                    {'event_id': identity, 'event_type': event_type,
                     'datum': datetime(2025, 1, day), 'recorded_role': role,
                     'custom_metadata': identity}
                    for identity, day in (('first', 1), ('middle', 2), ('last', 2))]
                ensure_project_periods(record)
                app.animals['Subject'] = record
                app._save_persistence()
                repository = ReportIndexRepository()
                original = copy.deepcopy(backend.load_core_data()['animals']['Subject']['events'])
                original_in_memory = copy.deepcopy(app.animals['Subject']['events'])
                before_index, _ = repository.get('Subject', record, str, date_format='%d.%m.%Y')
                self.assertEqual(before_index.event_counts_through(date(2025, 2, 1))[event_type], (3, 3))

                def save(dialog):
                    errors = []
                    with patch('sys.excepthook', lambda *args: errors.append(args)):
                        next(b for b in dialog.findChildren(QPushButton)
                             if b.text() == app.messages['button.save']).click()
                        self.qt.processEvents()
                    self.assertEqual(errors, [])

                # Cancelling an editor after deleting a row must not persist it.
                cancelled = self.open_dialog(app, method)
                self.event_editor_rows_247(cancelled)['middle'][2].click()
                self.qt.processEvents()
                cancelled.reject()
                self.assertEqual(backend.load_core_data()['animals']['Subject']['events'], original)
                self.assertEqual(app.animals['Subject']['events'], original_in_memory)
                # A rejected backend transaction must not partially commit a
                # deletion alongside an illegal edit of another saved event.
                rejected = copy.deepcopy(backend.load_core_data())
                rejected_events = rejected['animals']['Subject']['events']
                rejected_events[:] = [e for e in rejected_events if e['event_id'] != 'middle']
                rejected_events[0]['event_type'] = 'illegal_retype'
                with self.assertRaises(MAIN.ConflictError):
                    backend.save_core_data(rejected)
                self.assertEqual(backend.load_core_data()['animals']['Subject']['events'], original)
                # Remove middle, then first, then last: also tests an empty table.
                expected_ids = {'first', 'middle', 'last'}
                for deleted in ('middle', 'first', 'last'):
                    dialog = self.open_dialog(app, method)
                    rows = self.event_editor_rows_247(dialog)
                    # Built-in confirmation No must leave the saved row intact.
                    if method != '_dlg_basic_animal_role':
                        app._show_message_raw = lambda *_a, **_k: MAIN.QMessageBox.StandardButton.No
                        rows[deleted][2].click()
                        self.assertIn(deleted, self.event_editor_rows_247(dialog))
                        app._show_message_raw = lambda *_a, **_k: MAIN.QMessageBox.StandardButton.Yes
                    rows[deleted][2].click()
                    self.qt.processEvents()
                    remaining_rows = self.event_editor_rows_247(dialog)
                    if remaining_rows:
                        field = next(iter(remaining_rows.values()))[0]
                        valid_date = field.text()
                        disk_before_failed_save = backend.load_core_data()['animals']['Subject']['events']
                        memory_before_failed_save = copy.deepcopy(app.animals['Subject']['events'])
                        field.setText('not-a-date')
                        save(dialog)
                        self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)
                        self.assertEqual(backend.load_core_data()['animals']['Subject']['events'],
                                         disk_before_failed_save)
                        self.assertEqual(app.animals['Subject']['events'], memory_before_failed_save)
                        field.setText(valid_date)
                    save(dialog)
                    expected_ids.remove(deleted)
                    saved = backend.load_core_data()['animals']['Subject']
                    self.assertEqual({e['event_id'] for e in saved['events']}, expected_ids)
                    self.assertEqual(saved['events'], [e for e in original if e['event_id'] in expected_ids])
                    plotted_record = app.animals['Subject']
                    index, _ = repository.get('Subject', plotted_record, str, date_format='%d.%m.%Y')
                    self.assertEqual(index.event_counts_through(date(2025, 2, 1)).get(event_type, (0, 0)),
                                     (len(expected_ids), len(expected_ids)))
                    if event_type.startswith('custom.'):
                        contexts = app._plot_event_tooltip_contexts(
                            plotted_record, MAIN.plot_event_entries(plotted_record), scoped=False, active_period='')
                        self.assertEqual([contexts[id(e)][1] for e in plotted_record['events']],
                                         list(range(1, len(expected_ids) + 1)))
                    unchanged = self.open_dialog(app, method)
                    save(unchanged)
                    self.assertEqual(backend.load_core_data()['animals']['Subject']['events'], saved['events'])
                    if deleted == 'middle':
                        # A genuinely new row after deletion uses the next count,
                        # not a sequence number retained from removed history.
                        addition = self.open_dialog(app, method)
                        rows_before = self.event_editor_rows_247(addition)
                        combo = next(iter(rows_before.values()))[1]
                        event_scroll = combo.parentWidget().parentWidget().parentWidget()
                        add_button = next(b for b in event_scroll.findChildren(QPushButton)
                                          if b.text().casefold() in ('new event', 'add events'))
                        add_button.click()
                        self.qt.processEvents()
                        new_row = self.event_editor_rows_247(addition)['']
                        new_row[0].setText('03.01.2025')
                        self.assertGreaterEqual(new_row[1].findData(event_type), 0)
                        new_row[1].setCurrentIndex(new_row[1].findData(event_type))
                        save(addition)
                        added_record = app.animals['Subject']
                        new_id = next(e['event_id'] for e in added_record['events']
                                      if e['event_id'] not in expected_ids)
                        index, _ = repository.get('Subject', added_record, str, date_format='%d.%m.%Y')
                        self.assertEqual(index.event_counts_through(date(2025, 1, 3))[event_type], (3, 3))
                        if event_type.startswith('custom.'):
                            contexts = app._plot_event_tooltip_contexts(
                                added_record, MAIN.plot_event_entries(added_record), scoped=False, active_period='')
                            added = next(e for e in added_record['events'] if e['event_id'] == new_id)
                            self.assertEqual(contexts[id(added)][1], 3)
                        remove_addition = self.open_dialog(app, method)
                        self.event_editor_rows_247(remove_addition)[new_id][2].click()
                        self.qt.processEvents()
                        save(remove_addition)
                        self.assertEqual(backend.load_core_data()['animals']['Subject']['events'], saved['events'])
                backend.close()
                reopened = ProgTrackBackend(runtime_paths(root, application_root=root),
                                            acquire_process_lock=False)
                self.addCleanup(reopened.close)
                self.assertEqual(reopened.load_core_data()['animals']['Subject']['events'], [])

    def test_247_shared_identity_guard_covers_every_event_editor_adapter(self):
        adapters = (
            ("_dlg_offspring", "special_measurement"),
            ("_dlg_zuchttier", "pregnancy"),
            ("_dlg_versuchstier", "pgf"),
            ("_dlg_basic_animal_role", "custom.audit_event"),
            ("_dlg_female_animal", "pgf"),
        )
        for method_name, event_type in adapters:
            with self.subTest(adapter=method_name, event_type=event_type):
                source = inspect.getsource(getattr(MAIN.ProgTrackApp, method_name))
                self.assertIn("event_payload_from_editor_row(", source)
                self.assertIn("append_event_payload_once(", source)

        for event_type in ("pgf", "custom.audit_event"):
            with self.subTest(event_type=event_type):
                event_date = datetime(2025, 1, 1)
                existing_source = {
                    "event_id": f"persisted-{event_type}",
                    "event_type": event_type,
                    "datum": event_date,
                    "recorded_role": "egg_cell_donor",
                }
                first_row = SimpleNamespace(
                    _progtrack_source_event=copy.deepcopy(existing_source)
                )
                duplicate_reference = SimpleNamespace(
                    _progtrack_source_event=copy.deepcopy(existing_source)
                )
                payload = MAIN.event_payload_from_editor_row(
                    first_row, event_type, event_date, recording_role="egg_cell_donor"
                )
                duplicate_payload = MAIN.event_payload_from_editor_row(
                    duplicate_reference,
                    event_type,
                    event_date,
                    recording_role="egg_cell_donor",
                )
                collected, seen = [], set()
                self.assertTrue(MAIN.append_event_payload_once(
                    collected, seen, first_row, payload
                ))
                self.assertFalse(MAIN.append_event_payload_once(
                    collected, seen, duplicate_reference, duplicate_payload
                ))

                # Equal type/date/role values are not identity: separate new
                # editor rows remain separate, while one row appended twice is
                # still emitted once.
                distinct_rows = [SimpleNamespace(), SimpleNamespace()]
                for row in distinct_rows:
                    new_payload = MAIN.event_payload_from_editor_row(
                        row, event_type, event_date, recording_role="egg_cell_donor"
                    )
                    self.assertTrue(MAIN.append_event_payload_once(
                        collected, seen, row, new_payload
                    ))
                    self.assertFalse(MAIN.append_event_payload_once(
                        collected, seen, row, new_payload
                    ))
                self.assertEqual(len(collected), 3)
                self.assertEqual(
                    [item["event_type"] for item in collected],
                    [event_type, event_type, event_type],
                )

                legacy_source = {"event_type": event_type, "datum": event_date}
                first_projection = MAIN.event_source_for_editor_row(legacy_source)
                second_projection = MAIN.event_source_for_editor_row(legacy_source)
                self.assertTrue(legacy_source["event_id"])
                self.assertEqual(first_projection["event_id"], second_projection["event_id"])

    def test_247_female_editor_keeps_distinct_same_day_same_type_events_and_ids(self):
        from Plugins.core.backend import ProgTrackBackend
        from Plugins.core.project_periods import (
            current_event_counts,
            ensure_project_periods,
            event_counts_by_period,
        )

        with tempfile.TemporaryDirectory() as directory:
            temporary_root = Path(directory)
            backend = ProgTrackBackend(
                runtime_paths(temporary_root, application_root=temporary_root),
                acquire_process_lock=False,
            )
            self.addCleanup(backend.close)
            app = self.host()
            app.backend = backend
            app._role_dialog_blocks = lambda *_a, **_k: ["weight", "reference_weight"]
            app._project_names_for_species = lambda _species: ["Current Project"]
            app._retain_selection_after_animal_save = lambda *_a, **_k: None
            app._run_post_persistence_sync = lambda *_a, **_k: None
            app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(
                app, MAIN.ProgTrackApp
            )

            record = self.record(MAIN.Role.SPENDER.value)
            record.update({
                "project": "Current Project",
                "project_entry_date": "01.02.2025",
                "project_history": [{
                    "project": "Former Project",
                    "entry_date": "01.01.2025",
                    "leave_date": "01.02.2025",
                }],
                "max_pgf": 10,
                "events": [{
                    "event_id": "former-pgf",
                    "event_type": "pgf",
                    "datum": datetime(2025, 2, 1, 9, 0),
                    "recorded_role": MAIN.Role.SPENDER.value,
                    "project_period_id": "",
                    "project_period_project": "Former Project",
                    "custom_metadata": "keep-existing-payload",
                }],
            })
            periods = ensure_project_periods(record)
            former_period = next(period for period in periods if not period["current"])
            current_period = next(period for period in periods if period["current"])
            record["events"][0]["project_period_id"] = former_period["period_id"]
            app.animals["Subject"] = record

            def add_two_same_day_pgf_rows() -> None:
                dialog = self.open_dialog(app, "_dlg_female_animal")
                events_label = app.messages.get("dialog.tab.events", "Events")
                event_tab = next(
                    tab.widget(index)
                    for tab in dialog.findChildren(QTabWidget)
                    for index in range(tab.count())
                    if tab.tabText(index) == events_label
                )
                event_scroll = event_tab.findChild(QScrollArea)
                self.assertIsNotNone(event_scroll)
                add_label = app.messages["button.new_item"].format(
                    title=app.messages.get("dialog.female_animal.tab.events", "Events")
                )
                add_button = next(
                    button for button in event_scroll.findChildren(QPushButton)
                    if button.text() == add_label
                )
                add_button.click()
                add_button.click()
                add_button.click()
                self.qt.processEvents()

                event_combos = [
                    combo for combo in event_scroll.findChildren(QComboBox)
                    if hasattr(combo, "_progtrack_source_event")
                ]
                date_fields = event_scroll.findChildren(QLineEdit)
                self.assertEqual(len(event_combos), 4)
                self.assertEqual(len(date_fields), 4)
                for combo in event_combos:
                    pgf_index = combo.findData("pgf")
                    self.assertGreaterEqual(pgf_index, 0)
                    combo.setCurrentIndex(pgf_index)
                for field in date_fields:
                    field.setText("01.02.2025")
                # Simulate the same persisted event being represented twice in
                # the editor. Its repeated event_id must be emitted once, but
                # the two independent same-value rows must remain.
                event_combos[1]._progtrack_source_event = copy.deepcopy(
                    event_combos[0]._progtrack_source_event
                )
                save_button = next(
                    button for button in dialog.findChildren(QPushButton)
                    if button.text() == app.messages["button.save"]
                )
                save_button.click()
                self.qt.processEvents()

            add_two_same_day_pgf_rows()
            saved_events = app.animals["Subject"]["events"]
            self.assertEqual(len(saved_events), 3)
            self.assertEqual(sum(event["event_type"] == "pgf" for event in saved_events), 3)
            self.assertEqual(len({event["event_id"] for event in saved_events}), 3)
            former_event = next(event for event in saved_events if event["event_id"] == "former-pgf")
            self.assertEqual(former_event["project_period_id"], former_period["period_id"])
            self.assertEqual(former_event["recorded_role"], MAIN.Role.SPENDER.value)
            self.assertEqual(former_event["custom_metadata"], "keep-existing-payload")
            new_events = [event for event in saved_events if event["event_id"] != "former-pgf"]
            self.assertTrue(all(event["recorded_role"] == MAIN.Role.SPENDER.value for event in new_events))
            self.assertEqual(
                {event["project_period_id"] for event in new_events},
                {current_period["period_id"]},
            )
            self.assertEqual(current_event_counts(app.animals["Subject"]), {"pgf": 2})
            self.assertEqual(
                event_counts_by_period(app.animals["Subject"])[former_period["period_id"]],
                {"pgf": 1},
            )
            self.assertEqual(
                event_counts_by_period(app.animals["Subject"])[current_period["period_id"]],
                {"pgf": 2},
            )

            event_ids_before_noop_save = {event["event_id"] for event in saved_events}
            second_dialog = self.open_dialog(app, "_dlg_female_animal")
            second_save = next(
                button for button in second_dialog.findChildren(QPushButton)
                if button.text() == app.messages["button.save"]
            )
            second_save.click()
            self.qt.processEvents()
            round_trip = backend.load_core_data()["animals"]["Subject"]["events"]
            self.assertEqual(len(round_trip), 3)
            self.assertEqual({event["event_id"] for event in round_trip}, event_ids_before_noop_save)
            self.assertEqual(len({event["project_period_id"] for event in round_trip}), 2)
            self.assertEqual(
                sum(event["project_period_id"] == current_period["period_id"] for event in round_trip),
                2,
            )

    def test_271_flow_detail_dialogs_never_show_unowned_content(self):
        from Plugins.Flow_Track.flow_track_widget import FlowTrackWidget
        for method, efficiency_method, per_key, total_key in (
            ('_show_egg_donor_dialog', '_calculate_egg_donor_efficiency', 'per_surgery', 'total_gv_isolated'),
            ('_show_sperm_donor_dialog', '_calculate_sperm_donor_efficiency', 'per_donation', 'total_ivm_m2_inseminated'),
            ('_show_surrogate_dialog', '_calculate_surrogate_efficiency', 'per_transfer', 'total_embryos'),
            ('_edit_embryo', '_unused_efficiency', '', ''),
        ):
            with self.subTest(dialog=method):
                host = self.host()
                flow = FlowTrackWidget.__new__(FlowTrackWidget)
                flow.parent_app = SimpleNamespace(QtWidgets=MAIN.QtWidgets, QtCore=MAIN.QtCore,
                                                  Role=MAIN.Role, animals={})
                flow.widget = QWidget(host)
                flow.messages, flow.settings, flow.manual_data = {}, {}, {}
                flow._animal_export_name = lambda name: name
                flow._can = lambda *_a: True
                flow.manual_data = {'transfers_by_id': {'Subject': {'embryos': [
                    {'embryo_id': 'embryo', 'cryopreserved': True, 'freeze_date': '2026-01-01'}]}}}
                setattr(flow, efficiency_method, lambda _name, per=per_key, total=total_key: {
                    per: [], 'total': {total: 1}})
                flow._save_settings = lambda: None
                record = {'events': [
                    {'event_type': 'surgery', 'datum': datetime(2026, 1, 1)},
                    {'event_type': 'embryo_transfer', 'datum': datetime(2026, 1, 2)},
                ], 'sperm': [{'datum': datetime(2026, 1, 1)}]}
                unexpected, dialogs = [], []

                class Observer(MAIN.QtCore.QObject):
                    def eventFilter(self, widget, event):
                        if event.type() == QEvent.Type.Show and isinstance(widget, QWidget) and widget.isWindow():
                            unexpected.append((type(widget).__name__, widget.windowTitle()))
                        return False

                observer = Observer()
                self.qt.installEventFilter(observer)
                try:
                    with patch.object(QDialog, 'exec', lambda dlg: dialogs.append(dlg) or 0):
                        getattr(flow, method)('Subject', 'embryo' if method == '_edit_embryo' else record)
                    self.hosts.extend(dialogs)
                    self.qt.processEvents()
                    self.assertEqual(len(dialogs), 1)
                    self.assertEqual(unexpected, [])
                    dialogs[0].show()
                    self.qt.processEvents()
                    self.assertTrue(dialogs[0].isVisible())
                    self.assertEqual(unexpected, [(type(dialogs[0]).__name__, dialogs[0].windowTitle())])
                    dialogs[0].reject()
                    self.qt.processEvents()
                    self.assertFalse(dialogs[0].isVisible())
                finally:
                    self.qt.removeEventFilter(observer)

    def test_271_job_reload_disposes_old_controls_without_orphans(self):
        from PyQt6 import sip
        from Plugins.Master_Track.dialogs import EditJobsDialog
        host = self.host()
        dialog = EditJobsDialog(host, {}, SimpleNamespace(
            app=SimpleNamespace(lang='en'), get_job_timeout=lambda _name: 30))
        dialog.show()
        self.qt.processEvents()
        count = dialog._scroll_vbox.count()
        for _ in range(3):
            old = list(dialog._perm_checks.values())
            dialog._load_job(dialog.job_combo.currentText())
            self.assertTrue(all(widget.parentWidget() is not None for widget in old))
            self.qt.sendPostedEvents(None, QEvent.Type.DeferredDelete)
            self.qt.processEvents()
            self.assertTrue(all(sip.isdeleted(widget) for widget in old))
            self.assertEqual(dialog._scroll_vbox.count(), count)

    def test_271_manage_users_button_never_shows_before_owned_dialog(self):
        from Plugins.Master_Track.dialogs import ManageUsersDialog
        for allowed in (True, False):
            with self.subTest(can_create=allowed):
                host = self.host()
                shown = []
                class Observer(MAIN.QtCore.QObject):
                    def eventFilter(self, widget, event):
                        if event.type() == QEvent.Type.Show and isinstance(widget, QWidget) and widget.isWindow():
                            shown.append(widget)
                        return False
                observer = Observer()
                self.qt.installEventFilter(observer)
                try:
                    dialog = ManageUsersDialog(host, {}, SimpleNamespace(users=[]), 'fixture',
                                               can_create_users=allowed)
                    self.assertEqual(shown, [])
                    dialog.show()
                    self.qt.processEvents()
                    self.assertEqual(shown, [dialog])
                    button = next(b for b in dialog.findChildren(QPushButton) if b.text() == 'Create User')
                    self.assertEqual(button.isVisible(), allowed)
                    dialog.reject()
                finally:
                    self.qt.removeEventFilter(observer)

    def _320_sidebar_host(self):
        app = self._307_sqlite_host(seeded=True)
        self._316_admin(app)
        self._306_sidebar_shell(app)
        for name in ('btn_edit_animal', 'btn_load_sperm', 'btn_load_urine'):
            setattr(app, name, QPushButton(app))
        app._on_category_selected = MAIN.ProgTrackApp._on_category_selected.__get__(app)
        app._capture_death_lifecycle = MAIN.ProgTrackApp._capture_death_lifecycle.__get__(app)
        app.lst.itemSelectionChanged.connect(app._on_select)
        app.lst.itemClicked.connect(app._on_animal_item_clicked)
        app.btn_edit_animal.clicked.connect(app._on_edit_animal_from_all_tab)
        app._rebuild_category_tabs()
        app.category_tab.setCurrentIndex(app._all_category_tab_index())
        app._on_category_selected(app.category_tab.currentIndex())
        central = QWidget(app)
        layout = QVBoxLayout(central)
        layout.addWidget(app.lst)
        layout.addWidget(app.btn_edit_animal)
        app.setCentralWidget(central)
        app.resize(800, 600)
        app.show()
        self.qt.processEvents()
        names = app.backend.records.get('seed', 'shire-colony')['named_animals']
        return app, names['Adamanta'], names['Frodo']

    def _320_click(self, app, key, *, ctrl=False, drain=True):
        from PyQt6.QtTest import QTest
        item = next(app.lst.item(i) for i in range(app.lst.count())
                    if app.lst.item(i).data(Qt.ItemDataRole.UserRole) == key)
        app.lst.scrollToItem(item)
        QTest.mouseClick(app.lst.viewport(), Qt.MouseButton.LeftButton,
                         Qt.KeyboardModifier.ControlModifier if ctrl else Qt.KeyboardModifier.NoModifier,
                         app.lst.visualItemRect(item).center())
        if drain:
            self.qt.processEvents()

    def test_320_real_click_refresh_retains_current_row_and_edit_target(self):
        app, first, last = self._320_sidebar_host()
        self._320_click(app, first)
        self._320_click(app, last, ctrl=True)
        expected_plot = list(app._plot_selection_order)
        self.assertEqual(app._selected_animal_for_edit(), last)
        self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)
        app._refresh_list(update_tab_visibility=True)
        self.qt.processEvents()
        self.assertIsNotNone(app.lst.currentItem(), 'Save rebuild lost the current row')
        self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)
        app.lst.clearFocus()
        app.lst.setFocus()
        self.qt.processEvents()
        self.assertEqual(app._selected_animal_for_edit(), last)
        self.assertEqual(app._plot_selection_order, expected_plot)
        self.assertEqual(set(app.selected_animals), {first, last})

    def _320_run_editor(self, app, key, action, *, primary=False, explicit=False):
        """Real modal loop, dispatcher, lease, Save and shared finalization."""
        execute = app._exec_animal_editor_dialog
        callback_errors = []
        dialog_results = []
        def run(dialog):
            self.hosts.append(dialog)
            def interact():
                try:
                    if action == 'cancel':
                        dialog.reject()
                        return
                    if action in ('change', 'invalid', 'failure'):
                        page = self.weight_page(app, dialog)
                        old = app.animals[key]['gewicht'][0]['wert']
                        value = next(w for w in page.findChildren(QLineEdit)
                                     if w.text() == format_measurement(old))
                        value.setText('invalid' if action == 'invalid' else format_measurement(old + .1))
                    if action == 'decline':
                        next(w for w in dialog.findChildren(QCheckBox)
                             if w.text() == app.messages['checkbox.in_experiment']).setChecked(False)
                        next(b for b in dialog.findChildren(QPushButton)
                             if b.text() == app.messages['dialog.versuchstier.button.new_event']).click()
                        date, combo, _delete = self.event_editor_rows_247(dialog)['']
                        date.setText('01.10.2026')
                        combo.setCurrentIndex(combo.findData('measurement'))
                    next(b for b in dialog.findChildren(QPushButton)
                         if b.text() == app.messages['button.save']).click()
                except BaseException:
                    callback_errors.append(traceback.format_exc())
                finally:
                    dialog_results.append(dialog.result())
                    if dialog.isVisible():
                        dialog.reject()  # Failed/invalid attempt: no hanging modal.
            MAIN.QTimer.singleShot(0, interact)
            return execute(dialog)
        with patch.object(app, '_exec_animal_editor_dialog', side_effect=run), \
             patch.object(app, '_open_animal_editor_with_lease',
                          wraps=app._open_animal_editor_with_lease) as opened, \
             patch('sys.excepthook', lambda typ, exc, tb: callback_errors.append(
                 ''.join(traceback.format_exception(typ, exc, tb)))):
            if explicit:
                app._open_animal_editor_with_lease(
                    key, app._dialog_for_role_value(app._animal_role_value(app.animals[key])), False)
            else:
                (app.btn_edit if primary else app.btn_edit_animal).click()
        for _ in range(3):
            self.qt.processEvents()
        self.assertEqual(opened.call_count, 1)
        self.assertEqual(opened.call_args.args[0], key)
        return dialog_results, callback_errors

    def test_320_actual_modal_saves_cancel_and_next_dispatch_preserve_last_subject(self):
        app, first, last = self._320_sidebar_host()
        for multi in (False, True):
            self._320_click(app, first)
            self._320_click(app, last, ctrl=multi)
            plot_order = list(app._plot_selection_order)
            expected = {first, last} if multi else {last}
            for action in ('cancel', 'noop', 'change', 'noop'):
                with self.subTest(multi=multi, action=action):
                    result, errors = self._320_run_editor(app, last, action)
                    self.assertFalse(errors or app.errors, (errors, app.errors))
                    if action != 'cancel':
                        self.assertEqual(result, [QDialog.DialogCode.Accepted])
                    app.lst.setFocus()
                    self.qt.processEvents()
                    self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)
                    self.assertEqual(app._selected_animal_for_edit(), last)
                    self.assertEqual(set(app.selected_animals), expected)
                    self.assertEqual(app._plot_selection_order, plot_order)
            # Next actual Edit dispatch must still open Frodo, not just resolve a helper.
            self.assertFalse(self._320_run_editor(app, last, 'cancel')[1])
        with app.backend.adapter.transaction() as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM entity_leases').fetchone()[0], 0)

    def test_320_category_custom_and_explicit_other_editor_keep_interaction_order(self):
        app, first, last = self._320_sidebar_host()
        self._320_click(app, first)
        self._320_click(app, last, ctrl=True)
        # Saving a selected but not interaction-last animal must not promote it.
        self.assertFalse(self._320_run_editor(app, first, 'noop', explicit=True)[1])
        self.assertEqual(app._selected_animal_for_edit(), last)
        self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)
        # Nor should an explicit unselected editor add itself to plots/selection.
        self._320_click(app, last)
        self.assertFalse(self._320_run_editor(app, first, 'noop', explicit=True)[1])
        self.assertEqual(app.selected_animals, [last])
        self.assertEqual(app._selected_animal_for_edit(), last)
        # Exercise the primary category-specific Edit dispatcher.
        app.category_tab.setCurrentIndex(5)
        app._on_category_selected(5)
        self.assertFalse(self._320_run_editor(app, last, 'noop', primary=True)[1])
        self.assertEqual(app._selected_animal_for_edit(), last)
        # A configured-role fixture uses the real basic-role editor and Save.
        raw = app.backend.load_core_data()
        raw['animals'][last]['rolle'] = 'custom.selection_test'
        app.backend.save_core_data(raw)
        MAIN.ProgTrackApp._load_persistence(app)
        app._role_dialog_blocks = lambda *_a, **_k: ['weight', 'reference_weight']
        app.category_tab.setCurrentIndex(app._all_category_tab_index())
        app._on_category_selected(app.category_tab.currentIndex())
        self._320_click(app, last)
        result, errors = self._320_run_editor(app, last, 'change')
        self.assertFalse(errors or app.errors, (errors, app.errors))
        self.assertEqual(result, [QDialog.DialogCode.Accepted])
        self.assertEqual(app._selected_animal_for_edit(), last)
        self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)

    def test_320_hidden_deselected_deleted_and_session_state_are_not_resurrected(self):
        app, first, last = self._320_sidebar_host()
        self._320_click(app, first)
        self._320_click(app, last, ctrl=True, drain=False)
        # Interaction is settled before any queued callback or immediate Edit.
        self.assertEqual(app._selected_animal_for_edit(), last)
        app._animal_name_filter = app.animals[first]['name']
        app._refresh_list()
        self.assertEqual(set(app.selected_animals), {first, last})
        self.assertIsNone(app.lst.currentItem(), 'do not invent a row for a hidden target')
        self.assertEqual(app._selected_animal_for_edit(), last)
        app._animal_name_filter = ''
        app._refresh_list()
        self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), last)
        # Revoked visibility, unlike a name filter, invalidates the selection.
        with patch.object(app, '_animal_visible_to_current_user',
                          side_effect=lambda record: record.get('ipid') != last):
            app._refresh_list()
            self.assertEqual(app.selected_animals, [first])
            self.assertEqual(app._selected_animal_for_edit(), first)
            self.assertEqual(app.lst.currentItem().data(Qt.ItemDataRole.UserRole), first)
        app._refresh_list()
        self._320_click(app, last, ctrl=True)
        self.assertEqual(app._selected_animal_for_edit(), last)
        self._320_click(app, last, ctrl=True)
        self.assertEqual(app.selected_animals, [first])
        app._refresh_list()
        self.assertEqual(app._selected_animal_for_edit(), first)
        # Remove an active subject without a persistence command: projection-only negative fixture.
        app.animals.pop(first)
        app._refresh_list()
        self.assertIsNone(app._selected_animal_for_edit())
        self.assertIsNone(app.lst.currentItem())
        self._320_click(app, last, drain=False)
        app._clear_session_animal_selection()
        self.qt.processEvents()
        self.assertEqual(app._edit_selection_order, [])
        self.assertEqual(app.selected_animals, [])
        self.assertIsNone(app.lst.currentItem())

    def test_320_invalid_declined_and_failed_save_keep_selection_and_release_lease(self):
        app, first, last = self._320_sidebar_host()
        app._is_projects_track_active = lambda: True
        self._320_click(app, first)
        self._320_click(app, last, ctrl=True)
        baseline = self._316_sql_state(app)
        audit = []
        app._master_audit = lambda *args: audit.append(args)
        for action in ('invalid', 'failure', 'decline'):
            with self.subTest(action=action):
                save = app.backend.save_core_data
                def persist(*args, **kwargs):
                    if action == 'failure':
                        raise OSError('test-320 write failure')
                    return save(*args, **kwargs)
                message = app._show_message_raw
                def answer(*args, **kwargs):
                    if action == 'decline' and args[0] == app.messages['confirm.event_save.title']:
                        return MAIN.QMessageBox.StandardButton.No
                    return message(*args, **kwargs)
                with patch.object(app.backend, 'save_core_data', side_effect=persist) as write, \
                     patch.object(app, '_show_message_raw', side_effect=answer):
                    result, errors = self._320_run_editor(app, last, action)
                self.assertEqual(result, [QDialog.DialogCode.Rejected])
                self.assertEqual(write.call_count, 1 if action == 'failure' else 0)
                if action == 'failure':
                    self.assertTrue(errors and 'test-320 write failure' in errors[0], errors)
                else:
                    self.assertFalse(errors)
                if action == 'invalid':
                    self.assertTrue(app.errors, 'invalid input must be reported')
                app.errors.clear()
                self.assertEqual(self._316_sql_state(app), baseline)
                self.assertFalse(audit)
                self.assertEqual(app._selected_animal_for_edit(), last)
                self.assertEqual(app._edit_selection_order, [first, last])
                self.assertEqual(app._plot_selection_order, [first, last])
                with app.backend.adapter.transaction() as connection:
                    self.assertEqual(connection.execute('SELECT COUNT(*) FROM entity_leases').fetchone()[0], 0)
                # Restore the unsaved test candidate before the next independent negative.
                MAIN.ProgTrackApp._load_persistence(app)

    def test_320_modal_completion_cannot_restore_a_previous_session_target(self):
        app, first, last = self._320_sidebar_host()
        self._320_click(app, last)
        for change_actor in (False, True):
            with self.subTest(change_actor=change_actor):
                dialog = QDialog(app)
                self.hosts.append(dialog)
                def finish():
                    dialog._progtrack_successful_animal_edit = (last, last)
                    app._prepare_actor_transition()
                    if change_actor:
                        app.master_track._current_username = 'Vet'
                    dialog.accept()
                MAIN.QTimer.singleShot(0, finish)
                with patch.object(app, '_complete_successful_animal_edit',
                                  wraps=app._complete_successful_animal_edit) as complete:
                    self.assertFalse(app._exec_animal_editor_dialog(dialog))
                    complete.assert_not_called()
                self.qt.processEvents()
                self.assertEqual(app.selected_animals, [])
                self.assertEqual(app._edit_selection_order, [])
                self.assertIsNone(app.lst.currentItem())
                self.assertFalse(hasattr(dialog, '_progtrack_successful_animal_edit'))

    def test_270_sidebar_rebuild_blocks_intermediate_selection_signals(self):
        app = self.host()
        app.lst = QListWidget()
        item = QListWidgetItem('Subject')
        item.setData(Qt.ItemDataRole.UserRole, 'Subject')
        app.lst.addItem(item)
        item.setSelected(True)
        app.selected_animals = ['Subject']
        app._edit_selection_order = ['Subject']
        received = []
        app.lst.itemSelectionChanged.connect(lambda: received.append(True))
        def rebuild(*_args):
            app.lst.clear()
            item = QListWidgetItem('Subject')
            item.setData(Qt.ItemDataRole.UserRole, 'Subject')
            app.lst.addItem(item)
            item.setSelected(True)
        app._populate_animal_list = rebuild
        MAIN.ProgTrackApp._refresh_list(app)
        self.assertEqual(received, [])
        self.assertFalse(app.lst.signalsBlocked())
        self.assertEqual(app._edit_selection_order, ['Subject'])
        self.assertTrue(app.lst.item(0).isSelected())

    def test_270_real_sidebar_refresh_preserves_order_and_permission_state(self):
        app = self.host()
        app.category_tab = QTabBar(app)
        app._category_custom_role_values = []
        app._custom_category_roles = lambda: []
        app.cmb_arch = QComboBox(app)
        app._sync_heritage_category_tab_state = lambda: None
        app._animal_visible_to_current_user = lambda _rec: True
        app._project_visibility_scope = lambda **_kwargs: (True, set())
        app._get_status = lambda _name: ''
        app._status_icon_ids = lambda *_args: []
        app._status_text_without_icon_tokens = lambda _status: ''
        app.animals = {'First': self.record(), 'Second': self.record()}
        app._rebuild_category_tabs()
        app.category_tab.setCurrentIndex(app._all_category_tab_index())
        app.selected_animals = ['First', 'Second']
        app._plot_selection_order = ['First', 'Second']
        app._edit_selection_order = ['Second', 'First']
        for name in ('btn_restore', 'btn_delete', 'btn_edit', 'btn_edit_animal'):
            setattr(app, name, QPushButton(app))
        app._can_use_primary_edit_button = lambda: True
        app._master_can = lambda action: True
        app._populate_animal_list = lambda *args: MAIN.ProgTrackApp._populate_animal_list(app, *args)
        MAIN.ProgTrackApp._refresh_list(app, update_tab_visibility=True)
        self.assertEqual(app.selected_animals, ['First', 'Second'])
        self.assertEqual(app._edit_selection_order, ['Second', 'First'])
        self.assertEqual(app._selected_animal_for_edit(), 'First')
        self.assertTrue(app.btn_edit.isEnabled())
        self.assertEqual(len(app.lst.selectedItems()), 2)
        # The All-tab primary action is Edit Role, while the adjacent action
        # opens Edit Animal. A Vet may have only the latter permission.
        app._master_can = lambda action: action == 'core.edit_animal_core'
        app._can_use_primary_edit_button = lambda: False
        MAIN.ProgTrackApp._refresh_list(app)
        self.assertFalse(app.btn_edit.isEnabled())
        self.assertTrue(app.btn_edit_animal.isEnabled())

    def test_270_master_button_recalculation_keeps_edit_permissions_separate(self):
        for tab_index in range(7):
            for master_available in (False, True):
                for edit_role, edit_animal in ((False, True), (True, False)):
                    with self.subTest(tab_index=tab_index,
                                      master_available=master_available,
                                  edit_role=edit_role, edit_animal=edit_animal):
                        app = self.host()
                        permissions = {
                            'core.edit_animal_role': edit_role,
                            'core.edit_animal_core': edit_animal,
                        }
                        app.authorization = SimpleNamespace(
                            can=lambda action, values=permissions: values.get(action, False),
                            trusted_local=False,
                        )
                        app._master_can = lambda action, values=permissions: values.get(action, False)
                        app.master_track = (
                            SimpleNamespace(can=lambda action, values=permissions: values.get(action, False))
                            if master_available else None
                        )
                        app.category_tab = SimpleNamespace(
                            currentIndex=lambda index=tab_index: index
                        )
                        app._all_category_tab_index = lambda: 0
                        app.animals = {'Subject': self.record()}
                        app.selected_animals = ['Subject']
                        app.btn_edit = QPushButton(app)
                        app.btn_edit_animal = QPushButton(app)

                        MAIN.ProgTrackApp._apply_master_button_states(app)

                        expected_primary = edit_role if tab_index == 0 else edit_animal
                        self.assertEqual(app.btn_edit.isEnabled(), expected_primary)
                        self.assertEqual(app.btn_edit_animal.isEnabled(), edit_animal)

    def test_270_female_save_refresh_keeps_all_tab_edit_action_and_target(self):
        with tempfile.TemporaryDirectory() as directory:
            temporary_root = Path(directory)
            backend = ProgTrackBackend(
                runtime_paths(temporary_root, application_root=temporary_root),
                acquire_process_lock=False,
            )
            self.addCleanup(backend.close)
            app = self.host()
            app.backend = backend
            app._master_can = lambda action: action != 'core.edit_animal_role'
            app._role_dialog_blocks = lambda *_args, **_kwargs: ['weight', 'reference_weight']
            app._project_names_for_species = lambda _species: ['Current Project']
            app._run_post_persistence_sync = lambda *_args, **_kwargs: None
            app._save_persistence = MAIN.ProgTrackApp._save_persistence.__get__(
                app, MAIN.ProgTrackApp
            )
            app._refresh_list = lambda update_tab_visibility=False, force_heritage_visible=False: (
                MAIN.ProgTrackApp._refresh_list(
                    app, update_tab_visibility, force_heritage_visible
                )
            )
            app._animal_visible_to_current_user = lambda _record: True
            app._project_visibility_scope = lambda **_kwargs: (True, set())
            app._animal_matches_sex_filter = lambda _record: True
            app._current_animal_name_filter_text = lambda: ''
            app._get_status = lambda _name: ''
            app._status_icon_ids = lambda *_args: []
            app._status_text_without_icon_tokens = lambda _status: ''
            app._sync_heritage_category_tab_state = lambda: None
            app._sync_animal_list_row_widths = lambda: None
            app._schedule_animal_list_row_width_sync = lambda: None
            app._open_animal_editor_with_lease = lambda *args: opened.append(args)
            app._dialog_for_role_value = lambda _role: 'female-editor'
            app.category_tab = QTabBar(app)
            app._category_custom_role_values = []
            app._custom_category_roles = lambda: []
            app.cmb_arch = QComboBox(app)
            app.btn_new = QPushButton(app)
            app.btn_archive = QPushButton(app)
            app.btn_load_blood = QPushButton(app)
            app.btn_load_urine = QPushButton(app)
            app.btn_load_weights = QPushButton(app)
            app.btn_load_sperm = QPushButton(app)
            app.btn_edit = QPushButton(app)
            app.btn_edit_animal = QPushButton(app)
            app.btn_restore = QPushButton(app)
            app.btn_delete = QPushButton(app)
            app.btn_edit_animal.clicked.connect(app._on_edit_animal_from_all_tab)
            app._rebuild_category_tabs()
            all_index = app._all_category_tab_index()
            app.category_tab.setCurrentIndex(all_index)
            app._apply_sidebar_button_visibility_for_category(all_index)

            rachel = self.record(MAIN.Role.AMME.value)
            rachel.update({
                'name': 'Rachel', 'birth_date': '14.02.2017',
                'species': 'Callithrix jacchus', 'origin': 'DPZ',
                'sex': 'Female', 'id': 'RACHEL-TEST-ID',
                'project': 'Current Project', 'project_entry_date': '01.01.2025',
            })
            app.animals = {'Rachel': rachel}
            app.selected_animals = ['Rachel']
            app._plot_selection_order = ['Rachel']
            app._edit_selection_order = ['Rachel']
            app._refresh_list(update_tab_visibility=True)
            self.qt.processEvents()

            self.assertFalse(app.btn_edit.isEnabled(), 'Veti cannot edit animal roles')
            self.assertTrue(app.btn_edit_animal.isEnabled(), 'Veti can edit animal data')
            dialog = self.open_dialog(app, '_dlg_female_animal', name='Rachel')
            weight_page = self.weight_page(app, dialog)
            weight_value = next(
                field for field in weight_page.findChildren(QLineEdit)
                if field.text() == '27.1'
            )
            weight_value.setText('440.1')
            save_button = next(
                button for button in dialog.findChildren(QPushButton)
                if button.text() == app.messages['button.save']
            )
            save_button.click()
            self.qt.processEvents()
            self.qt.processEvents()

            self.assertEqual(app.selected_animals, ['Rachel'])
            self.assertEqual(app.animals['Rachel']['gewicht'][0]['wert'], 440.1)
            self.assertEqual(
                [item.data(Qt.ItemDataRole.UserRole) for item in app.lst.selectedItems()],
                ['Rachel'],
            )
            self.assertFalse(app.btn_edit.isEnabled())
            self.assertTrue(app.btn_edit_animal.isEnabled())
            opened = []
            app.btn_edit_animal.click()
            self.assertEqual(opened[0][0], 'Rachel')

    def test_270_saved_animal_retains_other_subjects_and_last_interaction(self):
        app = self.host()
        app.animals = {'First': self.record(), 'Renamed': self.record()}
        app.selected_animals = ['First', 'Old']
        app._plot_selection_order = ['First', 'Old']
        app._edit_selection_order = ['Old', 'First']
        app._retain_selection_after_animal_save('Old', 'Renamed')
        self.assertEqual(app.selected_animals, ['First', 'Renamed'])
        self.assertEqual(app._plot_selection_order, ['First', 'Renamed'])
        self.assertEqual(app._edit_selection_order, ['Renamed', 'First'])
        self.assertEqual(app._selected_animal_for_edit(), 'First')

    def test_272_session_reset_clears_only_runtime_selection(self):
        app = self.host()
        app.lst = QListWidget()
        app.lst.addItem('Subject')
        app.lst.item(0).setSelected(True)
        attrs = ('selected_animals', '_selected_archived', '_selected_heritage_only',
                 '_plot_selection_order', '_edit_selection_order', 'last_plotted_animals')
        for attr in attrs:
            setattr(app, attr, ['Subject'])
        maps = {'key': {'Subject': [1, 2]}}
        app.backend.records.put('configuration', 'layout-witness', maps)
        app._clear_session_animal_selection()
        for attr in attrs:
            self.assertEqual(getattr(app, attr), [])
        self.assertEqual(app.lst.selectedItems(), [])
        self.assertEqual(app.backend.records.get('configuration', 'layout-witness'), maps)

    def test_272_logout_and_login_success_failure_boundaries(self):
        for path in ('logout', 'login-success', 'login-cancelled'):
            with self.subTest(path=path):
                app = self.host()
                app.selected_animals = ['Subject']
                app._selected_heritage_only = ['Dummy']
                app.lst.addItem('Subject')
                app.lst.item(0).setSelected(True)
                received = []
                for method in ('_save_master_session', '_prepare_reports_for_user_context_change',
                               '_refresh_ui', '_reload_user_style_settings', '_load_messages',
                               '_apply_all_plugin_states', '_restore_session_ui',
                               '_refresh_master_menu_states', '_update_master_status_bar',
                               '_apply_master_button_states'):
                    setattr(app, method, lambda *_a, **_k: None)
                app._on_select = lambda: received.append(list(app.selected_animals))
                app._load_disabled_plugins = lambda: set()
                app._get_global_language_fallback = lambda: 'en'
                app.master_track.logout = lambda **_kwargs: None
                app.master_track.login_interactive = lambda **_kwargs: path == 'login-success'
                app.master_track.load_session = lambda: {}
                if path == 'logout':
                    app._do_master_logout()
                else:
                    app._do_master_login()
                expected = ['Subject'] if path == 'login-cancelled' else []
                self.assertEqual(app.selected_animals, expected)
                self.assertEqual(app._selected_heritage_only, ['Dummy'] if expected else [])
                self.assertEqual(received, [] if expected else [[]])

    def test_273_project_heading_never_uses_origin_or_identity(self):
        record = {'name': 'Subject', 'origin': 'DPZ', 'ipid': 'DPZ:Subject', 'project': ''}
        self.assertEqual(MAIN.ProgTrackApp._plot_statistics_label(record, 'Count: 1'), 'Count: 1')
        record['project'] = 'Project A'
        self.assertEqual(MAIN.ProgTrackApp._plot_statistics_label(record, 'Count: 1'), 'Project A: Count: 1')
        record['project'] = ''
        record['project_history'] = [{'project': 'Project A'}]
        self.assertEqual(MAIN.ProgTrackApp._plot_statistics_label(record, 'Count: 1'), 'Count: 1')


if __name__ == '__main__':
    unittest.main()
